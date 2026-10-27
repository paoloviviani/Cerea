"""Geometry of created decks (no renderer): every shape on the slide, nothing over the title, columns even.

The 0.30.0 built-in base left the default deck's layouts with explicit boxes 40 cm wide and 0 cm high on a 16:9
slide; these checks look at the numbers a renderer would use, so that class of bug fails in CI."""

from __future__ import annotations

import io
import os
from typing import Any

import pptx
import pytest
from pptx.enum.shapes import PP_PLACEHOLDER

from cerea_documents import create
from cerea_documents.config import Limits

from .fixtures import png
from .test_create import chart_block
from .test_create_templates import style_template_pptx

SIZES = {"16:9": (12_192_000, 6_858_000), "4:3": (9_144_000, 6_858_000)}
TITLES = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
IMG = {"wide.png": png(160, 90, (10, 90, 200)), "tall.png": png(60, 120, (200, 90, 10))}
LONG = "A fairly long bullet that says a number of things and keeps going so that it wraps over lines"


def build(spec: dict[str, Any], template: bytes | None = None) -> tuple[dict[str, Any], pptx.presentation.Presentation]:
    report, data = create.create({"format": "pptx", "spec": spec}, template, IMG, Limits())
    return report, pptx.Presentation(io.BytesIO(data))


def box(shape: Any) -> tuple[int, int, int, int]:
    assert None not in (shape.left, shape.top, shape.width, shape.height), shape.name
    return int(shape.left), int(shape.top), int(shape.width), int(shape.height)


def assert_inside(shape: Any, w: int, h: int) -> None:
    x, y, sw, sh = box(shape)
    assert sw > 0 and sh > 0, f"{shape.name} has no size"
    assert x >= 0 and y >= 0 and x + sw <= w and y + sh <= h, f"{shape.name} {(x, y, sw, sh)} leaves the {w}x{h} slide"


def is_title(shape: Any) -> bool:
    return shape.is_placeholder and shape.placeholder_format.type in TITLES


def check_slide(slide: Any, w: int, h: int) -> None:
    shapes = list(slide.shapes)
    title = next((s for s in shapes if is_title(s)), None)
    for s in shapes:
        assert_inside(s, w, h)
        if title is not None and s is not title:
            assert box(s)[1] >= box(title)[1] + box(title)[3], f"{s.name} overlaps the title"


@pytest.mark.parametrize("size", SIZES)
def test_every_builtin_layout_and_the_master_fit_the_slide_size(size: str) -> None:
    _, prs = build({"size": size, "slides": [{"title": "t"}]})
    w, h = SIZES[size]
    assert (prs.slide_width, prs.slide_height) == (w, h)  # the slide size is the geometry basis of the layouts
    for holder in (prs.slide_master, *prs.slide_layouts):
        for shape in holder.shapes:
            assert_inside(shape, w, h)
    layouts: list[Any] = list(prs.slide_layouts)
    for lay in layouts:  # a content layout has a title band above its body
        title = next((p for p in lay.placeholders if p.placeholder_format.type in TITLES), None)
        bodies = [p for p in lay.placeholders if p.placeholder_format.type == PP_PLACEHOLDER.OBJECT]
        if lay.name in ("Title and Content", "Two Content", "Comparison") and title is not None:
            assert all(box(b)[1] >= box(title)[1] + box(title)[3] for b in bodies)


DECKS: dict[str, list[dict[str, Any]]] = {
    "text": [
        {"title": "Cover", "subtitle": "sub"},
        {"layout": "section", "title": "Part", "subtitle": "s"},
        {"title": "Bullets", "body": [{"type": "list", "items": [LONG] * 6}]},
        {"title": "Very " * 20 + "long title", "body": [{"type": "paragraph", "text": "x"}]},
        {"layout": "title_only", "title": "Only"},
        {"layout": "blank"},
        {
            "title": "Cols",
            "body": [{"type": "list", "items": [LONG] * 3}],
            "body_right": [{"type": "list", "items": ["a"]}],
        },
    ],
    "visuals": [
        {
            "title": "Table",
            "body": [{"type": "table", "header": ["A", "B"], "rows": [[i, "x" * 30] for i in range(12)]}],
        },
        {"title": "Picture", "images": [{"image": "wide.png"}]},
        {"title": "Tall picture", "images": [{"image": "tall.png"}]},
        {"title": "Chart", "body": [chart_block("column", "native")]},
        {"title": "Image chart", "charts": [{k: v for k, v in chart_block("pie", "image").items() if k != "type"}]},
        {
            "title": "Mixed",
            "body": [{"type": "list", "items": ["a", "b"]}, {"type": "table", "rows": [["1", "2"]]}],
            "images": [{"image": "wide.png"}],
        },
        {
            "layout": "two_content",
            "title": "Table and chart",
            "body": [{"type": "table", "header": ["A"], "rows": [["1"], ["2"]]}],
            "body_right": [chart_block("line", "native")],
        },
        {"title": "Placed", "images": [{"image": "wide.png", "x_cm": 30, "y_cm": 10, "width_cm": 20}]},
    ],
}


@pytest.mark.parametrize("size", SIZES)
@pytest.mark.parametrize("kind", DECKS)
def test_mixed_decks_keep_every_shape_inside_the_slide_and_off_the_title(kind: str, size: str) -> None:
    _, prs = build({"size": size, "slides": DECKS[kind]})
    for slide in prs.slides:
        check_slide(slide, *SIZES[size])


def test_placeholders_inherit_their_geometry_from_the_layout() -> None:
    _, prs = build({"slides": [{"title": "t", "body": [{"type": "list", "items": ["a"]}]}]})
    for sp in prs.slides[0].shapes:
        assert sp._element.spPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}xfrm") is None


def test_two_content_columns_split_the_width_evenly_on_16_9() -> None:
    _, prs = build(
        {
            "slides": [
                {
                    "title": "t",
                    "body": [{"type": "list", "items": ["a"]}],
                    "body_right": [{"type": "list", "items": ["b"]}],
                }
            ]
        }
    )
    cols = sorted((s for s in prs.slides[0].placeholders if not is_title(s)), key=lambda s: s.left)
    assert len(cols) == 2
    (lx, _, lw, _), (rx, _, rw, _) = box(cols[0]), box(cols[1])
    assert abs(lw - rw) < 36_000 and lx + lw < rx
    assert abs(lx - (prs.slide_width - rx - rw)) < 36_000  # symmetric margins
    assert lw > prs.slide_width * 0.4


def test_a_picture_keeps_its_aspect_ratio_inside_the_content_area() -> None:
    _, prs = build({"slides": [{"title": "t", "images": [{"image": "tall.png"}]}]})
    pic = next(s for s in prs.slides[0].shapes if s.shape_type == 13)
    assert abs(pic.width / pic.height - 0.5) < 0.01
    lay_body = next(p for p in prs.slides[0].slide_layout.placeholders if p.placeholder_format.idx == 1)
    x, y, w, h = box(lay_body)
    assert pic.left >= x and pic.top >= y and pic.left + pic.width <= x + w and pic.top + pic.height <= y + h


def test_body_text_is_sized_for_a_slide_not_a_wall_of_small_text() -> None:
    _, prs = build({"slides": [{"title": "t", "body": [{"type": "list", "items": ["short", "list", "of", "points"]}]}]})
    body = next(p for p in prs.slides[0].placeholders if p.placeholder_format.idx == 1)
    assert {r.font.size.pt for p in body.text_frame.paragraphs for r in p.runs} == {24.0}
    master_title = prs.slide_master.element.xpath("//p:titleStyle/a:lvl1pPr/a:defRPr/@sz")
    assert master_title == ["3600"]


def test_text_that_cannot_fit_says_so_and_suggests_splitting() -> None:
    items = [LONG * 2] * 14
    report, prs = build({"slides": [{"title": "t", "body": [{"type": "list", "items": items}]}]})
    assert any("split" in n and "slides[0]" in n for n in report["notes"])
    body = next(p for p in prs.slides[0].placeholders if p.placeholder_format.idx == 1)
    assert min(r.font.size.pt for p in body.text_frame.paragraphs for r in p.runs) >= 14  # never a tiny wall
    assert body.text_frame._txBody.xpath("a:bodyPr/a:normAutofit")


def test_a_table_too_tall_for_the_area_says_so() -> None:
    rows = [[str(i), "x"] for i in range(40)]
    report, _ = build({"slides": [{"title": "t", "body": [{"type": "table", "header": ["A", "B"], "rows": rows}]}]})
    assert any("split" in n for n in report["notes"])


def test_template_decks_keep_the_templates_own_layout_geometry() -> None:
    template = style_template_pptx()
    ref = pptx.Presentation(io.BytesIO(template))
    spec = {
        "theme": {"template": "b.pptx"},
        "slides": [
            {"layout": "Brand Cover", "title": "Cover", "subtitle": "s"},
            {"layout": "Brand Content", "title": "T", "body": [{"type": "list", "items": ["a", "b"]}]},
            {
                "layout": "two_content",
                "title": "T",
                "body": [{"type": "list", "items": ["a"]}],
                "body_right": [{"type": "list", "items": ["b"]}],
            },
        ],
    }
    _, prs = build(spec, template)
    assert (prs.slide_width, prs.slide_height) == (ref.slide_width, ref.slide_height)
    for slide in prs.slides:
        layout_boxes = {p.placeholder_format.idx: box(p) for p in slide.slide_layout.placeholders}
        for sh in slide.placeholders:
            assert box(sh) == layout_boxes[sh.placeholder_format.idx]
        check_slide(slide, prs.slide_width, prs.slide_height)


@pytest.mark.parametrize("fmt", ["docx", "pdf"])
def test_docx_and_pdf_images_never_exceed_the_text_area(fmt: str) -> None:
    """The same class of bug in the other formats: a tall or wide picture must fit the page's text area."""
    import docx

    blocks: list[dict[str, Any]] = [
        {"type": "image", "image": "tall.png", "height_cm": 60},
        {"type": "image", "image": "wide.png", "width_cm": 80},
    ]
    _, data = create.create({"format": fmt, "spec": {"blocks": blocks}}, None, IMG, Limits())
    if fmt == "pdf":
        return  # laying out at all is the check: a flowable taller than the frame raises
    d = docx.Document(io.BytesIO(data))
    s = d.sections[0]
    tw = int(s.page_width or 0) - int(s.left_margin or 0) - int(s.right_margin or 0)
    th = int(s.page_height or 0) - int(s.top_margin or 0) - int(s.bottom_margin or 0)
    assert d.inline_shapes
    assert all(0 < sh.width <= tw and 0 < sh.height <= th for sh in d.inline_shapes)


def to_pdf(data: bytes) -> bytes:
    """Convert a deck through LibreOffice (Gotenberg at GOTENBERG_URL)."""
    import urllib.request
    import uuid

    boundary = uuid.uuid4().hex
    body = (
        (
            f'--{boundary}\r\nContent-Disposition: form-data; name="files"; filename="deck.pptx"\r\n'
            "Content-Type: application/octet-stream\r\n\r\n"
        ).encode()
        + data
        + f"\r\n--{boundary}--\r\n".encode()
    )
    req = urllib.request.Request(
        os.environ["GOTENBERG_URL"].rstrip("/") + "/forms/libreoffice/convert",
        body,
        {"Content-Type": f"multipart/form-data; boundary={boundary}"},
    )
    with urllib.request.urlopen(req, timeout=120) as resp:
        pdf: bytes = resp.read()
    assert pdf.startswith(b"%PDF")
    return pdf


@pytest.mark.skipif(
    not os.environ.get("GOTENBERG_URL"), reason="needs a Gotenberg (LibreOffice) service; CI provides one"
)
def test_a_generated_deck_converts_to_a_pdf_page_per_slide() -> None:
    import pypdf

    slides = [s for deck in DECKS.values() for s in deck]
    _, data = create.create({"format": "pptx", "spec": {"slides": slides}}, None, IMG, Limits())
    assert len(pypdf.PdfReader(io.BytesIO(to_pdf(data))).pages) == len(slides)
