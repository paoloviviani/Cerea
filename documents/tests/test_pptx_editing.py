"""Whole-slide editing of PowerPoint files (replace_slide, add_slide, update_shape), the fit estimate for
shapes, readable array limits, titles and slide numbers that follow the theme."""

from __future__ import annotations

import io
from typing import Any

import pptx
import pytest

from cerea_documents import core, create
from cerea_documents.config import Limits
from cerea_documents.errors import DocumentError, InvalidRequest

from .helpers import do_edit, refuses
from .test_create_templates import style_template_pptx
from .test_pptx_shapes import CM, IMG, build, rect

DARK = {"slide_background": {"color": "0F172A"}, "colors": {"text": "E2E8F0", "heading": "FFFFFF", "muted": "94A3B8"}}


def label(**kw: Any) -> dict[str, Any]:
    return {"kind": "text", "x_cm": 2, "y_cm": 5, "w_cm": 6, "h_cm": 0.6, "text": "Label", "size_pt": 10, **kw}


def fit_notes(report: dict[str, Any]) -> list[str]:
    return [n for n in report.get("notes", []) if "will not fit" in n]


# ---- the fit estimate for shapes -----------------------------------------------------------------


def test_a_one_line_10pt_label_in_a_short_box_is_not_a_warning() -> None:
    report, _, _ = build([{"layout": "blank", "shapes": [label()]}])
    assert not fit_notes(report)
    # the estimate assumes PowerPoint's default insets (0.254 and 0.127 cm); none is overridden in the file
    _, prs, _ = build([{"layout": "blank", "shapes": [label()]}])
    body = prs.slides[0].shapes[0]._element.xpath("p:txBody/a:bodyPr")[0]
    assert [body.get(k) for k in ("lIns", "rIns", "tIns", "bIns")] == [None] * 4


def test_a_short_label_that_runs_past_its_box_over_nothing_is_not_a_warning() -> None:
    long = label(text="word " * 16, w_cm=4, h_cm=0.6)  # three lines in a 0.6 cm box, with empty slide around it
    report, _, _ = build([{"layout": "blank", "shapes": [long]}])
    assert not fit_notes(report)


def test_a_paragraph_that_overflows_into_the_text_below_it_warns_once_with_the_height_needed() -> None:
    para = label(
        text="A sentence that goes on and on over several lines of a narrow box. " * 2, w_cm=7, h_cm=1.2, size_pt=14
    )
    below = label(y_cm=6.6, text="Next label")
    report, _, _ = build([{"layout": "blank", "shapes": [para, below]}])
    notes = fit_notes(report)
    assert len(notes) == 1 and notes[0].startswith("slides[0].shapes[0]:")
    assert "cm of height needed" in notes[0] and "overlap the text of shapes[1]" in notes[0]
    assert "(8 at least)" in notes[0]


def test_text_that_overflows_its_filled_card_warns() -> None:
    card = rect(w_cm=5, h_cm=2, text="Card text that is much longer than the card can hold at all. " * 2, size_pt=14)
    report, _, _ = build([{"layout": "blank", "shapes": [card]}])
    notes = fit_notes(report)
    assert len(notes) == 1 and "spill out of its filled shape" in notes[0]


def test_text_in_a_box_that_sits_on_a_card_warns_when_it_runs_out_of_the_card() -> None:
    card = rect(w_cm=8, h_cm=3, fill="1E293B")
    inside = label(
        x_cm=2.5, y_cm=5.3, w_cm=7, h_cm=1.0, text="A sentence of some length that takes lines. " * 3, size_pt=14
    )
    report, _, _ = build([{"layout": "blank", "shapes": [card, inside]}])
    assert any("run out of shapes[0], the shape behind it" in n for n in fit_notes(report))


def test_text_that_would_leave_the_slide_warns() -> None:
    low = label(y_cm=18.5, w_cm=3, h_cm=0.5, text="word " * 30, size_pt=12)
    report, _, _ = build([{"layout": "blank", "shapes": [low]}])
    assert any("run off the slide" in n for n in fit_notes(report))


def test_a_run_size_counts_in_the_estimate() -> None:
    big = {
        "kind": "rect",
        "x_cm": 2,
        "y_cm": 5,
        "w_cm": 8,
        "h_cm": 1.2,
        "runs": [{"text": "Big", "size_pt": 40}],
        "size_pt": 10,
    }
    report, _, _ = build([{"layout": "blank", "shapes": [big]}])
    assert fit_notes(report)  # 40 pt does not fit 1.2 cm; at the shape's 10 pt it would


# ---- the spec conveniences -----------------------------------------------------------------------


def test_a_shape_colour_is_the_default_colour_of_its_text_and_a_run_size_overrides_the_shape_size() -> None:
    shape = rect(
        fill="FFFFFF",
        color="C0504D",
        size_pt=12,
        runs=[{"text": "Big ", "size_pt": 20}, {"text": "own", "color": "00AA00"}, {"text": " rest"}],
    )
    _, prs, _ = build([{"layout": "blank", "shapes": [shape]}])
    runs = prs.slides[0].shapes[0].text_frame.paragraphs[0].runs
    assert [r.font.size.pt for r in runs] == [20, 12, 12]
    assert [str(r.font.color.rgb) for r in runs] == ["C0504D", "00AA00", "C0504D"]


def test_shape_text_may_be_8_pt() -> None:
    build([{"layout": "blank", "shapes": [label(size_pt=8)]}])
    with pytest.raises(InvalidRequest, match="size_pt"):
        build([{"layout": "blank", "shapes": [label(size_pt=7)]}])
    with pytest.raises(InvalidRequest, match="size_pt"):
        build([{"layout": "blank", "shapes": [rect(runs=[{"text": "x", "size_pt": 7}])]}])


def test_title_and_content_is_an_alias_of_title_content() -> None:
    _, prs, _ = build([{"layout": "title_and_content", "title": "T", "body": [{"type": "paragraph", "text": "b"}]}])
    assert prs.slides[0].slide_layout.name == "Title and Content"
    _, prs, _ = build(
        [{"layout": "title_and_content", "title": "T", "body": [{"type": "paragraph", "text": "b"}]}],
        theme={"template": "t.pptx"},
        template=style_template_pptx(),
    )
    assert prs.slides[0].slide_layout.name == "Brand Content"


# ---- readable limits -------------------------------------------------------------------------------


def test_too_many_shapes_is_a_short_sentence_naming_the_slide_and_the_limit() -> None:
    shapes = [rect(x_cm=1 + (i % 10), y_cm=1 + i // 10, w_cm=0.5, h_cm=0.5) for i in range(41)]
    with pytest.raises(InvalidRequest) as ei:
        build([{"layout": "blank"}] * 9 + [{"layout": "blank", "shapes": shapes}])
    msg = str(ei.value)
    assert msg == "slides[9].shapes: 41 shapes; 40 at most per slide"
    assert "kind" not in msg and len(msg) < 100


def test_forty_shapes_are_accepted() -> None:
    shapes = [rect(x_cm=1 + (i % 10), y_cm=1 + i // 10, w_cm=0.5, h_cm=0.5) for i in range(40)]
    build([{"layout": "blank", "shapes": shapes}])


def test_a_minimum_count_is_a_short_sentence_too() -> None:
    with pytest.raises(InvalidRequest) as ei:
        create.create({"format": "pptx", "spec": {"slides": []}}, None, {}, Limits())
    assert str(ei.value) == "spec: slides: 0 slides; at least 1 needed"


def test_a_big_value_of_the_wrong_type_is_described_never_echoed() -> None:
    big = ["x" * 50] * 200
    with pytest.raises(InvalidRequest) as ei:
        build([{"layout": "blank", "title": big}])
    assert "a list of 200 items" in str(ei.value) and "xxxxx" not in str(ei.value)
    with pytest.raises(InvalidRequest) as ei:
        build([{"layout": "blank", "shapes": [rect(fill={"k": "v" * 100})]}])
    assert "vvvvv" not in str(ei.value)
    with pytest.raises(InvalidRequest) as ei:
        build([{"layout": "blank", "shapes": [rect(valign="x" * 500)]}])
    assert "xxxxx" not in str(ei.value)
    with pytest.raises(InvalidRequest) as ei:
        build([{"layout": "blank", "title": "t" * 600}])
    assert "ttttt" not in str(ei.value) and "600 characters" in str(ei.value)


def test_a_big_value_in_an_edit_operation_is_not_echoed_either() -> None:
    deck = build([{"layout": "blank"}])[2]
    with pytest.raises(DocumentError) as ei:
        do_edit(deck, [{"op": "replace_text", "find": ["y" * 80] * 40, "replace": "x"}], fmt="pptx")
    assert "yyyyy" not in str(ei.value) and "a list of 40 items" in str(ei.value)


# ---- titles and numbers follow the theme ----------------------------------------------------------


def master_title(prs: Any) -> Any:
    return prs.slide_master._element.xpath("p:txStyles/p:titleStyle/a:lvl1pPr/a:defRPr")[0]


def test_the_master_title_takes_the_themes_h1_size() -> None:
    slides = [{"layout": "title_content", "title": "A title", "body": [{"type": "paragraph", "text": "b"}]}]
    _, prs, _ = build(slides, theme={"heading_sizes_pt": {"h1": 28}})
    assert master_title(prs).get("sz") == "2800"
    assert prs.slides[0].shapes.title.text_frame.paragraphs[0].runs[0].font.size is None  # inherits the master's
    _, prs, _ = build(slides)
    assert master_title(prs).get("sz") == "3600" and master_title(prs).get("b") == "1"


def test_a_long_title_still_steps_down_from_the_themes_size() -> None:
    long = "A title long enough that it cannot stay on two lines at the larger size " * 2
    _, prs, _ = build([{"layout": "title_content", "title": long}], theme={"heading_sizes_pt": {"h1": 28}})
    size = prs.slides[0].shapes.title.text_frame.paragraphs[0].runs[0].font.size
    assert size is not None and size.pt == 24


def test_heading_bold_false_makes_slide_titles_plain() -> None:
    _, prs, _ = build([{"layout": "title_only", "title": "T"}], theme={"heading_bold": False})
    assert master_title(prs).get("b") == "0"


def numbers(prs: Any) -> list[list[str]]:
    return [s._element.xpath("//a:fld[@type='slidenum']/a:t/text()") for s in prs.slides]


def test_slide_numbers_are_a_field_on_every_slide_but_the_first_in_the_muted_colour() -> None:
    slides = [{"layout": "title", "title": "Cover"}, {"layout": "blank"}, {"layout": "title_content", "title": "T"}]
    _, prs, _ = build(slides, theme={"slide_numbers": True, "colors": {"muted": "94A3B8"}})
    assert numbers(prs) == [[], ["2"], ["3"]]
    fld = prs.slides[1]._element.xpath("//a:fld[@type='slidenum']")[0]
    assert fld.xpath("a:rPr/@sz") == ["1000"]
    assert fld.xpath("a:rPr/a:solidFill/a:srgbClr/@val") == ["94A3B8"]
    ph = prs.slides[1]._element.xpath("//p:ph[@type='sldNum']")
    assert len(ph) == 1  # the layout's own slide-number placeholder, so the master positions it bottom right
    _, plain, _ = build(slides)
    assert numbers(plain) == [[], [], []]


def test_a_template_without_a_number_placeholder_gets_a_text_box_with_the_field() -> None:
    template = pptx.Presentation(io.BytesIO(style_template_pptx()))
    for lay in template.slide_layouts:
        for ph in list(lay.placeholders):
            if ph.placeholder_format.type == pptx.enum.shapes.PP_PLACEHOLDER.SLIDE_NUMBER:
                ph._element.getparent().remove(ph._element)
    out = io.BytesIO()
    template.save(out)
    slides = [{"layout": "Brand Cover", "title": "Cover"}, {"layout": "Brand Content", "title": "T"}]
    _, prs, _ = build(slides, theme={"template": "t.pptx", "slide_numbers": True}, template=out.getvalue())
    assert numbers(prs) == [[], ["2"]]
    box = next(s for s in prs.slides[1].shapes if s.name == "Slide Number")
    assert box.left + box.width <= prs.slide_width and box.top + box.height <= prs.slide_height
    assert box.left > prs.slide_width // 2 and box.top > prs.slide_height // 2  # bottom right


def test_a_template_with_its_own_number_placeholder_keeps_its_look() -> None:
    slides = [{"layout": "Brand Cover", "title": "Cover"}, {"layout": "Brand Content", "title": "T"}]
    _, prs, _ = build(slides, theme={"template": "t.pptx", "slide_numbers": True}, template=style_template_pptx())
    assert numbers(prs) == [[], ["2"]]
    fld = prs.slides[1]._element.xpath("//a:fld[@type='slidenum']")[0]
    assert not fld.xpath("a:rPr/@sz")


# ---- replace_slide, add_slide, update_shape ------------------------------------------------------


def cards(n: int = 3) -> list[dict[str, Any]]:
    return [
        {
            "kind": "rounded_rect", "x_cm": 1.6 + i * 10.4, "y_cm": 5, "w_cm": 9.6, "h_cm": 6, "fill": "1E293B",
            "runs": [{"text": f"Step {i}\n", "bold": True}, {"text": "A short sentence"}], "size_pt": 16, "color": "F8FAFC",
        }
        for i in range(n)
    ]  # fmt: skip


def deck(**theme: Any) -> bytes:
    slides = [
        {"layout": "title", "title": "Cover", "notes": "cover notes"},
        {"layout": "blank", "shapes": [label(), *cards()], "notes": "keep these notes"},
        {"layout": "title_content", "title": "Plain", "body": [{"type": "list", "items": ["a", "b"]}]},
    ]
    return build(slides, theme={**DARK, "slide_numbers": True, **theme})[2]


def prs_of(data: bytes) -> Any:
    return pptx.Presentation(io.BytesIO(data))


NEW = {
    "layout": "blank",
    "shapes": [
        {"kind": "text", "x_cm": 1.6, "y_cm": 1.5, "w_cm": 20, "h_cm": 1.5, "text": "Two columns", "size_pt": 28},
        {"kind": "rect", "x_cm": 1.6, "y_cm": 4, "w_cm": 14, "h_cm": 10, "fill": "334155", "text": "Left"},
        label(y_cm=15),
    ],
}


def test_replace_slide_rebuilds_one_slide_with_the_decks_theme_and_leaves_the_rest_as_it_was() -> None:
    src = deck()
    rep, out = do_edit(src, [{"op": "replace_slide", "slide": 2, "spec": NEW}], fmt="pptx")
    assert rep["verify"]["changed"] == ["ppt/slides/slide2.xml"] and not rep["verify"]["added"]
    p = prs_of(out)
    s = p.slides[1]
    assert [sh.name for sh in s.shapes] == ["text 0", "rect 1", "text 2", "Slide Number Placeholder"]
    assert s.slide_layout.name == "Blank" and len(p.slides) == 3
    # the theme came from the file: background, text colour, number colour and the number itself
    assert s._element.xpath("p:cSld/p:bg/p:bgPr/a:solidFill/a:srgbClr/@val") == ["0F172A"]
    assert s.shapes[0].text_frame.paragraphs[0].runs[0].font.color.rgb == pptx.dml.color.RGBColor.from_string("E2E8F0")
    assert numbers(p)[1] == ["2"]
    fld = s._element.xpath("//a:fld[@type='slidenum']/a:rPr/a:solidFill/a:srgbClr/@val")
    assert fld == ["94A3B8"]
    assert s.notes_slide.notes_text_frame.text == "keep these notes"  # kept: the spec gave none
    r = rep["operations"][0]["result"]
    assert r["theme_from_file"]["style"] == "built-in" and r["theme_from_file"]["defaults"] == []
    assert "notes" not in r  # nothing to warn about: the small label sits clear of the shape


def test_replace_slide_takes_new_notes_and_runs_the_same_checks_as_doc_create() -> None:
    src = deck()
    _, out = do_edit(src, [{"op": "replace_slide", "slide": 2, "spec": {**NEW, "notes": "new notes"}}], fmt="pptx")
    assert prs_of(out).slides[1].notes_slide.notes_text_frame.text == "new notes"
    too_many = {"layout": "blank", "shapes": [rect(w_cm=0.5, h_cm=0.5)] * 41}
    refuses(
        src,
        [{"op": "replace_slide", "slide": 2, "spec": too_many}],
        "slide 2.shapes: 41 shapes; 40 at most per slide",
        fmt="pptx",
    )
    off = {"layout": "blank", "shapes": [rect(x_cm=30, w_cm=6, text="x")]}
    refuses(
        src,
        [{"op": "replace_slide", "slide": 2, "spec": off}],
        "slide 2.shapes[0] (rect) reaches outside the slide",
        fmt="pptx",
    )
    refuses(src, [{"op": "replace_slide", "slide": 9, "spec": NEW}], "3 slides", fmt="pptx")
    refuses(src, [{"op": "replace_slide", "slide": 2, "spec": {"layout": "nope"}}], "unknown layout", fmt="pptx")


def test_replace_slide_reports_fit_warnings_with_the_slide_in_the_result() -> None:
    bad = {"layout": "blank", "shapes": [rect(w_cm=4, h_cm=1.5, text="word " * 40, size_pt=14)]}
    rep, _ = do_edit(deck(), [{"op": "replace_slide", "slide": 2, "spec": bad}], fmt="pptx")
    notes = rep["operations"][0]["result"]["notes"]
    assert len(notes) == 1 and notes[0].startswith("slide 2.shapes[0]: the text will not fit")


def test_replace_slide_keeps_position_and_sections_and_replaces_media_and_charts() -> None:
    src = deck()
    chart = {
        "layout": "title_content", "title": "Numbers",
        "body": [{"type": "chart", "chart_type": "bar", "categories": ["a", "b"], "series": [{"name": "s", "values": [1, 2]}]}],
    }  # fmt: skip
    rep, out = do_edit(src, [{"op": "replace_slide", "slide": 3, "spec": chart}], fmt="pptx")
    assert any(n.startswith("ppt/charts/chart") for n in rep["verify"]["added"])
    assert [s.shapes.title.text if s.shapes.title else None for s in prs_of(out).slides] == ["Cover", None, "Numbers"]
    # replacing it again drops the first chart: nothing is left behind
    rep2, out2 = do_edit(
        out, [{"op": "replace_slide", "slide": 3, "spec": {"layout": "title_only", "title": "No chart"}}], fmt="pptx"
    )
    assert not [n for n in core.outline(out2, "pptx", Limits())["slides"][2]["shapes"] if n["kind"] == "chart"]
    assert any(n.startswith("ppt/charts/") for n in rep2["verify"]["removed"])


def test_replace_slide_with_a_picture_background_or_shape_uses_the_files_images() -> None:
    spec = {"layout": "blank", "background": {"image": "wide.png", "fit": "cover"}, "shapes": [label()]}
    rep, out = do_edit(
        deck(), [{"op": "replace_slide", "slide": 2, "spec": spec}], fmt="pptx", images={"wide.png": IMG["wide.png"]}
    )
    assert any(n.startswith("ppt/media/") for n in rep["verify"]["added"])
    assert prs_of(out).slides[1]._element.xpath("p:cSld/p:bg/p:bgPr/a:blipFill")


def test_replace_slide_on_a_deck_without_slide_numbers_adds_none_and_on_the_first_slide_none() -> None:
    src = build([{"layout": "title", "title": "Cover"}, {"layout": "blank"}])[2]
    _, out = do_edit(src, [{"op": "replace_slide", "slide": 2, "spec": NEW}], fmt="pptx")
    assert numbers(prs_of(out)) == [[], []]
    src = deck()
    _, out = do_edit(
        src, [{"op": "replace_slide", "slide": 1, "spec": {"layout": "title", "title": "New cover"}}], fmt="pptx"
    )
    assert numbers(prs_of(out))[0] == []


def test_replace_slide_in_a_templated_deck_uses_the_files_own_layouts() -> None:
    slides = [{"layout": "Brand Cover", "title": "Cover"}, {"layout": "Brand Content", "title": "T"}]
    data = build(slides, theme={"template": "t.pptx", "slide_numbers": True}, template=style_template_pptx())[2]
    rep, out = do_edit(
        data,
        [
            {
                "op": "replace_slide",
                "slide": 2,
                "spec": {"layout": "Brand Content", "title": "Again", "body": [{"type": "paragraph", "text": "x"}]},
            }
        ],
        fmt="pptx",
    )
    s = prs_of(out).slides[1]
    assert s.slide_layout.name == "Brand Content" and s.shapes.title.text == "Again"
    assert rep["operations"][0]["result"]["theme_from_file"]["style"] == "template"
    assert numbers(prs_of(out))[1] == ["2"]
    refuses(
        data,
        [{"op": "replace_slide", "slide": 2, "spec": {"layout": "Nope"}}],
        "the template has no layout",
        fmt="pptx",
    )


def test_replace_slide_says_when_a_theme_value_could_not_be_recovered() -> None:
    src = build([{"layout": "title", "title": "Cover"}, {"layout": "blank"}])[2]  # no body text to read a colour from
    rep, _ = do_edit(src, [{"op": "replace_slide", "slide": 2, "spec": {"layout": "blank"}}], fmt="pptx")
    defaults = rep["operations"][0]["result"]["theme_from_file"]["defaults"]
    assert any("text colour" in d for d in defaults) and any("background" in d for d in defaults)


def test_add_slide_inserts_after_a_slide_in_the_same_section_and_numbers_it() -> None:
    src = deck()
    rep, out = do_edit(
        src,
        [
            {
                "op": "add_slide",
                "after": 1,
                "spec": {
                    "layout": "title_and_content",
                    "title": "Added",
                    "body": [{"type": "paragraph", "text": "hi"}],
                    "notes": "added notes",
                },
            }
        ],
        fmt="pptx",
    )
    r = rep["operations"][0]["result"]
    assert r["new_slide_position"] == 2
    p = prs_of(out)
    assert [s.slide_layout.name for s in p.slides] == ["Title Slide", "Title and Content", "Blank", "Title and Content"]
    assert p.slides[1].shapes.title.text == "Added" and p.slides[1].notes_slide.notes_text_frame.text == "added notes"
    assert numbers(p)[1] == ["2"] and rep["verify"]["ok"]
    _, first = do_edit(src, [{"op": "add_slide", "after": 0, "spec": {"layout": "title", "title": "Zero"}}], fmt="pptx")
    assert prs_of(first).slides[0].shapes.title.text == "Zero" and numbers(prs_of(first))[0] == []
    refuses(src, [{"op": "add_slide", "after": 7, "spec": NEW}], "3 slides", fmt="pptx")


def test_slides_added_in_one_request_stay_behind_their_slide_when_reordered() -> None:
    _, out = do_edit(
        deck(),
        [
            {"op": "add_slide", "after": 1, "spec": {"layout": "title_only", "title": "A"}},
            {"op": "add_slide", "after": 0, "spec": {"layout": "title_only", "title": "Z"}},
            {"op": "reorder_slides", "order": [2, 1, 3]},
        ],
        fmt="pptx",
    )
    titles = [s.shapes.title.text if s.shapes.title else "-" for s in prs_of(out).slides]
    assert titles == ["Z", "-", "Cover", "A", "Plain"]


def test_update_shape_moves_and_resizes_by_shape_id_and_keeps_text_on_the_slide() -> None:
    src = deck()
    outline = core.outline(src, "pptx", Limits())["slides"][1]["shapes"]
    card = next(s for s in outline if s["name"] == "rounded_rect 1")
    rep, out = do_edit(src, [{"op": "update_shape", "slide": 2, "shape": card["id"], "x_cm": 2, "w_cm": 5}], fmt="pptx")
    moved = next(s for s in prs_of(out).slides[1].shapes if s.shape_id == card["id"])
    assert (moved.left, moved.width, moved.top, moved.height) == (2 * CM, 5 * CM, card["box"][1], card["box"][3])
    assert rep["verify"]["changed"] == ["ppt/slides/slide2.xml"]
    refuses(
        src,
        [{"op": "update_shape", "slide": 2, "shape": card["id"], "x_cm": 33}],
        "reaches outside the slide",
        fmt="pptx",
    )
    refuses(src, [{"op": "update_shape", "slide": 2, "shape": 999, "x_cm": 1}], "no shape with id 999", fmt="pptx")
    refuses(src, [{"op": "update_shape", "slide": 2, "shape": card["id"]}], "at least one of", fmt="pptx")


def test_replace_slide_and_add_slide_leave_every_other_part_byte_identical() -> None:
    from .fixtures import parts

    src = deck()
    _, out = do_edit(src, [{"op": "replace_slide", "slide": 2, "spec": NEW}], fmt="pptx")
    a, b = parts(src), parts(out)
    assert {n for n in a if a[n] != b.get(n)} == {"ppt/slides/slide2.xml"} and set(b) == set(a)
    _, out = do_edit(src, [{"op": "add_slide", "after": 1, "spec": NEW}], fmt="pptx")
    b = parts(out)
    changed = {n for n in a if a[n] != b.get(n)}
    assert changed == {"[Content_Types].xml", "ppt/presentation.xml", "ppt/_rels/presentation.xml.rels"}
    assert set(b) - set(a) == {"ppt/slides/slide4.xml", "ppt/slides/_rels/slide4.xml.rels"}
