"""Fixture generators, ported from the in-sandbox helper checks. Nothing is committed as a binary:
every document is built here, then post-processed at the XML level to carry what the general
libraries would lose on load and save."""

from __future__ import annotations

import io
import re
import struct
import zipfile
import zlib

import docx
import openpyxl
import openpyxl.utils.datetime
import pptx
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from openpyxl.worksheet.datavalidation import DataValidation
from pptx.dml.color import RGBColor as PRGB
from pptx.util import Inches
from pptx.util import Pt as PPt

# ---- images -----------------------------------------------------------------------------


def png(w: int, h: int, rgb: tuple[int, int, int]) -> bytes:
    raw = b"".join(b"\x00" + bytes(rgb) * w for _ in range(h))

    def chunk(t: bytes, d: bytes) -> bytes:
        c = struct.pack(">I", len(d)) + t + d
        return c + struct.pack(">I", zlib.crc32(t + d))

    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def jpeg(w: int, h: int) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (w, h), (5, 5, 5)).save(buf, "JPEG")
    return buf.getvalue()


# ---- Word ---------------------------------------------------------------------------------

DOCX_CONTENT = {"word/document.xml", "word/header1.xml", "word/footer1.xml"}


def make_docx() -> bytes:
    d = docx.Document()
    d.add_heading("Statement of work", level=1)
    p = d.add_paragraph()
    a = p.add_run("Dear Ac")
    a.bold = True
    b = p.add_run("me L")
    b.italic = True
    b.font.color.rgb = RGBColor(0xC0, 0, 0)
    p.add_run("td, welcome")
    q = d.add_paragraph()
    x = q.add_run("Prepared for {{cli")
    x.bold = True
    y = q.add_run("ent}} on {{ date }}.")
    y.italic = True
    d.add_paragraph("Open points", style="List Bullet")
    c = d.add_paragraph()
    c.add_run("Review the ")
    r2 = c.add_run("pricing")
    r2.font.highlight_color = WD_COLOR_INDEX.YELLOW
    d.add_comment(r2, text="Check the numbers", author="Reviewer", initials="RV")
    d.add_paragraph("Keep {{other}} as it is.")
    t = d.add_table(rows=3, cols=2)
    t.style = "Table Grid"
    for i, row in enumerate(t.rows):
        for j, cell in enumerate(row.cells):
            run = cell.paragraphs[0].add_run(["Item", "Cost"][j] if i == 0 else ["Widget {{client}}", "10"][j])
            run.bold = i == 0
            run.font.size = Pt(9)
            if i == 2:
                run.font.color.rgb = RGBColor(0, 0x66, 0)
            tcpr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "DDEEFF" if i == 0 else "F2F2F2")
            tcpr.append(shd)
    d.add_paragraph("After the table.")
    sec = d.sections[0]
    sec.header.is_linked_to_previous = False
    hp = sec.header.paragraphs[0]
    h1 = hp.add_run("Report for {{cli")
    h1.bold = True
    hp.add_run("ent}}")
    sec.footer.is_linked_to_previous = False
    fp = sec.footer.paragraphs[0]
    f1 = fp.add_run("Confidential - {{foot")
    f1.italic = True
    fp.add_run("er_note}}")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_docx_with_image() -> bytes:
    d = docx.Document()
    d.add_paragraph("Before")
    d.add_picture(io.BytesIO(png(40, 20, (200, 30, 30))), width=Inches(2))
    d.add_paragraph("After")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ---- Excel ---------------------------------------------------------------------------------

SHEETS = ["Cover", "T1.1", "T2.4", "Lists", "Summary"]
TARGET = "T2.4"
MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
X14 = (
    '<extLst><ext uri="{78C0D931-6437-407d-A8EE-F0AAD7539E65}" '
    'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
    "<x14:conditionalFormattings><x14:conditionalFormatting "
    'xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main">'
    '<x14:cfRule type="dataBar" id="{A1B2C3D4-0000-0000-0000-%(n)012d}"><x14:dataBar '
    'minLength="0" maxLength="100"><x14:cfvo type="autoMin"/><x14:cfvo type="autoMax"/>'
    "</x14:dataBar></x14:cfRule><xm:sqref>B2:B30</xm:sqref></x14:conditionalFormatting>"
    "</x14:conditionalFormattings></ext>"
    '<ext uri="{CCE6A557-97BC-48b9-258B-7F2B2D6A1F5B}" '
    'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
    '<x14:dataValidations xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main" '
    'count="1"><x14:dataValidation type="list" allowBlank="1" showErrorMessage="1">'
    "<x14:formula1><xm:f>Lists!$A$2:$A$5</xm:f></x14:formula1><xm:sqref>E2:E20</xm:sqref>"
    "</x14:dataValidation></x14:dataValidations></ext></extLst>"
)


def make_xlsx(date1904: bool = False) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # type: ignore[arg-type]
    if date1904:
        wb.epoch = openpyxl.utils.datetime.CALENDAR_MAC_1904
    for n in SHEETS:
        ws = wb.create_sheet(n)
        ws["A1"] = "Header " + n
        for r in range(2, 30):
            ws.cell(r, 1, "item %d" % r)
            ws.cell(r, 2, r * 3)
            ws.cell(r, 3, "=B%d*2" % r)
        dv = DataValidation(type="list", formula1='"a,b,c"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add("D2:D20")
    t = wb[TARGET]
    t.merge_cells("A1:C1")
    for ref in ("E2", "E3"):
        t[ref].number_format = "dd/mm/yyyy"
    buf = io.BytesIO()
    wb.save(buf)

    with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as z:
        parts = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    out = []
    for name, d in parts:
        m = re.match(r"xl/worksheets/sheet(\d+)\.xml$", name)
        if m:
            k, s = int(m.group(1)), d.decode()
            s = s.replace("</worksheet>", X14 % {"n": k} + "</worksheet>")
            if name.endswith("sheet%d.xml" % (SHEETS.index(TARGET) + 1)):
                before = s
                s = re.sub(r'(<c r="C2"[^>]*>)<f>B2\*2</f>', r'\1<f t="shared" ref="C2:C10" si="0">B2*2</f>', s)
                for r in range(3, 11):
                    s = re.sub(r'(<c r="C%d"[^>]*>)<f>B%d\*2</f>' % (r, r), r'\1<f t="shared" si="0"/>', s)
                assert s != before and 't="shared"' in s, "shared formula injection failed"
            d = s.encode()
        out.append((name, d))
    chain = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><calcChain xmlns="{MAIN}">'
    for sid in range(1, len(SHEETS) + 1):
        for r in range(2, 30):
            chain += '<c r="C%d" i="%d"/>' % (r, sid) if r == 2 else '<c r="C%d"/>' % r
    chain += "</calcChain>"
    items = dict(out)
    items["[Content_Types].xml"] = (
        items["[Content_Types].xml"]
        .decode()
        .replace(
            "</Types>",
            '<Override PartName="/xl/calcChain.xml" ContentType="application/vnd.openxmlformats-'
            'officedocument.spreadsheetml.calcChain+xml"/></Types>',
        )
        .encode()
    )
    items["xl/_rels/workbook.xml.rels"] = (
        items["xl/_rels/workbook.xml.rels"]
        .decode()
        .replace(
            "</Relationships>",
            '<Relationship Id="rIdCC" Type="http://schemas.openxmlformats.org/officeDocument/'
            '2006/relationships/calcChain" Target="calcChain.xml"/></Relationships>',
        )
        .encode()
    )
    items["xl/calcChain.xml"] = chain.encode()
    order = [n for n, _ in out] + ["xl/calcChain.xml"]
    res = io.BytesIO()
    with zipfile.ZipFile(res, "w", zipfile.ZIP_DEFLATED) as z:
        for n in order:
            z.writestr(n, items[n])
    return res.getvalue()


# ---- PowerPoint -------------------------------------------------------------------------------

LOGO = png(40, 20, (200, 30, 30))
OWN = png(30, 30, (0, 120, 0))


def make_pptx() -> tuple[bytes, int]:
    prs = pptx.Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text_frame.text = "Annual review 2025"  # type: ignore[union-attr]
    run = s1.shapes.title.text_frame.paragraphs[0].runs[0]  # type: ignore[union-attr]
    run.font.bold = True
    run.font.color.rgb = PRGB(0x12, 0x34, 0x56)
    s1.placeholders[1].text_frame.text = "Prepared by the team"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Results"  # type: ignore[union-attr]
    body = s2.placeholders[1].text_frame
    body.text = "Revenue grew"
    body.add_paragraph().text = "Costs fell"
    box = s2.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(1))
    p = box.text_frame.paragraphs[0]
    r1 = p.add_run()
    r1.text = "Rev"
    r1.font.bold = True
    r1.font.size = PPt(20)
    r2 = p.add_run()
    r2.text = "enue 2025 up"
    r2.font.italic = True
    r2.font.size = PPt(20)
    r2.font.color.rgb = PRGB(0xC0, 0, 0)
    pic = s2.shapes.add_picture(io.BytesIO(LOGO), Inches(6), Inches(1), Inches(2), Inches(2))
    s2.notes_slide.notes_text_frame.text = "Say the number slowly: 2025"  # type: ignore[union-attr]
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "Detail"  # type: ignore[union-attr]
    s3.shapes.add_picture(io.BytesIO(LOGO), Inches(1), Inches(1), Inches(1), Inches(1))  # same media as slide 2
    s3.shapes.add_picture(io.BytesIO(OWN), Inches(3), Inches(1), Inches(1), Inches(1))  # media only slide 3 uses
    tbl = s3.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
    tbl.cell(0, 0).text = "Region"
    tbl.cell(0, 1).text = "Total 2025"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue(), pic.shape_id


# ---- helpers on packages ------------------------------------------------------------------------


def parts(data: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return {i.filename: z.read(i.filename) for i in z.infolist()}


def rezip(data: bytes, change: dict[str, bytes | None], add: dict[str, bytes] | None = None) -> bytes:
    """A copy of the package with parts replaced (bytes), removed (None) or added."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as zi, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for i in zi.infolist():
            if i.filename in change:
                if change[i.filename] is None:
                    continue
                zo.writestr(i.filename, change[i.filename] or b"")
            else:
                zo.writestr(i.filename, zi.read(i.filename))
        for n, b in (add or {}).items():
            zo.writestr(n, b)
    return out.getvalue()
