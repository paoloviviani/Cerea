"""The same engines through the HTTP API: authentication, shapes, errors, multipart in and out."""

from __future__ import annotations

import base64
import io
import json

import docx
import openpyxl
import pptx
import pytest
from starlette.testclient import TestClient

from cerea_documents.api import create_app
from cerea_documents.config import Settings
from cerea_documents.errors import ConfigError

from .conftest import AUTH, SECRET, parse_multipart, post
from .fixtures import TARGET, make_docx, make_pptx, make_xlsx, png


def test_healthz_is_open_and_capabilities_need_the_secret(client: TestClient) -> None:
    assert client.get("/healthz").json()["status"] == "ok"
    assert client.get("/v1/capabilities").status_code == 401
    assert client.get("/v1/capabilities", headers={"X-Cerea-Documents-Secret": "wrong"}).status_code == 401
    r = client.get("/v1/capabilities", headers=AUTH)
    assert r.status_code == 200
    caps = r.json()
    assert set(caps["formats"]) == {"docx", "xlsx", "pptx"}
    for fmt in caps["formats"].values():
        # every operation is described by its schema (the chat builds its tool text from these)
        assert sorted(fmt["operation_schemas"]) == fmt["operations"]
        assert all(sch["properties"]["op"] and sch["required"][0] == "op" for sch in fmt["operation_schemas"].values())
    assert caps["formats"]["docx"]["operations"] == sorted(
        [
            "replace_text",
            "fill_placeholders",
            "set_paragraph_text",
            "insert_paragraph_after",
            "delete_paragraph",
            "set_table_cell",
            "add_table_row",
            "insert_image_after",
            "replace_image",
        ]
    )
    assert caps["formats"]["xlsx"]["operations"] == ["add_sheet", "clear_cells", "insert_image", "set_cells"]
    assert caps["formats"]["pptx"]["operations"] == sorted(
        [
            "replace_text",
            "set_placeholder_text",
            "replace_image",
            "insert_image",
            "duplicate_slide",
            "delete_slide",
            "reorder_slides",
            "set_notes",
        ]
    )
    assert "add_comment" in caps["formats"]["docx"]["deferred_operations"] and caps["create"] is True
    assert caps["limits"]["max_zip_entries"] == 5000


def test_every_v1_endpoint_requires_the_secret(client: TestClient) -> None:
    for path in ("/v1/outline", "/v1/read", "/v1/edit"):
        r = client.post(path, files={"file": ("a.docx", make_docx())})
        assert r.status_code == 401 and r.json()["error"]["code"] == "unauthorized"


def test_service_refuses_to_start_without_a_secret_unless_dev_flag() -> None:
    with pytest.raises(ConfigError, match="CEREA_DOCUMENTS_SECRET"):
        create_app(Settings())
    app = create_app(Settings(allow_no_auth=True))
    with TestClient(app) as c:
        assert c.get("/v1/capabilities").status_code == 200  # dev flag: open


def test_secret_comparison_is_constant_time_api() -> None:
    import inspect

    from cerea_documents import api

    assert "hmac.compare_digest" in inspect.getsource(api)


def test_outline_docx(client: TestClient) -> None:
    r = post(client, "/v1/outline", "report.docx", make_docx())
    assert r.status_code == 200
    j = r.json()
    assert j["format"] == "docx" and j["blocks"][0]["id"] == "p1" and "client" in j["placeholders"]


def test_outline_and_read_xlsx_and_pptx(client: TestClient) -> None:
    r = post(client, "/v1/outline", "book.xlsx", make_xlsx())
    assert [s["name"] for s in r.json()["sheets"]][2] == TARGET
    r = post(client, "/v1/read", "book.xlsx", make_xlsx(), {"targets": [{"sheet": TARGET, "range": "A1:B2"}]})
    assert r.json()["targets"][0]["cells"]["B2"]["value"] == 6
    deck, _ = make_pptx()
    r = post(client, "/v1/outline", "deck.pptx", deck)
    assert len(r.json()["slides"]) == 3
    r = post(client, "/v1/read", "deck.pptx", deck, {"targets": [{"slide": 1}]})
    assert r.json()["targets"][0]["shapes"][0]["paragraphs"][0]["text"] == "Annual review 2025"


def test_edit_returns_multipart_with_file_diff_and_verify(client: TestClient) -> None:
    body = {"operations": [{"op": "replace_text", "find": "Acme Ltd", "replace": "Beta plc"}], "author": "Cerea"}
    r = post(client, "/v1/edit", "report.docx", make_docx(), body)
    assert r.status_code == 200 and r.headers["content-type"].startswith("multipart/form-data")
    result, file = parse_multipart(r)
    assert result["format"] == "docx" and result["diff"][-1] == "+ Dear Beta plc, welcome"
    assert result["verify"]["ok"] and result["verify"]["changed"] == ["word/document.xml"]
    assert result["file_name"] == "report.docx" and result["engine"].startswith("cerea-documents/")
    assert docx.Document(io.BytesIO(file)).paragraphs[1].text == "Dear Beta plc, welcome"


def test_edit_json_response_with_base64_file(client: TestClient) -> None:
    body = {"operations": [{"op": "set_cells", "sheet": TARGET, "cells": {"B2": 1}}]}
    r = post(client, "/v1/edit", "book.xlsx", make_xlsx(), body, params={"response": "json"})
    j = r.json()
    wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(j["file_base64"])))
    assert wb[TARGET]["B2"].value == 1 and j["file_bytes"] > 0


def test_edit_with_images_part(client: TestClient) -> None:
    deck, _ = make_pptx()
    body = {"operations": [{"op": "insert_image", "slide": 1, "image": "logo.png"}]}
    r = post(client, "/v1/edit", "deck.pptx", deck, body, images={"logo.png": png(30, 30, (1, 1, 1))})
    result, file = parse_multipart(r)
    assert result["verify"]["added"] == ["ppt/media/image3.png"] or len(result["verify"]["added"]) == 1
    assert len([s for s in pptx.Presentation(io.BytesIO(file)).slides[0].shapes if s.shape_type == 13]) == 1


def test_errors_are_json_with_code_and_operation_index(client: TestClient) -> None:
    body = {
        "operations": [
            {"op": "replace_text", "find": "Acme", "replace": "y"},
            {"op": "set_table_cell", "table": "t9", "row": 0, "col": 0, "text": "z"},
        ]
    }
    r = post(client, "/v1/edit", "report.docx", make_docx(), body)
    assert r.status_code == 422
    e = r.json()["error"]
    assert e["code"] in ("operation_failed",) and e["operation_index"] == 1


def test_actionable_error_lists_what_exists(client: TestClient) -> None:
    body = {"operations": [{"op": "set_table_cell", "table": "t9", "row": 0, "col": 0, "text": "z"}]}
    r = post(client, "/v1/edit", "report.docx", make_docx(), body)
    e = r.json()["error"]
    assert (
        e["code"] == "operation_failed" and "tables t1..t1" in e["message"] and e["details"]["nothing_written"] is True
    )


def test_schema_validation_before_any_work(client: TestClient) -> None:
    r = post(client, "/v1/edit", "report.docx", make_docx(), {"operations": [{"op": "replace_text", "replace": "y"}]})
    e = r.json()["error"]
    assert (
        r.status_code == 422
        and e["code"] == "invalid_request"
        and e["reason"] == "schema"
        and e["operation_index"] == 0
    )
    assert "find" in e["message"] and "find" in e["details"]["required"]
    r = post(client, "/v1/edit", "report.docx", make_docx(), {"operations": []})
    assert r.status_code == 422 and "operations" in r.json()["error"]["message"]
    r = post(client, "/v1/edit", "report.docx", make_docx(), {"operations": [{"op": "nope"}]})
    assert "available" in r.json()["error"]["details"]
    r = post(
        client,
        "/v1/edit",
        "report.docx",
        make_docx(),
        {"operations": [{"op": "replace_text", "find": "a", "replace": "b"}], "extra": 1},
    )
    assert r.status_code == 422
    r = post(client, "/v1/edit", "report.docx", make_docx(), None)
    assert r.status_code == 422 and r.json()["error"]["reason"] == "body"


def test_bad_json_and_nan_in_body(client: TestClient) -> None:
    r = client.post("/v1/edit", files={"file": ("a.docx", make_docx())}, data={"body": "{not json"}, headers=AUTH)
    assert r.status_code == 422 and r.json()["error"]["reason"] == "body"
    r = client.post(
        "/v1/edit",
        files={"file": ("a.xlsx", make_xlsx())},
        data={"body": '{"operations":[{"op":"set_cells","sheet":"T2.4","cells":{"A1":NaN}}]}'},
        headers=AUTH,
    )
    assert r.status_code == 422


def test_format_must_be_knowable_and_match_content(client: TestClient) -> None:
    body = {"operations": [{"op": "replace_text", "find": "a", "replace": "b"}]}
    r = post(client, "/v1/edit", "noextension", make_docx(), body)
    assert r.status_code == 422 and r.json()["error"]["reason"] == "format"
    r = post(
        client,
        "/v1/edit",
        "noextension",
        make_docx(),
        {"operations": [{"op": "replace_text", "find": "zzz", "replace": "b"}], "format": "docx"},
    )
    assert r.status_code == 422 and r.json()["error"]["code"] == "operation_failed"  # no match for 'zzz'
    r = post(client, "/v1/edit", "book.xlsx", make_docx(), body)
    assert r.status_code == 415 and r.json()["error"]["reason"] == "format_mismatch"


def test_not_multipart(client: TestClient) -> None:
    r = client.post("/v1/outline", content=b"x", headers={**AUTH, "content-type": "application/json"})
    assert r.status_code == 422 and r.json()["error"]["reason"] == "content_type"


def test_secret_never_reaches_the_worker() -> None:
    from cerea_documents.runner import _env

    assert SECRET not in json.dumps(_env()) and "CEREA_DOCUMENTS_SECRET" not in _env()


def test_chat_capabilities_fixture_matches_the_service(client: TestClient) -> None:
    """The chat's contract test reads this fixture; it must be what the service says today.
    Regenerate: uv run python -c "import json; from cerea_documents.core import capabilities; ..."
    (see the header of src/lib/server/documents/documentsContract.spec.ts)."""
    import json
    from pathlib import Path

    path = Path(__file__).resolve().parents[2] / "src/lib/server/documents/capabilities.fixture.json"
    live = client.get("/v1/capabilities", headers=AUTH).json()
    live.pop("limits")
    saved = json.loads(path.read_text())
    saved.pop("limits")
    assert saved == live, "the chat's capabilities fixture is stale: regenerate it"
