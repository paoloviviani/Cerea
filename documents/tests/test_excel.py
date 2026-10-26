"""Excel engine, ported from the 0.29.0 xlsx_patch checks, run against the service functions."""

from __future__ import annotations

import io
import warnings
import zipfile
from typing import Any

import openpyxl
import pytest

from cerea_documents import core
from cerea_documents.errors import DocumentError

from .fixtures import SHEETS, TARGET, make_xlsx, parts, png
from .helpers import LIMITS, do_edit, refuses


@pytest.fixture(scope="module")
def src() -> bytes:
    return make_xlsx()


def edit(src: bytes, ops: list[dict[str, Any]]) -> tuple[dict[str, Any], bytes]:
    return do_edit(src, ops, fmt="xlsx")


def wb(data: bytes) -> Any:
    return openpyxl.load_workbook(io.BytesIO(data))


def test_fixture_carries_what_openpyxl_would_drop(src: bytes) -> None:
    p = parts(src)
    assert b"x14:" in p["xl/worksheets/sheet2.xml"] and "xl/calcChain.xml" in p


def test_outline(src: bytes) -> None:
    o = core.outline(src, "xlsx", LIMITS)
    assert [s["name"] for s in o["sheets"]] == SHEETS
    t = next(s for s in o["sheets"] if s["name"] == TARGET)
    assert t["used_range"] == "A1:E29" and t["merged_ranges"] == 1 and t["state"] == "visible"


def test_read_range_values_and_formulas(src: bytes) -> None:
    r = core.read(src, "xlsx", {"targets": [{"sheet": TARGET, "range": "A1:C3"}]}, LIMITS)
    cells = r["targets"][0]["cells"]
    assert cells["A1"]["value"] == "Header T2.4" and cells["B2"]["value"] == 6
    assert cells["C2"]["formula"] == "=B2*2"
    assert cells["C3"]["formula"] == "(shared formula)"
    with pytest.raises(DocumentError, match="no sheet 'nope'; the sheets are"):
        core.read(src, "xlsx", {"targets": [{"sheet": "nope"}]}, LIMITS)


def test_values_dates_formulas_new_row_style_copy_clear(src: bytes) -> None:
    ops = [
        {
            "op": "set_cells",
            "sheet": TARGET,
            "cells": {
                "B2": "Hello & <world>",
                "B3": 42,
                "B4": 3.14,
                "E2": {"date": "2026-01-31"},
                "E3": {"datetime": "2026-02-01T13:30:00"},
                "F20": "  spaced\nline",
                "G5": True,
                "H5": {"formula": "=B5*10"},
                "A40": "new row",
                "A41": {"value": "styled", "style_from": "A2"},
                "D3": None,
            },
        }
    ]
    rep, out = edit(src, ops)
    r = rep["operations"][0]["result"]
    assert r["formulas_added"] == 1 and r["formula_cells_replaced"] == []
    assert set(rep["verify"]["changed"]) == {"xl/worksheets/sheet3.xml"}  # openpyxl already wrote fullCalcOnLoad
    assert rep["verify"]["byte_identical"] == rep["verify"]["parts"] - 1
    assert rep["recalculate_on_open"] is True
    a, b = parts(src), parts(out)
    for n in a:
        if n not in rep["verify"]["changed"]:
            assert a[n] == b[n], n
    for n in ("xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml", "xl/worksheets/sheet4.xml"):
        assert b"x14:dataBar" in b[n], n
    assert b'fullCalcOnLoad="1"' in b["xl/workbook.xml"]
    ws = wb(out)[TARGET]
    assert ws["B2"].value == "Hello & <world>" and ws["B3"].value == 42 and ws["B4"].value == 3.14
    assert ws["E2"].value.isoformat() == "2026-01-31T00:00:00" and ws["E2"].number_format == "dd/mm/yyyy"
    assert ws["E3"].value.isoformat() == "2026-02-01T13:30:00"
    assert ws["F20"].value == "  spaced\nline" and ws["G5"].value is True and ws["H5"].value == "=B5*10"
    assert ws["A40"].value == "new row" and ws["A41"].value == "styled" and ws["A41"].style_id == ws["A2"].style_id
    assert ws["D3"].value is None and ws["C3"].value == "=B3*2"
    assert ws.max_row == 41 and "A1:C1" in {str(m) for m in ws.merged_cells.ranges}
    assert any(line == "T2.4!B2: 6 -> 'Hello & <world>'" for line in rep["diff"]), rep["diff"][:5]


def test_replacing_a_formula_cell_updates_calcchain(src: bytes) -> None:
    rep, out = edit(src, [{"op": "set_cells", "sheet": TARGET, "cells": {"C4": 99}}])
    assert set(rep["verify"]["changed"]) == {"xl/calcChain.xml", "xl/worksheets/sheet3.xml"}
    before = parts(src)["xl/calcChain.xml"].decode().count("<c ")
    chain = parts(out)["xl/calcChain.xml"].decode()
    assert chain.count("<c ") == before - 1 and 'r="C4"' not in chain.split('i="3"')[1].split('i="4"')[0]
    assert '<c r="C2" i="3"/>' in chain


def test_1904_date_system() -> None:
    f = make_xlsx(date1904=True)
    _, out = edit(f, [{"op": "set_cells", "sheet": TARGET, "cells": {"E2": {"date": "2026-01-31"}}}])
    assert wb(out)[TARGET]["E2"].value.isoformat() == "2026-01-31T00:00:00"


def test_refusals(src: bytes) -> None:
    def bad(cells: dict[str, Any], needle: str, sheet: str = TARGET) -> None:
        refuses(src, [{"op": "set_cells", "sheet": sheet, "cells": cells}], needle, fmt="xlsx")

    bad({"C2": 5}, "master of a shared formula")
    bad({"C2": {"formula": "=1"}}, "master of a shared formula")
    bad({"B1": 1}, "merged range A1:C1")
    bad({"B1": 1, "A1": 1}, "merged range A1:C1")
    bad({"A1": 1}, "no sheet", sheet="nope")
    bad({"ZZZZ1": 1}, "does not match")
    bad({"A1": {"value": 1, "formula": "=1"}}, "exactly one of")
    bad({"A1": {"date": "31/01/2026"}}, "ISO 8601")
    with pytest.raises(DocumentError, match="does not match"):
        edit(src, [{"op": "set_cells", "sheet": TARGET, "cells": {"not a ref": 1}}])
    _, out = edit(src, [{"op": "set_cells", "sheet": TARGET, "cells": {"A1": "top-left is fine"}}])
    assert wb(out)[TARGET]["A1"].value == "top-left is fine"


def test_clear_cells_keeps_style_and_updates_chain(src: bytes) -> None:
    rep, out = edit(src, [{"op": "clear_cells", "sheet": TARGET, "ranges": ["B3:C4", "E2"]}])
    r = rep["operations"][0]["result"]
    assert r["cells_cleared"] == 4 and set(r["formula_cells_cleared"]) == {"C3", "C4"}  # E2 holds no value: not counted
    ws = wb(out)[TARGET]
    assert ws["B2"].value == 6 and ws["B3"].value is None and ws["E2"].value is None
    assert ws["E2"].number_format == "dd/mm/yyyy"


def test_clear_refuses_a_shared_formula_master_cleanly(src: bytes) -> None:
    refuses(src, [{"op": "clear_cells", "sheet": TARGET, "ranges": ["C2"]}], "master of a shared formula", fmt="xlsx")


def test_add_sheet_appends_and_leaves_the_rest_byte_identical(src: bytes) -> None:
    rep, out = edit(src, [{"op": "add_sheet", "name": "Notes", "cells": {"A1": "hello", "B2": 5}}])
    assert rep["operations"][0]["result"]["position"] == len(SHEETS) + 1
    w = wb(out)
    assert w.sheetnames == [*SHEETS, "Notes"] and w["Notes"]["A1"].value == "hello" and w["Notes"]["B2"].value == 5
    v = rep["verify"]
    assert "xl/worksheets/sheet6.xml" in v["added"]
    for n, data in parts(src).items():
        if n not in ("xl/workbook.xml", "xl/_rels/workbook.xml.rels", "[Content_Types].xml"):
            assert parts(out)[n] == data, n
    assert "+ sheet 'Notes' added" in rep["diff"]
    refuses(src, [{"op": "add_sheet", "name": "cover"}], "already exists", fmt="xlsx")
    refuses(src, [{"op": "add_sheet", "name": "a/b"}], "not a valid sheet name", fmt="xlsx")
    refuses(src, [{"op": "add_sheet", "name": "x" * 32}], "not a valid sheet name", fmt="xlsx")
    _, out2 = edit(src, [{"op": "add_sheet", "name": "N1"}, {"op": "set_cells", "sheet": "N1", "cells": {"A1": 1}}])
    assert wb(out2)["N1"]["A1"].value == 1


def test_insert_image_creates_a_drawing_and_reuses_it() -> None:
    src = make_xlsx()
    img = png(60, 30, (200, 0, 0))
    rep, out = do_edit(
        src,
        [
            {"op": "insert_image", "sheet": TARGET, "cell": "E5", "image": "a.png", "width_px": 120, "alt": "chart"},
            {"op": "insert_image", "sheet": TARGET, "cell": "E12", "image": "b.png"},
        ],
        fmt="xlsx",
        images={"a.png": img, "b.png": png(10, 10, (0, 0, 9))},
    )
    p = parts(out)
    assert "xl/drawings/drawing1.xml" in p and "xl/media/image1.png" in p and "xl/media/image2.png" in p
    assert p["xl/drawings/drawing1.xml"].count(b"<xdr:oneCellAnchor") == 2
    ws = wb(out)[TARGET]
    assert len(ws._images) == 2 and ws._images[0].anchor._from.col == 4 and ws._images[0].anchor._from.row == 4
    assert b"x14:dataBar" in p["xl/worksheets/sheet3.xml"]
    assert set(rep["verify"]["changed"]) >= {"xl/worksheets/sheet3.xml", "[Content_Types].xml"}
    o = core.outline(out, "xlsx", LIMITS)
    assert next(s for s in o["sheets"] if s["name"] == TARGET)["images"] == 2
    # the sheet now has a drawing, so a second request reuses it
    _, out2 = do_edit(
        out,
        [{"op": "insert_image", "sheet": TARGET, "cell": "A50", "image": "c.png"}],
        fmt="xlsx",
        images={"c.png": img},
    )
    assert parts(out2)["xl/drawings/drawing1.xml"].count(
        b"<xdr:oneCellAnchor"
    ) == 3 and "xl/drawings/drawing2.xml" not in parts(out2)


def test_track_changes_is_refused_for_excel(src: bytes) -> None:
    with pytest.raises(DocumentError, match="only available for Word"):
        do_edit(src, [{"op": "set_cells", "sheet": TARGET, "cells": {"A1": 1}}], fmt="xlsx", track=True)


def test_openpyxl_round_trip_would_lose_the_extensions(src: bytes) -> None:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        buf = io.BytesIO()
        wb(src).save(buf)
    assert b"x14:" not in parts(buf.getvalue())["xl/worksheets/sheet2.xml"]


def test_verify_catches_tampering(src: bytes) -> None:
    from cerea_documents import excel
    from cerea_documents.config import Limits
    from cerea_documents.errors import PatchError
    from cerea_documents.pkg import Pkg
    from cerea_documents.verify import verify

    from . import fixtures

    _, good = edit(src, [{"op": "set_cells", "sheet": TARGET, "cells": {"B2": 1}}])
    bad = fixtures.rezip(
        good, {"xl/worksheets/sheet2.xml": parts(good)["xl/worksheets/sheet2.xml"].replace(b"x14:", b"x15:")}
    )
    with pytest.raises(PatchError, match="sheet2.xml"):
        verify(Pkg(src, Limits()), bad, Limits(), excel, {"xl/worksheets/sheet3.xml", "xl/workbook.xml"}, "xlsx")
    with zipfile.ZipFile(io.BytesIO(good)) as z:
        assert z.testzip() is None


def test_calcpr_is_added_when_the_workbook_has_none(src: bytes) -> None:
    import re

    from . import fixtures

    wbx = parts(src)["xl/workbook.xml"].decode()
    stripped = fixtures.rezip(src, {"xl/workbook.xml": re.sub(r"<calcPr[^>]*/>", "", wbx).encode()})
    assert b"calcPr" not in parts(stripped)["xl/workbook.xml"]
    rep, out = edit(stripped, [{"op": "set_cells", "sheet": TARGET, "cells": {"H5": {"formula": "=1+1"}}}])
    assert (
        b'<calcPr fullCalcOnLoad="1"/>' in parts(out)["xl/workbook.xml"]
        and "xl/workbook.xml" in rep["verify"]["changed"]
    )
    assert wb(out)[TARGET]["H5"].value == "=1+1"
