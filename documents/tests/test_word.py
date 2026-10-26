"""Word engine, ported from the 0.29.0 docx helper checks, run against the service functions."""

from __future__ import annotations

import io
import re
import zipfile
from typing import Any

import docx
import pytest
from docx.shared import Pt, RGBColor

from cerea_documents import core
from cerea_documents.errors import DocumentError

from .fixtures import DOCX_CONTENT, jpeg, make_docx, make_docx_with_image, parts, png
from .helpers import LIMITS, do_edit, refuses, untouched


@pytest.fixture(scope="module")
def src() -> bytes:
    return make_docx()


def edit(src: bytes, ops: list[dict[str, Any]], **kw: Any) -> tuple[dict[str, Any], bytes]:
    return do_edit(src, ops, fmt="docx", **kw)


def opened(data: bytes) -> Any:
    return docx.Document(io.BytesIO(data))


def runs_of(p: Any) -> list[tuple[str, bool, bool]]:
    return [(r.text, bool(r.bold), bool(r.italic)) for r in p.runs]


def test_outline_ids_styles_placeholders(src: bytes) -> None:
    o = core.outline(src, "docx", LIMITS)
    ids = {b["id"]: b for b in o["blocks"]}
    assert ids["p1"]["type"] == "heading" and ids["p1"]["level"] == 1
    assert ids["p2"]["text"] == "Dear Acme Ltd, welcome"
    assert ids["t1"]["type"] == "table" and ids["t1"]["rows"] == 3 and ids["t1"]["cols"] == 2
    assert ids["t1"]["cells"][1][0] == "Widget {{client}}"
    assert o["placeholders"] == ["client", "date", "other", "footer_note"]
    assert o["comments"] == 1 and any("Report for" in "".join(v) for v in o["headers"].values())
    assert o["sections"] and o["format"] == "docx"


def test_read_runs_and_table(src: bytes) -> None:
    r = core.read(src, "docx", {"targets": ["p2", "t1"]}, LIMITS)
    p2 = r["targets"][0]
    assert p2["text"] == "Dear Acme Ltd, welcome"
    assert [x["text"] for x in p2["runs"]] == ["Dear Ac", "me L", "td, welcome"]
    assert p2["runs"][0]["bold"] is True and p2["runs"][1]["italic"] is True and p2["runs"][1]["color"] == "C00000"
    assert r["targets"][1]["rows"][1][0]["text"] == "Widget {{client}}"
    with pytest.raises(DocumentError, match=r"paragraphs p1\.\.p"):
        core.read(src, "docx", {"targets": ["p99"]}, LIMITS)


def test_replace_across_runs_keeps_formatting(src: bytes) -> None:
    rep, out = edit(src, [{"op": "replace_text", "find": "Acme Ltd", "replace": "Beta plc"}])
    assert rep["operations"][0]["result"]["replaced"] == 1
    assert rep["verify"]["changed"] == ["word/document.xml"]
    p2 = opened(out).paragraphs[1]
    assert p2.text == "Dear Beta plc, welcome"
    got = runs_of(p2)
    assert "".join(t for t, b, i in got if b and not i) == "Dear Beta plc"
    assert [t for t, b, i in got if not b and not i] == [", welcome"]
    assert not any(i for _t, _b, i in got), "the italic colour run must be fully consumed"
    assert rep["diff"] == ["word/document.xml:", "- Dear Acme Ltd, welcome", "+ Dear Beta plc, welcome"]
    total = untouched(src, out, {"word/document.xml"})
    assert parts(out)["word/comments.xml"] == parts(src)["word/comments.xml"]
    assert all(
        k in parts(out)["word/document.xml"].decode()
        for k in ("commentRangeStart", "commentRangeEnd", "commentReference")
    )
    assert rep["verify"]["byte_identical"] == total - 1


def test_fill_placeholders_everywhere_and_reports_leftovers(src: bytes) -> None:
    rep, out = edit(
        src, [{"op": "fill_placeholders", "values": {"client": "Beta plc", "date": "26 Oct 2026", "unused": "x"}}]
    )
    r = rep["operations"][0]["result"]
    assert r["filled"] == {"client": 4, "date": 1}
    assert r["missing"] == ["unused"] and r["unfilled"] == ["footer_note", "other"]
    d2 = opened(out)
    assert d2.paragraphs[2].text == "Prepared for Beta plc on 26 Oct 2026."
    assert d2.tables[0].rows[1].cells[0].text == "Widget Beta plc"
    assert d2.sections[0].header.paragraphs[0].text == "Report for Beta plc"
    assert d2.sections[0].footer.paragraphs[0].text == "Confidential - {{footer_note}}"
    assert [(r.text, r.bold) for r in d2.sections[0].header.paragraphs[0].runs] == [
        ("Report for ", True),
        ("Beta plc", True),
    ]
    untouched(src, out, {"word/document.xml", "word/header1.xml"})
    refuses(src, [{"op": "fill_placeholders", "values": {"client": "x"}, "strict": True}], "strict mode", fmt="docx")
    refuses(
        src,
        [{"op": "fill_placeholders", "values": {"nothing": "x"}}],
        "found ['client', 'date', 'footer_note', 'other']",
        fmt="docx",
    )


def test_table_cell_and_row(src: bytes) -> None:
    _, o3 = edit(src, [{"op": "set_table_cell", "table": "t1", "row": 1, "col": 1, "text": "12.5"}])
    cell = opened(o3).tables[0].rows[1].cells[1]
    assert cell.text == "12.5" and cell.paragraphs[0].runs[0].font.size == Pt(9)
    _, o3 = edit(o3, [{"op": "add_table_row", "table": "t1", "values": ["Gadget", "7"]}])
    t = opened(o3).tables[0]
    assert len(t.rows) == 4 and [c.text for c in t.rows[3].cells] == ["Gadget", "7"]
    run = t.rows[3].cells[0].paragraphs[0].runs[0]
    assert run.font.size == Pt(9) and run.font.color.rgb == RGBColor(0, 0x66, 0), "run formatting not copied"
    assert t.rows[3].cells[0]._tc.xml.count("F2F2F2") == 1, "cell shading not copied"
    refuses(src, [{"op": "set_table_cell", "table": "t1", "row": 9, "col": 0, "text": "x"}], "has 3 rows", fmt="docx")
    refuses(src, [{"op": "add_table_row", "table": "t1", "values": ["a", "b", "c"]}], "2 cells", fmt="docx")
    refuses(
        src,
        [
            {"op": "set_table_cell", "table": "t1", "row": 0, "col": 0, "text": "x"},
            {"op": "set_paragraph_text", "paragraph": "t1", "text": "x"},
        ],
        "does not match",
        fmt="docx",
    )


def test_insert_paragraph_after_and_set_text(src: bytes) -> None:
    _, out = edit(src, [{"op": "insert_paragraph_after", "anchor": "p1", "text": "Background", "style": "Heading 2"}])
    ps = opened(out).paragraphs
    assert ps[1].text == "Background" and ps[1].style.name == "Heading 2" and ps[2].text.startswith("Dear")
    _, out = edit(src, [{"op": "insert_paragraph_after", "anchor": "p4", "text": "Second point"}])
    ps = opened(out).paragraphs
    assert ps[4].text == "Second point" and ps[4].style.name == "List Bullet"
    _, out = edit(src, [{"op": "insert_paragraph_after", "anchor": "t1", "text": "After the table, new"}])
    assert [p.text for p in opened(out).paragraphs][-2:] == ["After the table, new", "After the table."]
    refuses(
        src,
        [{"op": "insert_paragraph_after", "anchor": "p2", "text": "x", "style": "No Such Style"}],
        "available:",
        fmt="docx",
    )
    refuses(src, [{"op": "insert_paragraph_after", "anchor": "p99", "text": "x"}], "no block 'p99'", fmt="docx")
    _, out = edit(src, [{"op": "set_paragraph_text", "paragraph": "p2", "text": "Hello team"}])
    first = opened(out).paragraphs[1]
    assert first.text == "Hello team" and first.runs[0].bold, "formatting of the first run is kept"


def test_ids_refer_to_the_original_document_within_a_request(src: bytes) -> None:
    ops: list[dict[str, Any]] = [
        {"op": "insert_paragraph_after", "anchor": "p1", "text": "New first"},
        {"op": "set_paragraph_text", "paragraph": "p2", "text": "Still the original second"},
        {"op": "delete_paragraph", "paragraph": "p5"},
    ]
    _, out = edit(src, ops)
    texts = [p.text for p in opened(out).paragraphs]
    assert texts[:3] == ["Statement of work", "New first", "Still the original second"]
    assert "Review the pricing" not in texts
    refuses(
        src,
        [{"op": "delete_paragraph", "paragraph": "p5"}, {"op": "set_paragraph_text", "paragraph": "p5", "text": "x"}],
        "deleted by an earlier operation",
        fmt="docx",
    )


def test_delete_paragraph(src: bytes) -> None:
    rep, out = edit(src, [{"op": "delete_paragraph", "paragraph": "p4"}])
    assert "Open points" not in [p.text for p in opened(out).paragraphs]
    assert "- Open points" in rep["diff"]
    _, out = edit(src, [{"op": "delete_paragraph", "paragraph": "p4"}], track=True, author="Cerea Bot")
    xml = parts(out)["word/document.xml"].decode()
    assert re.search(r"<w:del [^>]*w:author=\"Cerea Bot\"", xml) and "Open points" in xml
    last = f"p{core.outline(src, 'docx', LIMITS)['blocks'][-1]['id'][1:]}"
    refuses(src, [{"op": "delete_paragraph", "paragraph": last}], "closes the document after a table", fmt="docx")


def test_tracked_changes(src: bytes) -> None:
    rep, out = edit(
        src, [{"op": "replace_text", "find": "Acme Ltd", "replace": "Beta plc"}], track=True, author="Cerea Bot"
    )
    xml = parts(out)["word/document.xml"].decode()
    dels = re.findall(r"<w:del [^>]*w:author=\"Cerea Bot\"[^>]*>.*?</w:del>", xml)
    assert "".join(re.findall(r"<w:delText[^>]*>([^<]*)</w:delText>", "".join(dels))) == "Acme Ltd"
    assert re.search(r'<w:ins [^>]*w:date="\d{4}-\d\d-\d\dT[^"]*Z"[^>]*>.*?<w:t[^>]*>Beta plc</w:t>', xml)
    ids = re.findall(r'<w:(?:ins|del) w:id="(\d+)"', xml)
    assert len(ids) == len(set(ids)) and len(ids) >= 4
    assert rep["verify"]["tracked"] == {"ins": 1, "del": 3}
    assert "+ Dear Beta plc, welcome" in rep["diff"]
    _, o3 = edit(src, [{"op": "add_table_row", "table": "t1", "values": ["Tracked", "1"]}], track=True)
    assert "<w:trPr>" in parts(o3)["word/document.xml"].decode()
    _, o3 = edit(src, [{"op": "insert_paragraph_after", "anchor": "p2", "text": "Tracked paragraph"}], track=True)
    assert re.search(r"<w:ins [^>]*>.*?Tracked paragraph", parts(o3)["word/document.xml"].decode())


def test_chained_ops_one_request_and_verify_scope(src: bytes) -> None:
    ops: list[dict[str, Any]] = [
        {"op": "fill_placeholders", "values": {"client": "Beta plc", "footer_note": "Draft"}},
        {"op": "replace_text", "find": "Acme Ltd", "replace": "Beta plc", "scope": "body"},
    ]
    rep, _out = edit(src, ops)
    v = rep["verify"]
    assert set(v["changed"]) == DOCX_CONTENT and v["byte_identical"] == v["parts"] - 3
    assert "+ Confidential - Draft" in rep["diff"]
    refuses(src, [{"op": "replace_text", "find": "no such text", "replace": "x"}], "no match for", fmt="docx")
    refuses(src, [{"op": "replace_text", "find": "Acme", "replace": "x", "scope": "p99"}], "no block", fmt="docx")


def test_all_or_none_names_the_failing_operation(src: bytes) -> None:
    ops: list[dict[str, Any]] = [
        {"op": "replace_text", "find": "Acme Ltd", "replace": "Beta plc"},
        {"op": "set_table_cell", "table": "t1", "row": 40, "col": 0, "text": "x"},
    ]
    with pytest.raises(DocumentError) as ei:
        edit(src, ops)
    e = ei.value
    assert e.operation_index == 1 and e.details["nothing_written"] and e.details["applied_before_failure"] == 1
    assert "operation 1 (set_table_cell)" in str(e)


def test_schema_errors_name_the_operation_and_list_what_exists(src: bytes) -> None:
    with pytest.raises(DocumentError, match=r"unknown operation 'nope'.*available:"):
        edit(src, [{"op": "nope"}])
    with pytest.raises(DocumentError, match=r"operations\[0\] \(replace_text\): .*find"):
        edit(src, [{"op": "replace_text", "replace": "x"}])
    with pytest.raises(DocumentError, match="is planned but not implemented"):
        edit(src, [{"op": "add_comment"}])
    with pytest.raises(DocumentError, match="Additional properties"):
        edit(src, [{"op": "replace_text", "find": "a", "replace": "b", "regex": True}])


def test_insert_image_after_and_caption() -> None:
    src = make_docx()
    rep, out = edit(
        src,
        [
            {
                "op": "insert_image_after",
                "anchor": "p2",
                "image": "logo.png",
                "width_cm": 4,
                "alt": "Company logo",
                "caption": "Figure 1",
            }
        ],
        images={"logo.png": png(80, 40, (10, 20, 30))},
    )
    p = parts(out)
    assert any(n.startswith("word/media/image") and n.endswith(".png") for n in p)
    d = opened(out)
    assert (
        len(d.inline_shapes) == 1
        and abs(d.inline_shapes[0].width.cm - 4) < 0.01
        and abs(d.inline_shapes[0].height.cm - 2) < 0.01
    )
    assert d.paragraphs[3].text == "Figure 1"
    v = rep["verify"]
    assert "word/media/image1.png" in v["added"] and "word/_rels/document.xml.rels" in v["changed"]
    assert "images: 0 -> 1" in rep["diff"]
    untouched(src, out, {"word/document.xml", "word/_rels/document.xml.rels", "[Content_Types].xml"} | set(v["added"]))
    o = core.outline(out, "docx", LIMITS)
    assert o["images"][0]["alt"] == "Company logo" and o["images"][0]["size_cm"] == [4.0, 2.0]
    _, out2 = edit(
        src,
        [{"op": "insert_image_after", "anchor": "p2", "image": "p.jpg"}],
        images={"p.jpg": jpeg(60, 30)},
        track=True,
    )
    x2 = parts(out2)["word/document.xml"].decode()
    assert re.search(r"<w:ins [^>]*><w:r><w:drawing", x2) and "word/media/image1.jpeg" in parts(out2)
    refuses(
        src, [{"op": "insert_image_after", "anchor": "p2", "image": "missing.png"}], "images sent: none", fmt="docx"
    )


def test_replace_image_crops_and_drops_old_media() -> None:
    src = make_docx_with_image()
    o = core.outline(src, "docx", LIMITS)
    assert o["images"][0]["id"] == "i1" and o["images"][0]["part"] == "word/media/image1.png"
    wide = png(80, 20, (0, 0, 200))
    rep, out = edit(src, [{"op": "replace_image", "image_id": "i1", "image": "new.png"}], images={"new.png": wide})
    r = rep["operations"][0]["result"]
    assert r["cropped"] and r["removed_media"] == ["word/media/image1.png"]
    p = parts(out)
    assert "word/media/image1.png" not in p and "word/media/image2.png" in p and p["word/media/image2.png"] == wide
    assert b"a:srcRect" in p["word/document.xml"]
    d = opened(out)
    assert abs(d.inline_shapes[0].width.cm - 2 * 2.54) < 0.01
    _, out2 = edit(
        src, [{"op": "replace_image", "image_id": "i1", "image": "n.png", "fit": "stretch"}], images={"n.png": wide}
    )
    assert b"a:srcRect" not in parts(out2)["word/document.xml"]
    refuses(
        src,
        [{"op": "replace_image", "image_id": "i7", "image": "n.png"}],
        "images i1..i1",
        fmt="docx",
        images={"n.png": wide},
    )
    refuses(
        make_docx(),
        [{"op": "replace_image", "image_id": "i1", "image": "n.png"}],
        "no images to replace",
        fmt="docx",
        images={"n.png": wide},
    )


def test_chained_edits_remain_openable_and_idempotent_on_untouched_parts(src: bytes) -> None:
    _, o1 = edit(src, [{"op": "replace_text", "find": "Acme Ltd", "replace": "Beta plc"}])
    rep, o2 = edit(o1, [{"op": "set_table_cell", "table": "t1", "row": 1, "col": 1, "text": "99"}])
    assert opened(o2).tables[0].rows[1].cells[1].text == "99"
    assert rep["verify"]["reopened_with"] == "python-docx"
    with zipfile.ZipFile(io.BytesIO(o2)) as z:
        assert z.testzip() is None


def test_wrong_kind_of_file_is_refused() -> None:
    with pytest.raises(DocumentError, match="named as docx"):
        core.outline(make_xlsx_bytes(), "docx", LIMITS)


def make_xlsx_bytes() -> bytes:
    from .fixtures import make_xlsx

    return make_xlsx()
