"""Creation from a spec (ADR 0103 step 3): every block, chart, image and theme override, per format."""

from __future__ import annotations

import io
import time
import zipfile
from typing import Any

import docx
import openpyxl
import pptx
import pypdf
import pytest
from docx.shared import RGBColor

from cerea_documents import create
from cerea_documents.config import Limits
from cerea_documents.errors import InvalidRequest

from .fixtures import jpeg, png

LIMITS = Limits()
LOGO = png(120, 60, (200, 40, 40))
CATS = ["North", "South", "East"]
SERIES = [{"name": "2025", "values": [3, 5, 2]}, {"name": "2026", "values": [4, 6, 3]}]
CHART_TYPES = ["bar", "column", "stacked", "line", "area", "pie", "doughnut", "scatter"]


def build(
    fmt: str, spec: dict[str, Any], images: dict[str, bytes] | None = None, template: bytes | None = None
) -> tuple[dict[str, Any], bytes]:
    return create.create({"format": fmt, "spec": spec}, template, images or {}, LIMITS)


def names(data: bytes) -> list[str]:
    return zipfile.ZipFile(io.BytesIO(data)).namelist()


def pdf_text(data: bytes) -> str:
    return "\n".join(p.extract_text() for p in pypdf.PdfReader(io.BytesIO(data)).pages)


def paragraphs(data: bytes) -> list[str]:
    return [p.text for p in docx.Document(io.BytesIO(data)).paragraphs]


def chart_block(kind: str, render: str) -> dict[str, Any]:
    block: dict[str, Any] = {
        "type": "chart",
        "chart_type": kind,
        "title": f"{kind} chart",
        "render": render,
        "x_title": "Region",
        "y_title": "Units",
    }
    if kind == "scatter":
        block["series"] = [
            {"name": "a", "points": [[1, 2], [2, 4], [3, 5]]},
            {"name": "b", "points": [[1, 1], [2, 3], [3, 2]]},
        ]
    elif kind in ("pie", "doughnut"):
        block |= {"categories": CATS, "series": SERIES[:1]}
    else:
        block |= {"categories": CATS, "series": SERIES}
    return block


# ---- every block type ---------------------------------------------------------------------------

ALL_BLOCKS: list[dict[str, Any]] = [
    {"type": "heading", "level": 1, "text": "Quarterly report"},
    {
        "type": "paragraph",
        "runs": [
            {"text": "Plain, "},
            {"text": "bold", "bold": True},
            {"text": " italic", "italic": True},
            {"text": " underlined", "underline": True},
            {"text": " code", "code": True},
            {"text": " red", "color": "C00000"},
            {"text": " link", "link": "https://example.com/x"},
        ],
    },
    {"type": "paragraph", "text": "Centered", "align": "center"},
    {
        "type": "list",
        "ordered": False,
        "items": ["alpha", {"text": "beta", "children": ["beta one", {"text": "beta two", "children": ["deep"]}]}],
    },
    {"type": "list", "ordered": True, "items": ["first", "second"]},
    {"type": "list", "ordered": True, "items": ["restarts at one"]},
    {
        "type": "table",
        "header": ["Region", "Units"],
        "rows": [["North", 3], ["South", {"text": 5, "bold": True}], ["East", None]],
        "widths": [2, 1],
        "align": ["left", "right"],
        "caption": "Units by region",
    },
    {"type": "image", "image": "logo.png", "width_cm": 4, "caption": "Our logo", "alt": "a red logo", "align": "right"},
    chart_block("column", "native"),
    {"type": "quote", "text": "Stay curious.", "attribution": "Someone"},
    {"type": "code", "text": "def f(x):\n    return x + 1"},
    {"type": "toc"},
    {"type": "page_break"},
    {"type": "heading", "level": 2, "text": "Appendix"},
]


def test_docx_every_block_type() -> None:
    report, data = build(
        "docx",
        {
            "title": "Q report",
            "language": "en-GB",
            "header": {"text": "Confidential"},
            "footer": {"text": "Page {page} of {pages}"},
            "blocks": ALL_BLOCKS,
        },
        {"logo.png": LOGO},
    )
    d = docx.Document(io.BytesIO(data))
    text = [p.text for p in d.paragraphs]
    assert text[0] == "Quarterly report" and d.paragraphs[0].style.name == "Heading 1"
    assert "Plain, bold italic underlined code red link" in text
    assert any(p.alignment == 1 for p in d.paragraphs)  # centered
    assert "beta two" in text and "deep" in text
    table = d.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["Region", "Units"] and table.rows[3].cells[1].text == ""
    assert (
        "Units by region" in text
        and "Our logo" in text
        and "Stay curious." in text
        and "def f(x):\n    return x + 1" in text
    )
    assert d.sections[0].header.paragraphs[0].text == "Confidential"
    footer_xml = d.sections[0].footer._element.xml
    assert 'w:instr="PAGE"' in footer_xml and 'w:instr="NUMPAGES"' in footer_xml
    doc_xml = d.element.xml
    assert doc_xml.count("<w:hyperlink") == 1 and "TOC" in doc_xml and 'w:type="page"' in doc_xml
    assert d.core_properties.title == "Q report" and d.core_properties.language == "en-GB"
    assert "word/charts/chart1.xml" in names(data) and "word/media/image1.png" in names(data)
    # two numbered lists use separate numbering instances, so the second restarts at 1
    ids = {p._p.pPr.numPr.numId.val for p in d.paragraphs if p._p.pPr is not None and p._p.pPr.numPr is not None}
    assert len(ids) == 3
    assert report["verify"]["ok"] and report["verify"]["reopened_with"] == "python-docx"
    assert report["charts"] == [
        {"type": "column", "title": "column chart", "series": 2, "requested": "native", "render": "native"}
    ]
    assert {b["type"] for b in report["outline"]["blocks"]} >= {"heading", "paragraph", "table"}


def test_pdf_every_block_type() -> None:
    report, data = build(
        "pdf",
        {
            "title": "Q report",
            "header": {"text": "Confidential"},
            "footer": {"text": "Page {page} of {pages}"},
            "blocks": ALL_BLOCKS,
        },
        {"logo.png": LOGO},
    )
    assert data.startswith(b"%PDF-")
    text = pdf_text(data)
    for needle in (
        "Quarterly report",
        "Plain,",
        "beta two",
        "deep",
        "Units by region",
        "Our logo",
        "Stay curious.",
        "Someone",
        "return x + 1",
        "Contents",
        "Appendix",
        "Confidential",
    ):
        assert needle in text, needle
    reader = pypdf.PdfReader(io.BytesIO(data))
    pages = len(reader.pages)
    assert pages >= 2 and f"Page 1 of {pages}" in text and f"Page {pages} of {pages}" in text
    assert reader.metadata is not None and reader.metadata.title == "Q report"
    assert report["verify"] == {"ok": True, "pages": pages, "reopened_with": "pypdf"}
    assert report["charts"][0]["render"] == "image" and report["outline"]["headings"] == [
        "Quarterly report",
        "Appendix",
    ]


def test_pdf_renders_non_latin_text_with_embedded_fonts() -> None:
    _, data = build(
        "pdf",
        {
            "blocks": [
                {"type": "paragraph", "text": "Привет мир, \u03b3\u03b5\u03b9\u03ac \u03c3\u03bf\u03c5, café, €100"}
            ]
        },
    )
    text = pdf_text(data)
    assert "Привет мир" in text and "\u03b3\u03b5\u03b9\u03ac" in text and "café" in text and "€100" in text
    assert b"DejaVu" in data


def test_xlsx_every_feature() -> None:
    sheet: dict[str, Any] = {
        "name": "Sales",
        "header": True,
        "rows": [["Region", "Units", "Price"], ["North", 3, 9.5], ["South", 5, 8], ["East", 2, 10]],
        "cells": [
            {"ref": "E1", "value": "Total", "bold": True, "fill": "FFF2CC"},
            {"ref": "E2", "formula": "=SUM(B2:B4)", "number_format": "#,##0"},
            {"ref": "E3", "date": "2026-01-31"},
            {
                "ref": "E4",
                "value": "=B2*C2",
                "number_format": "0.00",
                "italic": True,
                "color": "C00000",
                "align": "right",
                "wrap": True,
            },
        ],
        "column_widths": {"A": 18, "B": 10},
        "freeze": "A2",
        "tables": [{"name": "SalesTable", "range": "A1:C4"}],
        "merge": ["G1:H1"],
        "charts": [{"chart_type": "column", "title": "Units", "data_range": "A1:B4", "anchor": "J2"}],
        "images": [{"image": "logo.png", "at": "J20", "width_px": 120}],
    }
    other = {"name": "Notes", "rows": [["a", "b"]], "autofilter": "A1:B1"}
    report, data = build("xlsx", {"title": "Sales", "sheets": [sheet, other]}, {"logo.png": LOGO})
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["Sales"]
    assert ws["A2"].value == "North" and ws["E2"].value == "=SUM(B2:B4)" and ws["E4"].value == "=B2*C2"
    assert ws["A1"].font.b and ws["E1"].fill.start_color.rgb == "FFFFF2CC"
    assert ws["E2"].number_format == "#,##0" and ws["E3"].number_format == "yyyy-mm-dd"
    assert ws.freeze_panes == "A2" and ws.column_dimensions["A"].width == 18
    assert "SalesTable" in ws.tables and ws.merged_cells.ranges and len(ws._charts) == 1 and len(ws._images) == 1
    assert wb["Notes"].auto_filter.ref == "A1:B1"
    assert wb.calculation.fullCalcOnLoad is True
    assert report["verify"]["reopened_with"] == "openpyxl"
    assert report["outline"]["sheets"][0]["tables"][0]["name"] == "SalesTable"


def test_pptx_every_slide_kind() -> None:
    spec = {
        "title": "Deck",
        "slides": [
            {"layout": "title", "title": "Big idea", "subtitle": "2026 plan", "notes": "Welcome"},
            {
                "title": "Points",
                "body": [
                    {"type": "paragraph", "text": "Intro"},
                    {"type": "list", "items": ["a", {"text": "b", "children": ["b1"]}]},
                    {"type": "quote", "text": "q"},
                    {"type": "code", "text": "x = 1"},
                ],
            },
            {"layout": "section", "title": "Part two", "subtitle": "Details"},
            {
                "layout": "two_content",
                "title": "Compare",
                "body": [{"type": "paragraph", "text": "left"}],
                "body_right": [{"type": "table", "header": ["A", "B"], "rows": [[1, 2]]}],
            },
            {
                "title": "Chart and picture",
                "body": [chart_block("pie", "native")],
                "images": [{"image": "logo.png", "fit": "placeholder"}],
            },
            {
                "layout": "title_only",
                "title": "Placed",
                "images": [{"image": "logo.png", "x_cm": 2, "y_cm": 4, "width_cm": 6}],
                "charts": [
                    {k: v for k, v in chart_block("line", "image").items() if k != "type"}
                    | {"x_cm": 10, "y_cm": 4, "width_cm": 12}
                ],
            },
            {"layout": "blank"},
        ],
    }
    report, data = build("pptx", spec, {"logo.png": LOGO})
    prs = pptx.Presentation(io.BytesIO(data))
    assert len(prs.slides) == 7 and prs.slide_width == 12_192_000
    assert (
        prs.slides[0].shapes.title.text == "Big idea" and prs.slides[0].notes_slide.notes_text_frame.text == "Welcome"
    )
    body = next(s for s in prs.slides[1].placeholders if s.placeholder_format.idx == 1)
    assert [p.text for p in body.text_frame.paragraphs][:2] == ["Intro", "a"]
    assert any(getattr(s, "has_table", False) and s.has_table for s in prs.slides[3].shapes)
    assert any(s.has_chart for s in prs.slides[4].shapes if hasattr(s, "has_chart"))
    assert sum(1 for s in prs.slides[4].shapes if s.shape_type == 13) == 1  # the picture
    assert [s.shape_type for s in prs.slides[5].shapes].count(13) == 2  # the logo and the image chart
    assert not [s for s in prs.slides[0].placeholders if not s.text_frame.text.strip()]  # no empty placeholders left
    assert report["verify"]["reopened_with"] == "python-pptx" and len(report["outline"]["slides"]) == 7


# ---- charts: every type x render mode x format ---------------------------------------------------


def chart_spec(fmt: str, kind: str, render: str) -> dict[str, Any]:
    block = chart_block(kind, render)
    if fmt in ("docx", "pdf"):
        return {"blocks": [block]}
    if fmt == "pptx":
        return {"slides": [{"title": "Chart", "body": [block]}]}
    scatter = kind == "scatter"
    rows = (
        [["x", "a", "b"], [1, 2, 1], [2, 4, 3], [3, 5, 2]]
        if scatter
        else [
            ["Region", "2025", "2026"],
            *[[c, SERIES[0]["values"][i], SERIES[1]["values"][i]] for i, c in enumerate(CATS)],
        ]
    )
    rng = "A1:B4" if kind in ("pie", "doughnut") else "A1:C4"
    return {
        "sheets": [
            {
                "name": "Data",
                "rows": rows,
                "charts": [
                    {
                        "chart_type": kind,
                        "title": f"{kind} chart",
                        "data_range": rng,
                        "anchor": "F2",
                        "render": render,
                        "x_title": "Region",
                        "y_title": "Units",
                    }
                ],
            }
        ]
    }


@pytest.mark.parametrize("fmt", ["docx", "xlsx", "pptx", "pdf"])
@pytest.mark.parametrize("render", ["native", "image"])
@pytest.mark.parametrize("kind", CHART_TYPES)
def test_chart_matrix(fmt: str, render: str, kind: str) -> None:
    report, data = build(fmt, chart_spec(fmt, kind, render))
    got = report["charts"]
    assert len(got) == 1 and got[0]["type"] == kind
    listing = names(data) if fmt != "pdf" else []
    expected = "image" if (render == "image" or fmt == "pdf") else "native"
    assert got[0]["render"] == expected, got
    if fmt == "pdf":
        assert b"/Subtype /Image" in data or b"/Image" in data
        return
    has_chart = any("/charts/chart" in n and n.endswith(".xml") for n in listing)
    has_png = any(n.endswith(".png") and "/media/" in n for n in listing)
    assert has_chart == (expected == "native") and has_png == (expected == "image"), (listing, got)


def test_native_chart_data_stays_editable_in_the_package() -> None:
    """The chart keeps its data in an embedded workbook, so Office can edit it."""
    for fmt, prefix in (("docx", "word"), ("pptx", "ppt")):
        _, data = build(fmt, chart_spec(fmt, "column", "native"))
        books = [n for n in names(data) if n.startswith(f"{prefix}/embeddings/") and n.endswith(".xlsx")]
        assert len(books) == 1
        wb = openpyxl.load_workbook(io.BytesIO(zipfile.ZipFile(io.BytesIO(data)).read(books[0])))
        assert [c.value for c in wb.active[1]][:3] == [None, "2025", "2026"]


def test_image_charts_are_deterministic_pngs() -> None:
    _, a = build("docx", chart_spec("docx", "stacked", "image"))
    _, b = build("docx", chart_spec("docx", "stacked", "image"))
    png_a = zipfile.ZipFile(io.BytesIO(a)).read("word/media/image1.png")
    assert png_a == zipfile.ZipFile(io.BytesIO(b)).read("word/media/image1.png") and png_a.startswith(b"\x89PNG")


def test_xlsx_image_chart_refuses_formulas_with_advice() -> None:
    spec = {
        "sheets": [
            {
                "name": "D",
                "rows": [["k", "v"], ["a", "=1+1"]],
                "charts": [{"chart_type": "column", "data_range": "A1:B2", "anchor": "D2", "render": "image"}],
            }
        ]
    }
    with pytest.raises(InvalidRequest, match='use render "native"'):
        build("xlsx", spec)


def test_native_chart_failure_falls_back_to_an_image_and_says_so(monkeypatch: pytest.MonkeyPatch) -> None:
    from cerea_documents import charts

    def boom(*_a: object, **_k: object) -> None:
        raise RuntimeError("no chart writer")

    monkeypatch.setattr(charts, "native_part_for_word", boom)
    report, data = build("docx", chart_spec("docx", "column", "native"))
    assert report["charts"][0]["render"] == "image" and "embedded as an image" in report["charts"][0]["note"]
    assert any(n.endswith(".png") for n in names(data)) and "native chart could not be written" in report["notes"][0]


# ---- images ---------------------------------------------------------------------------------------


def test_images_in_every_format_and_a_missing_one_is_named() -> None:
    jpg = jpeg(80, 40)
    _, d = build("docx", {"blocks": [{"type": "image", "image": "a.jpg", "height_cm": 2}]}, {"a.jpg": jpg})
    assert docx.Document(io.BytesIO(d)).inline_shapes[0].height.cm == pytest.approx(2, abs=0.01)
    _, p = build("pdf", {"blocks": [{"type": "image", "image": "a.jpg"}]}, {"a.jpg": jpg})
    assert b"/DCTDecode" in p
    with pytest.raises(InvalidRequest, match=r"images sent: \['a.jpg'\]"):
        build("docx", {"blocks": [{"type": "image", "image": "missing.png"}]}, {"a.jpg": jpg})


def test_oversized_image_is_clamped_to_the_text_width() -> None:
    _, d = build(
        "docx", {"blocks": [{"type": "image", "image": "l", "width_cm": 90}], "page": {"size": "A5"}}, {"l": LOGO}
    )
    assert docx.Document(io.BytesIO(d)).inline_shapes[0].width.cm < 14.8


# ---- theme --------------------------------------------------------------------------------------


def test_theme_overrides_reach_the_styles() -> None:
    theme = {
        "body_font": "Georgia",
        "heading_font": "Verdana",
        "body_size_pt": 12,
        "heading_sizes_pt": {"h1": 30},
        "colors": {"heading": "C00000", "text": "222222"},
        "accent": "00AA55",
        "line_spacing": 1.5,
        "paragraph_spacing_pt": 10,
        "table_style": "grid",
    }
    blocks = [
        {"type": "heading", "level": 1, "text": "H"},
        {"type": "paragraph", "text": "p"},
        {"type": "table", "header": ["a"], "rows": [["b"]]},
    ]
    _, data = build("docx", {"theme": theme, "blocks": blocks})
    d = docx.Document(io.BytesIO(data))
    normal, h1 = d.styles["Normal"], d.styles["Heading 1"]
    assert normal.font.name == "Georgia" and normal.font.size.pt == 12 and normal.paragraph_format.line_spacing == 1.5
    assert normal.paragraph_format.space_after.pt == 10 and normal.font.color.rgb == RGBColor.from_string("222222")
    assert h1.font.name == "Verdana" and h1.font.size.pt == 30 and h1.font.color.rgb == RGBColor.from_string("C00000")
    assert "w:asciiTheme" not in h1.element.xml
    assert 'w:fill="00AA55"' in d.tables[0]._tbl.xml  # accent drives the table header
    assert '<w:insideV w:val="single"' in d.tables[0]._tbl.xml  # grid style
    # the default is neutral and different
    _, base = build("docx", {"blocks": blocks})
    assert docx.Document(io.BytesIO(base)).styles["Normal"].font.name == "Calibri"


def test_theme_in_the_other_formats() -> None:
    theme = {"body_font": "Georgia", "heading_font": "Verdana", "accent": "00AA55", "body_size_pt": 13}
    _, x = build("xlsx", {"theme": theme, "sheets": [{"name": "S", "rows": [["h", "v"], ["a", 1]], "header": True}]})
    ws = openpyxl.load_workbook(io.BytesIO(x))["S"]
    assert ws["A2"].font.name == "Georgia" and ws["A2"].font.sz == 13 and ws["A1"].fill.start_color.rgb == "FF00AA55"
    _, p = build("pptx", {"theme": theme, "slides": [{"title": "T", "body": [{"type": "paragraph", "text": "b"}]}]})
    theme_xml = zipfile.ZipFile(io.BytesIO(p)).read("ppt/theme/theme1.xml").decode()
    assert 'typeface="Verdana"' in theme_xml and 'typeface="Georgia"' in theme_xml and "00AA55" in theme_xml
    _, f = build("pdf", {"theme": {"body_font": "Times New Roman"}, "blocks": [{"type": "paragraph", "text": "serif"}]})
    assert b"DejaVuSerif" in f and b"DejaVuSans-Bold" not in f


def test_page_setup_in_docx_and_pdf() -> None:
    spec = {
        "page": {"size": "Letter", "orientation": "landscape", "margins_cm": {"top": 1, "left": 3}},
        "blocks": [{"type": "paragraph", "text": "x"}],
    }
    sec = docx.Document(io.BytesIO(build("docx", spec)[1])).sections[0]
    assert (
        round(sec.page_width.cm, 1) == 27.9
        and round(sec.page_height.cm, 1) == 21.6
        and sec.top_margin.cm == pytest.approx(1, abs=0.01)
        and sec.left_margin.cm == pytest.approx(3, abs=0.01)
    )
    box = pypdf.PdfReader(io.BytesIO(build("pdf", spec)[1])).pages[0].mediabox
    assert round(float(box.width) / 28.3465, 1) == 27.9


# ---- Markdown shorthand ------------------------------------------------------------------------

MD = """---
title: From Markdown
language: it
page_size: A5
footer: Page {page}
---
# Heading one

Some **bold**, *italic*, `code` and a [link](https://example.com).
A second line of the same paragraph.

## Lists
- one
  - nested
- two

1. first
2. second

| Name | Qty |
|:-----|----:|
| Pens | 3 |
| Ink  | 12 |

> A wise quote

```python
print("hi")
```

![a logo](logo.png)

[TOC]

\\newpage

Last paragraph.
"""


def test_markdown_becomes_the_same_tree() -> None:
    from cerea_documents.spec import validate_spec

    expanded = validate_spec("docx", {"markdown": MD}, LIMITS)
    kinds = [b["type"] for b in expanded["blocks"]]
    assert kinds == [
        "heading",
        "paragraph",
        "heading",
        "list",
        "list",
        "table",
        "quote",
        "code",
        "image",
        "toc",
        "page_break",
        "paragraph",
    ]
    assert expanded["title"] == "From Markdown" and expanded["language"] == "it" and expanded["page"] == {"size": "A5"}
    assert expanded["blocks"][1]["runs"][1] == {"text": "bold", "bold": True}
    assert expanded["blocks"][3]["items"] == ["one", {"text": "nested"}] or expanded["blocks"][3]["items"][0] == {
        "text": "one",
        "children": ["nested"],
    }
    assert expanded["blocks"][5]["align"] == ["left", "right"] and expanded["blocks"][5]["rows"] == [
        ["Pens", "3"],
        ["Ink", "12"],
    ]
    report, data = build("docx", {"markdown": MD}, {"logo.png": LOGO})
    d = docx.Document(io.BytesIO(data))
    assert d.core_properties.title == "From Markdown" and d.tables[0].rows[1].cells[1].text == "3"
    assert any(p.text == "nested" for p in d.paragraphs) and report["verify"]["ok"]
    _, pdf = build("pdf", {"markdown": MD}, {"logo.png": LOGO})
    assert "Last paragraph." in pdf_text(pdf)


@pytest.mark.parametrize(
    ("md", "needle"),
    [
        ("<script>alert(1)</script>", "raw HTML is not accepted"),
        ("text with <b>html</b> inline", "raw HTML is not accepted"),
        ("[x](javascript:alert(1))", "only http, https and mailto"),
        ("```\nnever closed", "never closed"),
        ("---\nbogus: 1\n---\nhi", "unknown keys"),
    ],
)
def test_markdown_refusals(md: str, needle: str) -> None:
    with pytest.raises(InvalidRequest, match=needle):
        build("docx", {"markdown": md})


# ---- determinism --------------------------------------------------------------------------------


def test_the_same_spec_gives_the_same_bytes() -> None:
    cases = {
        "docx": {"title": "t", "blocks": ALL_BLOCKS},
        "pdf": {"title": "t", "blocks": ALL_BLOCKS},
        "xlsx": chart_spec("xlsx", "column", "native") | {"title": "t"},
        "pptx": {"title": "t", "slides": [{"title": "a", "body": [chart_block("column", "native")], "notes": "n"}]},
    }
    first = {f: build(f, s, {"logo.png": LOGO})[1] for f, s in cases.items()}
    time.sleep(2.2)  # zip dates have two-second resolution: a stamp from 'now' would differ
    for fmt, spec in cases.items():
        assert build(fmt, spec, {"logo.png": LOGO})[1] == first[fmt], fmt
    assert zipfile.ZipFile(io.BytesIO(first["docx"])).infolist()[0].date_time == (2000, 1, 1, 0, 0, 0)
