"""PowerPoint engine, ported from the 0.29.0 pptx helper checks, run against the service functions."""

from __future__ import annotations

import io
from typing import Any

import pptx
import pytest
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from cerea_documents import core
from cerea_documents.errors import DocumentError

from .fixtures import jpeg, make_pptx, parts, png, rezip
from .helpers import LIMITS, do_edit, refuses, untouched


@pytest.fixture(scope="module")
def deck() -> tuple[bytes, int]:
    return make_pptx()


@pytest.fixture(scope="module")
def src(deck: tuple[bytes, int]) -> bytes:
    return deck[0]


@pytest.fixture(scope="module")
def pic_id(deck: tuple[bytes, int]) -> int:
    return deck[1]


def edit(src: bytes, ops: list[dict[str, Any]], **kw: Any) -> tuple[dict[str, Any], bytes]:
    return do_edit(src, ops, fmt="pptx", **kw)


def prs(data: bytes) -> Any:
    return pptx.Presentation(io.BytesIO(data))


def test_outline(src: bytes, pic_id: int) -> None:
    o = core.outline(src, "pptx", LIMITS)
    assert [s["layout"] for s in o["slides"]] == ["Title Slide", "Title and Content", "Title Only"]
    s2 = o["slides"][1]
    pic = next(s for s in s2["shapes"] if s["kind"] == "picture")
    assert pic["id"] == pic_id and pic["image"].startswith("ppt/media/") and pic["box"][2:] == [1828800, 1828800]
    assert any(
        s.get("placeholder", {}).get("idx") == 1 and s["text"] == ["Revenue grew", "Costs fell"] for s in s2["shapes"]
    )
    assert "2025" in s2["notes"] and "Total 2025" in str(o["slides"][2]["shapes"])
    assert o["slide_size_cm"] == [25.4, 19.05]


def test_read_slide_text_with_runs(src: bytes) -> None:
    r = core.read(src, "pptx", {"targets": [{"slide": 2}]}, LIMITS)["targets"][0]
    box = next(s for s in r["shapes"] if s["name"].startswith("TextBox"))
    assert [x["text"] for x in box["paragraphs"][0]["runs"]] == ["Rev", "enue 2025 up"]
    assert box["paragraphs"][0]["runs"][0]["bold"] is True and box["paragraphs"][0]["runs"][1]["size_pt"] == 20
    assert r["notes"] == ["Say the number slowly: 2025"]
    with pytest.raises(DocumentError, match="3 slides"):
        core.read(src, "pptx", {"targets": [{"slide": 9}]}, LIMITS)


def test_replace_across_runs(src: bytes) -> None:
    rep, out = edit(src, [{"op": "replace_text", "find": "enue 2025 up", "replace": "enue 2026 down"}])
    assert rep["operations"][0]["result"]["replaced"] == 1 and rep["verify"]["changed"] == ["ppt/slides/slide2.xml"]
    tb = next(s for s in prs(out).slides[1].shapes if s.name.startswith("TextBox"))
    runs = tb.text_frame.paragraphs[0].runs
    assert tb.text_frame.text == "Revenue 2026 down"
    assert [(r.text, bool(r.font.bold)) for r in runs] == [("Rev", True), ("enue 2026 down", False)]
    assert runs[1].font.italic and runs[1].font.size == Pt(20)
    assert rep["diff"] == [
        f"- slide2.xml #{tb.shape_id} {tb.name}: Revenue 2025 up",
        f"+ slide2.xml #{tb.shape_id} {tb.name}: Revenue 2026 down",
    ]
    untouched(src, out, {"ppt/slides/slide2.xml"})
    a, b = parts(src), parts(out)
    assert all(b[n] == a[n] for n in a if "slideLayout" in n or "slideMaster" in n or "theme" in n)


def test_replace_scope_and_notes(src: bytes) -> None:
    _, out = edit(src, [{"op": "replace_text", "find": "2025", "replace": "2026", "slides": [1, 3]}])
    assert prs(out).slides[0].shapes.title.text == "Annual review 2026"
    assert "2025" in prs(out).slides[1].notes_slide.notes_text_frame.text
    _, out = edit(src, [{"op": "replace_text", "find": "2025", "replace": "2026", "notes": True}])
    assert prs(out).slides[1].notes_slide.notes_text_frame.text == "Say the number slowly: 2026"
    refuses(src, [{"op": "replace_text", "find": "no such text", "replace": "x"}], "no match for", fmt="pptx")


def test_set_placeholder_text_keeps_run_formatting(src: bytes) -> None:
    _, out = edit(
        src,
        [
            {"op": "set_placeholder_text", "slide": 1, "placeholder": "title", "text": "Annual review 2026"},
            {"op": "set_placeholder_text", "slide": 2, "placeholder": 1, "text": ["One", "Two", "Three"]},
        ],
    )
    t = prs(out).slides[0].shapes.title.text_frame.paragraphs[0].runs[0]
    assert t.text == "Annual review 2026" and t.font.bold and t.font.color.rgb == RGBColor(0x12, 0x34, 0x56)
    assert prs(out).slides[1].placeholders[1].text_frame.text == "One\nTwo\nThree"
    refuses(
        src,
        [{"op": "set_placeholder_text", "slide": 1, "placeholder": "chart", "text": "x"}],
        "this slide has: ctrTitle (idx 0), subTitle (idx 1)",
        fmt="pptx",
    )
    refuses(
        src, [{"op": "set_placeholder_text", "slide": 9, "placeholder": "title", "text": "x"}], "3 slides", fmt="pptx"
    )


def test_replace_image_crop_shared_and_private_media(src: bytes, pic_id: int) -> None:
    wide = png(80, 20, (0, 0, 200))
    rep, out = edit(
        src, [{"op": "replace_image", "slide": 2, "shape": pic_id, "image": "n.png"}], images={"n.png": wide}
    )
    r = rep["operations"][0]["result"]
    assert r["cropped"] and r["removed_media"] == []  # slide 3 still uses the old file
    pnew = next(s for s in prs(out).slides[1].shapes if s.shape_id == pic_id)
    assert (pnew.left, pnew.top, pnew.width, pnew.height) == (Inches(6), Inches(1), Inches(2), Inches(2))
    assert pnew.image.blob == wide and pnew.crop_left > 0.3 and pnew.crop_top == 0
    v = rep["verify"]
    assert len(v["added"]) == 1 and v["added"][0].startswith("ppt/media/image") and not v["removed"]
    assert "ppt/slides/_rels/slide2.xml.rels" in v["changed"]
    o = core.outline(src, "pptx", LIMITS)
    pic = next(s for s in o["slides"][1]["shapes"] if s["kind"] == "picture")
    private = next(s for s in o["slides"][2]["shapes"] if s["kind"] == "picture" and s["image"] != pic["image"])
    rep, _out2 = edit(
        src,
        [{"op": "replace_image", "slide": 3, "shape": private["id"], "image": "n.png", "fit": "stretch"}],
        images={"n.png": png(30, 30, (9, 9, 9))},
    )
    assert not rep["operations"][0]["result"]["cropped"] and rep["operations"][0]["result"]["removed_media"] == [
        private["image"]
    ]
    assert private["image"] in rep["verify"]["removed"] and len(rep["verify"]["added"]) == 1
    _, out3 = edit(
        src, [{"op": "replace_image", "slide": 2, "shape": pic_id, "image": "j.jpg"}], images={"j.jpg": jpeg(30, 60)}
    )
    assert any(n.endswith(".jpeg") for n in parts(out3))
    jp = next(x for x in prs(out3).slides[1].shapes if x.shape_id == pic_id)
    assert jp.image.content_type == "image/jpeg" and jp.width == Inches(2) and jp.crop_top > 0.2
    refuses(
        src,
        [{"op": "replace_image", "slide": 2, "shape": 9999, "image": "n.png"}],
        "no picture with id 9999",
        fmt="pptx",
        images={"n.png": wide},
    )
    with pytest.raises(DocumentError, match="not a PNG"):
        edit(src, [{"op": "replace_image", "slide": 2, "shape": pic_id, "image": "n.gif"}], images={"n.gif": b"GIF"})


def test_insert_image_default_and_explicit_box(src: bytes) -> None:
    rep, out = edit(
        src,
        [
            {"op": "insert_image", "slide": 1, "image": "a.png", "alt": "diagram"},
            {"op": "insert_image", "slide": 3, "image": "a.png", "box": {"x": 2, "y": 3, "width": 4}},
        ],
        images={"a.png": png(60, 30, (1, 2, 3))},
    )
    s1 = prs(out).slides[0]
    pics = [s for s in s1.shapes if s.shape_type == 13]
    assert len(pics) == 1 and pics[0]._element.xpath(".//p:cNvPr/@descr")[0] == "diagram"
    cx = prs(out).slide_width
    assert abs((pics[0].left + pics[0].width / 2) - cx / 2) < 2 and pics[0].width / pics[0].height == 2.0
    p3 = [s for s in prs(out).slides[2].shapes if s.shape_type == 13][-1]
    assert (p3.left, p3.top, p3.width, p3.height) == (720000, 1080000, 1440000, 720000)
    assert len({s.shape_id for s in s1.shapes}) == len(list(s1.shapes))
    assert len(rep["verify"]["added"]) == 2  # one media part per insertion


def test_duplicate_delete_reorder(src: bytes) -> None:
    rep, out = edit(src, [{"op": "duplicate_slide", "slide": 2}])
    d = prs(out)
    assert (
        len(d.slides) == 4
        and d.slides[2].shapes.title.text == "Results"
        and d.slides[2].slide_layout.name == "Title and Content"
    )
    assert d.slides[1].shapes[3].image.blob == d.slides[2].shapes[3].image.blob
    v = rep["verify"]
    assert v["slides"] == 4 and "ppt/slides/slide4.xml" in v["added"] and not any("media" in n for n in v["added"])
    assert rep["diff"][:2] == [
        "- order: slide1.xml, slide2.xml, slide3.xml",
        "+ order: slide1.xml, slide2.xml, slide4.xml, slide3.xml",
    ]
    untouched(
        src, out, {"ppt/presentation.xml", "ppt/_rels/presentation.xml.rels", "[Content_Types].xml"} | set(v["added"])
    )
    _, o2 = edit(src, [{"op": "delete_slide", "slide": 2}])
    assert len(prs(o2).slides) == 2
    rep, o2 = edit(src, [{"op": "delete_slide", "slide": 2}])
    v = rep["verify"]
    assert "ppt/slides/slide2.xml" in v["removed"] and any(n.startswith("ppt/notesSlides/") for n in v["removed"])
    assert not any("slideLayout" in n or "slideMaster" in n or "theme" in n for n in v["removed"] + v["changed"])
    rep, o3 = edit(src, [{"op": "reorder_slides", "order": [3, 1, 2]}])
    assert [s.shapes.title.text for s in prs(o3).slides] == ["Detail", "Annual review 2025", "Results"]
    assert rep["verify"]["changed"] == ["ppt/presentation.xml"]
    refuses(src, [{"op": "reorder_slides", "order": [1, 1, 2]}], "exactly once", fmt="pptx")
    refuses(src, [{"op": "delete_slide", "slide": 7}], "3 slides", fmt="pptx")
    refuses(
        src,
        [{"op": "delete_slide", "slide": 1}, {"op": "delete_slide", "slide": 1}],
        "deleted by an earlier",
        fmt="pptx",
    )


def test_slide_numbers_refer_to_the_original_deck_within_a_request(src: bytes) -> None:
    ops: list[dict[str, Any]] = [
        {"op": "duplicate_slide", "slide": 1},
        {"op": "delete_slide", "slide": 2},
        {"op": "set_placeholder_text", "slide": 3, "placeholder": "title", "text": "Still the third"},
        {"op": "reorder_slides", "order": [3, 1]},
    ]
    _, out = edit(src, ops)
    titles = [s.shapes.title.text for s in prs(out).slides]
    assert titles == ["Still the third", "Annual review 2025", "Annual review 2025"]


def test_set_notes_existing_and_new(src: bytes) -> None:
    rep, out = edit(
        src,
        [
            {"op": "set_notes", "slide": 2, "text": ["Line one", "Line two"]},
            {"op": "set_notes", "slide": 1, "text": "Fresh notes"},
        ],
    )
    d = prs(out)
    assert d.slides[1].notes_slide.notes_text_frame.text == "Line one\nLine two"
    assert d.slides[0].notes_slide.notes_text_frame.text == "Fresh notes"
    assert rep["operations"][1]["result"]["notes_page_created"] is True
    assert any(n.startswith("ppt/notesSlides/notesSlide") for n in rep["verify"]["added"])
    assert "notes: Fresh notes" in " ".join(rep["diff"])


def test_set_notes_refused_without_a_notes_master() -> None:
    d = pptx.Presentation()
    d.slides.add_slide(d.slide_layouts[0])
    buf = io.BytesIO()
    d.save(buf)
    refuses(buf.getvalue(), [{"op": "set_notes", "slide": 1, "text": "x"}], "no notes master", fmt="pptx")


def test_verify_catches_an_out_of_bounds_change(src: bytes) -> None:
    from cerea_documents import powerpoint
    from cerea_documents.config import Limits
    from cerea_documents.errors import PatchError
    from cerea_documents.pkg import Pkg
    from cerea_documents.verify import verify

    bad = rezip(src, {"ppt/slideLayouts/slideLayout1.xml": parts(src)["ppt/slideLayouts/slideLayout1.xml"] + b" "})
    with pytest.raises(PatchError, match="slideLayout1.xml"):
        verify(Pkg(src, Limits()), bad, Limits(), powerpoint, set(), "pptx")
