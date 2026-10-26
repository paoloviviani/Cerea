"""Templates the person uploads: taking styles from one (`theme.template`) and filling one
(`create_from_template`). Nothing here uses an admin template: the file always comes with the request."""

from __future__ import annotations

import io
import zipfile
from typing import Any

import docx
import openpyxl
import pptx
import pytest
from docx.shared import RGBColor
from openpyxl.styles import Font, NamedStyle
from pptx.util import Inches

from cerea_documents import create
from cerea_documents.config import Limits
from cerea_documents.errors import InvalidDocument, InvalidRequest, UnsafeDocument

from . import hostile
from .fixtures import png

LIMITS = Limits()
LOGO = png(100, 50, (10, 120, 200))


def style_template_docx() -> bytes:
    d = docx.Document()
    d.styles["Normal"].font.name = "Georgia"
    d.styles["Normal"].font.size = docx.shared.Pt(10)
    d.styles["Heading 1"].font.color.rgb = RGBColor.from_string("C00000")
    d.sections[0].header.paragraphs[0].text = "ACME HEADER"
    d.add_paragraph("Template body that must not survive")
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def style_template_pptx() -> bytes:
    prs = pptx.Presentation()
    prs.slide_layouts[1].name = "Brand Content"
    prs.slide_layouts[0].name = "Brand Cover"
    prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = "Template slide that must not survive"
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def style_template_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    wb.active["A1"] = "template cell that must not survive"
    wb.defined_names["Old"] = openpyxl.workbook.defined_name.DefinedName("Old", attr_text="Sheet!$A$1")
    wb.add_named_style(NamedStyle(name="Brand", font=Font(bold=True, color="FFC00000")))
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def make(
    fmt: str, spec: dict[str, Any], template: bytes, images: dict[str, bytes] | None = None
) -> tuple[dict[str, Any], bytes]:
    return create.create({"format": fmt, "spec": spec}, template, images or {}, LIMITS)


# ---- styles from the uploaded file --------------------------------------------------------------


def test_docx_takes_styles_headers_and_page_from_the_template() -> None:
    blocks = [
        {"type": "heading", "level": 1, "text": "Hello"},
        {"type": "paragraph", "text": "Body"},
        {"type": "table", "header": ["a"], "rows": [["b"]]},
    ]
    report, data = make("docx", {"theme": {"template": "brand.docx"}, "blocks": blocks}, style_template_docx())
    d = docx.Document(io.BytesIO(data))
    assert d.styles["Normal"].font.name == "Georgia" and d.styles["Normal"].font.size.pt == 10
    assert d.styles["Heading 1"].font.color.rgb == RGBColor.from_string("C00000")
    assert d.sections[0].header.paragraphs[0].text == "ACME HEADER"
    assert not any("must not survive" in p.text for p in d.paragraphs) and d.paragraphs[0].text == "Hello"
    assert d.paragraphs[0].style.name == "Heading 1" and report["styled_from"] == "template"
    # what the theme names explicitly wins over the template
    _, over = make(
        "docx", {"theme": {"template": "brand.docx", "body_font": "Verdana"}, "blocks": blocks}, style_template_docx()
    )
    o = docx.Document(io.BytesIO(over))
    assert o.styles["Normal"].font.name == "Verdana" and o.styles["Heading 1"].font.color.rgb == RGBColor.from_string(
        "C00000"
    )


def test_docx_template_without_the_needed_styles_degrades_and_says_so() -> None:
    d = docx.Document()
    for name in ("Heading 1", "Quote", "Caption", "List Paragraph"):
        d.styles[name].element.getparent().remove(d.styles[name].element)
    buf = io.BytesIO()
    d.save(buf)
    blocks = [
        {"type": "heading", "level": 1, "text": "H"},
        {"type": "quote", "text": "q"},
        {"type": "list", "items": ["a"]},
    ]
    report, data = make("docx", {"theme": {"template": "x.docx"}, "blocks": blocks}, buf.getvalue())
    assert [p.text for p in docx.Document(io.BytesIO(data)).paragraphs] == ["H", "q", "a"] and report["verify"]["ok"]


def test_pptx_uses_the_templates_layouts_by_name() -> None:
    spec = {
        "theme": {"template": "brand.pptx"},
        "slides": [
            {"layout": "Brand Cover", "title": "Cover", "subtitle": "Sub"},
            {"layout": "brand content", "title": "Body", "body": [{"type": "list", "items": ["a", "b"]}]},
            {"layout": "title_content", "title": "By shape", "body": [{"type": "paragraph", "text": "x"}]},
        ],
    }
    report, data = make("pptx", spec, style_template_pptx())
    prs = pptx.Presentation(io.BytesIO(data))
    assert [s.slide_layout.name for s in prs.slides] == ["Brand Cover", "Brand Content", "Brand Content"]
    assert len(prs.slides) == 3 and all("must not survive" not in (s.shapes.title.text or "") for s in prs.slides)
    assert prs.slide_width == 9_144_000  # the template's own size is kept
    assert report["verify"]["ok"]


def test_pptx_unknown_layout_lists_the_templates_layouts() -> None:
    with pytest.raises(InvalidRequest) as ei:
        make(
            "pptx",
            {"theme": {"template": "b.pptx"}, "slides": [{"layout": "Nope", "title": "t"}]},
            style_template_pptx(),
        )
    assert "Brand Content" in str(ei.value) and "Brand Cover" in ei.value.details["available"]


def test_xlsx_takes_named_styles_from_the_template() -> None:
    sheet = {"name": "Data", "cells": [{"ref": "A1", "value": "Total", "style": "Brand"}, {"ref": "A2", "value": 1}]}
    report, data = make("xlsx", {"theme": {"template": "t.xlsx"}, "sheets": [sheet]}, style_template_xlsx())
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Data"] and wb["Data"]["A1"].style == "Brand" and wb["Data"]["A1"].font.b
    assert wb["Data"]["A2"].font.name == "Calibri" and "Old" not in wb.defined_names and report["verify"]["ok"]
    with pytest.raises(InvalidRequest, match=r"named style 'Missing'.*Brand"):
        make(
            "xlsx",
            {
                "theme": {"template": "t.xlsx"},
                "sheets": [{"name": "D", "cells": [{"ref": "A1", "value": 1, "style": "Missing"}]}],
            },
            style_template_xlsx(),
        )


def test_template_must_match_the_format_and_open() -> None:
    with pytest.raises(InvalidRequest, match="template is a pptx file but a docx"):
        make(
            "docx", {"theme": {"template": "x"}, "blocks": [{"type": "paragraph", "text": "x"}]}, style_template_pptx()
        )
    with pytest.raises(InvalidDocument):
        make("docx", {"theme": {"template": "x"}, "blocks": [{"type": "paragraph", "text": "x"}]}, b"not a zip")


# ---- hostile templates ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("name", "build", "code"),
    [
        ("macros", hostile.docm, UnsafeDocument),
        ("xxe", hostile.docx_xxe, UnsafeDocument),
        ("entity_bomb", hostile.docx_laughs, UnsafeDocument),
        ("zip_bomb", hostile.zip_bomb_ratio, UnsafeDocument),
        ("path_traversal", hostile.path_traversal, UnsafeDocument),
        ("many_entries", hostile.zip_many_entries, UnsafeDocument),
    ],
)
def test_hostile_templates_are_refused_for_styles_and_for_filling(name: str, build: Any, code: type[Exception]) -> None:
    data = build()
    with pytest.raises(code) as ei:
        make("docx", {"theme": {"template": "t.docx"}, "blocks": [{"type": "paragraph", "text": "x"}]}, data)
    if name == "macros":
        assert ei.value.reason == "active_content" and "vbaProject" in str(ei.value)  # type: ignore[attr-defined]
    with pytest.raises(code):
        create.create_from_template({"values": {"a": "b"}}, data, "docx", {}, LIMITS)


def test_macro_templates_are_refused_in_every_format() -> None:
    from .fixtures import rezip

    for fmt, src, part in (
        ("xlsx", style_template_xlsx(), "xl/vbaProject.bin"),
        ("pptx", style_template_pptx(), "ppt/vbaProject.bin"),
    ):
        bad = rezip(src, {}, {part: b"\x00macro"})
        with pytest.raises(UnsafeDocument, match="never carry active content"):
            create.create_from_template({"values": {"a": "b"}, "format": fmt}, bad, fmt, {}, LIMITS)


# ---- filling ------------------------------------------------------------------------------------


def fill_template_docx() -> bytes:
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "For {{client}}"
    d.sections[0].footer.paragraphs[0].text = "Ref {{ref}}"
    d.add_paragraph("Dear {{client}},")
    p = d.add_paragraph()
    for t in ("The total is {", "{amou", "nt}} euro."):  # split across runs
        p.add_run(t).bold = t.startswith("{")
    d.add_paragraph("{{logo}}")
    d.add_paragraph("Untouched {{spare}}")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Client"
    t.rows[0].cells[1].text = "{{client}}"
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def test_docx_placeholders_in_body_tables_headers_footers_and_an_image() -> None:
    values = {"client": "Beta plc", "ref": "A-17", "amount": 1250.5, "logo": {"image": "logo.png", "width_cm": 3}}
    report, data = create.create_from_template(
        {"values": values}, fill_template_docx(), "docx", {"logo.png": LOGO}, LIMITS
    )
    d = docx.Document(io.BytesIO(data))
    text = [p.text for p in d.paragraphs]
    assert "Dear Beta plc," in text and "The total is 1250.5 euro." in text and "Untouched {{spare}}" in text
    assert (
        d.sections[0].header.paragraphs[0].text == "For Beta plc"
        and d.sections[0].footer.paragraphs[0].text == "Ref A-17"
    )
    assert d.tables[0].rows[0].cells[1].text == "Beta plc"
    assert len(d.inline_shapes) == 1 and not any("{{logo}}" in t for t in text)
    assert report["unfilled"] == ["spare"] and report["filled"] == ["amount", "client", "logo", "ref"]
    assert report["verify"]["ok"] and report["diff"] and "word/media/image1.png" in report["verify"]["added"]
    assert report["styled_from"] == "template"


def test_unknown_placeholder_lists_what_exists_and_writes_nothing() -> None:
    with pytest.raises(InvalidRequest) as ei:
        create.create_from_template({"values": {"client": "x", "nope": "y"}}, fill_template_docx(), "docx", {}, LIMITS)
    assert (
        "['nope']" in str(ei.value)
        and "client" in ei.value.details["available"]
        and ei.value.reason == "unknown_placeholder"
    )
    with pytest.raises(InvalidRequest, match="holds only that placeholder"):
        create.create_from_template(
            {"values": {"client": {"image": "logo.png"}}}, fill_template_docx(), "docx", {"logo.png": LOGO}, LIMITS
        )


def fill_template_pptx() -> bytes:
    prs = pptx.Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Report for {{client}}"
    box = s.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1))
    box.text_frame.text = "Prepared {{date}}"
    logo = s.shapes.add_textbox(Inches(5), Inches(2), Inches(2), Inches(1))
    logo.text_frame.text = "{{logo}}"
    s.notes_slide.notes_text_frame.text = "Notes for {{client}}"
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def test_pptx_placeholders_notes_and_a_picture_in_a_marked_box() -> None:
    values = {"client": "Beta plc", "date": "26 Oct 2026", "logo": {"image": "logo.png"}}
    report, data = create.create_from_template(
        {"values": values}, fill_template_pptx(), "pptx", {"logo.png": LOGO}, LIMITS
    )
    prs = pptx.Presentation(io.BytesIO(data))
    slide = prs.slides[0]
    texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame]
    assert "Report for Beta plc" in texts and "Prepared 26 Oct 2026" in texts and "{{logo}}" not in " ".join(texts)
    assert slide.notes_slide.notes_text_frame.text == "Notes for Beta plc"
    pics = [s for s in slide.shapes if s.shape_type == 13]
    assert len(pics) == 1 and pics[0].left == Inches(5)
    assert report["unfilled"] == [] and report["verify"]["ok"]


def fill_template_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quote"
    ws["A1"] = "Quote for {{client}}"
    ws["A2"] = "{{total}}"
    ws["A3"] = "{{client}} owes {{total}} by {{due}}"
    ws["C1"], ws["C2"], ws["C3"] = "item", "pen", "ink"
    ws["E1"] = "{{logo}}"
    wb.defined_names["Items"] = openpyxl.workbook.defined_name.DefinedName("Items", attr_text="Quote!$C$2:$C$3")
    wb.defined_names["Owner"] = openpyxl.workbook.defined_name.DefinedName("Owner", attr_text="Quote!$B$1")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_xlsx_placeholder_cells_and_named_ranges() -> None:
    values = {
        "client": "Beta plc",
        "total": 99.5,
        "due": "Friday",
        "Items": ["paper", "toner"],
        "Owner": "Ada",
        "logo": {"image": "logo.png"},
    }
    report, data = create.create_from_template(
        {"values": values}, fill_template_xlsx(), "xlsx", {"logo.png": LOGO}, LIMITS
    )
    ws = openpyxl.load_workbook(io.BytesIO(data))["Quote"]
    assert ws["A1"].value == "Quote for Beta plc"
    assert ws["A2"].value == 99.5  # a cell that is only the placeholder keeps the number's type
    assert ws["A3"].value == "Beta plc owes 99.5 by Friday"
    assert (ws["C2"].value, ws["C3"].value, ws["B1"].value) == ("paper", "toner", "Ada")
    assert ws["E1"].value is None and len(ws._images) == 1
    assert report["verify"]["ok"] and report["unfilled"] == []
    with pytest.raises(InvalidRequest, match="3 values for the named range 'Items'|has 2 cells"):
        create.create_from_template({"values": {"Items": ["a", "b", "c"]}}, fill_template_xlsx(), "xlsx", {}, LIMITS)
    with pytest.raises(InvalidRequest, match=r"cell placeholder"):
        create.create_from_template({"values": {"client": ["a"]}}, fill_template_xlsx(), "xlsx", {}, LIMITS)
    with pytest.raises(InvalidRequest, match="has no .*named \\['nothing'\\]"):
        create.create_from_template({"values": {"nothing": 1}}, fill_template_xlsx(), "xlsx", {}, LIMITS)


def test_filling_touches_only_what_it_fills() -> None:
    src = fill_template_docx()
    _, out = create.create_from_template({"values": {"client": "Beta"}}, src, "docx", {}, LIMITS)
    a, b = zipfile.ZipFile(io.BytesIO(src)), zipfile.ZipFile(io.BytesIO(out))
    changed = {n for n in a.namelist() if a.read(n) != b.read(n)}
    assert changed <= {"word/document.xml", "word/header1.xml"}, changed


def test_a_filled_file_is_reproducible() -> None:
    import time

    src = fill_template_docx()
    first = create.create_from_template({"values": {"client": "Beta"}}, src, "docx", {}, LIMITS)[1]
    time.sleep(2.2)
    assert create.create_from_template({"values": {"client": "Beta"}}, src, "docx", {}, LIMITS)[1] == first
