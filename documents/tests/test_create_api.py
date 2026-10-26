"""Creation through the HTTP API: capabilities, multipart in and out, error shapes, hostile inputs."""

from __future__ import annotations

import base64
import io
import json
from collections.abc import Iterator
from typing import Any

import docx
import pypdf
import pytest
from jsonschema import Draft202012Validator
from starlette.testclient import TestClient

from cerea_documents import runner

from . import hostile
from .conftest import AUTH, make_client, parse_multipart
from .fixtures import png

LOGO = png(60, 30, (30, 90, 200))
SPEC = {
    "title": "API doc",
    "blocks": [
        {"type": "heading", "level": 1, "text": "Hello"},
        {"type": "image", "image": "logo.png"},
        {"type": "chart", "chart_type": "line", "categories": ["a", "b"], "series": [{"name": "s", "values": [1, 2]}]},
    ],
}


@pytest.fixture(scope="module")
def client() -> Iterator[TestClient]:
    saved = runner.STATS["peak_rss_kb"]
    with make_client() as c:
        yield c
    runner.STATS["peak_rss_kb"] = saved  # chart libraries make a bigger worker than the hostile suite's bound


def create(
    client: TestClient,
    body: dict[str, Any] | str,
    images: dict[str, bytes] | None = None,
    template: tuple[str, bytes] | None = None,
    path: str = "/v1/create",
    headers: dict[str, str] | None = None,
) -> Any:
    files: list[Any] = []
    if template:
        files.append(("file", (template[0], template[1], "application/octet-stream")))
    for n, b in (images or {}).items():
        files.append(("image", (n, b, "application/octet-stream")))
    files.append(
        ("body", (None, body if isinstance(body, str) else json.dumps(body)))
    )  # a multipart field, like the chat's FormData
    return client.post(path, files=files, headers=AUTH if headers is None else headers)


def test_capabilities_publish_the_creation_schemas(client: TestClient) -> None:
    caps = client.get("/v1/capabilities", headers=AUTH).json()
    assert (
        caps["create"] is True
        and caps["create_from_template"] is True
        and caps["create_formats"] == ["docx", "xlsx", "pptx", "pdf"]
    )
    schemas = caps["create_schemas"]
    assert set(schemas) == {"docx", "xlsx", "pptx", "pdf"}
    assert {"heading", "paragraph", "list", "table", "image", "chart", "page_break", "toc", "quote", "code"} == set(
        schemas["docx"]["blocks"]
    )
    assert set(schemas["pptx"]["blocks"]) == {"paragraph", "list", "table", "image", "chart", "quote", "code"}
    assert "sheet" in schemas["xlsx"] and "slide" in schemas["pptx"]
    theme = schemas["docx"]["document"]["properties"]["theme"]["properties"]
    assert "Calibri" in theme["body_font"]["enum"] and "template" in theme
    for entry in schemas.values():
        for schema in [entry["document"], *entry["blocks"].values()]:
            Draft202012Validator.check_schema(schema)
    assert schemas["docx"]["blocks"]["chart"]["properties"]["chart_type"]["enum"] == [
        "bar",
        "column",
        "stacked",
        "line",
        "area",
        "pie",
        "doughnut",
        "scatter",
    ]
    assert schemas["docx"]["blocks"]["chart"]["properties"]["render"]["enum"] == ["native", "image"]


def test_create_returns_result_and_file_as_multipart(client: TestClient) -> None:
    r = create(client, {"format": "docx", "spec": SPEC, "file_name": "my report"}, {"logo.png": LOGO})
    assert r.status_code == 200, r.text
    result, file = parse_multipart(r)
    assert result["file_name"] == "my report.docx" and result["format"] == "docx" and result["file_bytes"] == len(file)
    assert result["content_type"].endswith("wordprocessingml.document") and result["verify"]["ok"]
    assert result["charts"][0]["render"] == "native" and result["outline"]["blocks"][0]["text"] == "Hello"
    assert docx.Document(io.BytesIO(file)).paragraphs[0].text == "Hello"


def test_create_as_json_for_every_format(client: TestClient) -> None:
    specs = {
        "xlsx": {"sheets": [{"name": "S", "rows": [["a", 1]]}]},
        "pptx": {"slides": [{"title": "t"}]},
        "pdf": {"blocks": [{"type": "paragraph", "text": "pdf text"}]},
    }
    for fmt, spec in specs.items():
        r = create(client, {"format": fmt, "spec": spec}, headers={**AUTH, "Accept": "application/json"})
        assert r.status_code == 200, r.text
        j = r.json()
        raw = base64.b64decode(j["file_base64"])
        assert j["file_name"] == f"document.{fmt}" and len(raw) == j["file_bytes"]
        assert raw.startswith(b"%PDF") if fmt == "pdf" else raw.startswith(b"PK")
        if fmt == "pdf":
            assert (
                j["content_type"] == "application/pdf"
                and "pdf text" in pypdf.PdfReader(io.BytesIO(raw)).pages[0].extract_text()
            )


def test_file_names_are_made_safe(client: TestClient) -> None:
    r = create(
        client,
        {
            "format": "pdf",
            "spec": {"blocks": [{"type": "paragraph", "text": "x"}]},
            "file_name": "../../etc/pass wd.docx",
        },
        headers={**AUTH, "Accept": "application/json"},
    )
    name = r.json()["file_name"]
    assert "/" not in name and name.endswith(".pdf") and not name.endswith(".docx.pdf")


def test_create_with_a_style_template_file_part(client: TestClient) -> None:
    from .test_create_templates import style_template_docx

    spec = {"theme": {"template": "brand.docx"}, "blocks": [{"type": "paragraph", "text": "x"}]}
    r = create(
        client,
        {"format": "docx", "spec": spec},
        template=("brand.docx", style_template_docx()),
        headers={**AUTH, "Accept": "application/json"},
    )
    assert r.status_code == 200 and r.json()["styled_from"] == "template"


def test_create_from_template_over_http(client: TestClient) -> None:
    from .test_create_templates import fill_template_docx

    body = {"values": {"client": "Beta plc", "logo": {"image": "logo.png"}}}
    r = create(client, body, {"logo.png": LOGO}, ("tpl.docx", fill_template_docx()), "/v1/create_from_template")
    assert r.status_code == 200, r.text
    result, file = parse_multipart(r)
    assert result["file_name"] == "tpl.docx" and "Dear Beta plc," in [
        p.text for p in docx.Document(io.BytesIO(file)).paragraphs
    ]
    assert result["filled"] == ["client", "logo"] and result["verify"]["ok"]
    missing = create(client, body, path="/v1/create_from_template")
    assert missing.status_code == 422 and missing.json()["error"]["reason"] == "file"


def test_endpoints_need_the_secret(client: TestClient) -> None:
    for path in ("/v1/create", "/v1/create_from_template"):
        r = client.post(path, files=[("body", (None, "{}"))])
        assert r.status_code == 401 and r.json()["error"]["code"] == "unauthorized"


def error(r: Any) -> dict[str, Any]:
    assert r.headers["content-type"].startswith("application/json"), r.text[:200]
    return r.json()["error"]  # type: ignore[no-any-return]


def test_errors_are_actionable_json(client: TestClient) -> None:
    e = error(create(client, {"format": "docx", "spec": {"blocks": [{"type": "banner"}]}}))
    assert (
        e["code"] == "invalid_request"
        and "unknown block type 'banner'" in e["message"]
        and "heading" in e["details"]["available"]
    )
    e = error(create(client, {"format": "rtf", "spec": {}}))
    assert "is not one of" in e["message"]
    e = error(create(client, "not json"))
    assert e["reason"] == "body"
    e = error(create(client, {"format": "docx", "spec": {"blocks": [{"type": "image", "image": "nope.png"}]}}))
    assert "images sent: none" in e["message"] and e["reason"] == "spec"
    r = client.post("/v1/create", headers=AUTH)
    assert r.status_code == 422


# ---- hostile ------------------------------------------------------------------------------------


def test_oversized_spec_is_refused_without_work(client: TestClient) -> None:
    big = {"format": "docx", "spec": {"blocks": [{"type": "paragraph", "text": "x" * 15_000}] * 90}}
    r = create(client, big)
    assert r.status_code == 413 and error(r)["reason"] == "body_size"
    with make_client(max_request_bytes=20_000) as small:
        r = create(small, {"format": "pdf", "spec": {"blocks": [{"type": "paragraph", "text": "y" * 30_000}]}})
        assert r.status_code == 413 and error(r)["code"] == "too_large"


def test_deeply_nested_json_is_a_clean_422_not_a_crash(client: TestClient) -> None:
    for depth in (30, 5000, 200_000):
        text = '{"format":"docx","spec":{"blocks":[' + "[" * depth + "]" * depth + "]}}"
        r = create(client, text)
        assert r.status_code == 422 and error(r)["reason"] == "too_deep", (depth, r.text[:150])
    nested = {"text": "x"}
    for _ in range(40):
        nested = {"text": "x", "children": [nested]}
    r = create(client, {"format": "docx", "spec": {"blocks": [{"type": "list", "items": [nested]}]}})
    assert r.status_code == 422 and error(r)["reason"] == "too_deep"


@pytest.mark.parametrize(
    ("name", "image", "reason"),
    [
        ("huge.png", hostile.huge_png(), "image_pixels"),
        ("x.png", hostile.svg_image(), "image_type"),
        ("y.png", hostile.png_signature_jpeg_body(), "image_damaged"),
        ("g.gif", hostile.gif_garbage(), "image_damaged"),
    ],
)
def test_hostile_images_are_refused(client: TestClient, name: str, image: bytes, reason: str) -> None:
    r = create(client, {"format": "docx", "spec": {"blocks": [{"type": "image", "image": name}]}}, {name: image})
    assert r.status_code == 422 and error(r)["reason"] == reason, r.text[:200]


def test_hostile_templates_over_http(client: TestClient) -> None:
    spec = {"theme": {"template": "t"}, "blocks": [{"type": "paragraph", "text": "x"}]}
    for build, reason in (
        (hostile.docm, "active_content"),
        (hostile.docx_xxe, "doctype"),
        (hostile.zip_bomb_ratio, None),
        (hostile.truncated_zip, None),
    ):
        r = create(client, {"format": "docx", "spec": spec}, template=("t.docx", build()))
        assert r.status_code == 422, (build.__name__, r.status_code, r.text[:200])
        if reason:
            assert error(r)["reason"] == reason
    assert (
        error(
            create(client, {"values": {"a": "b"}}, template=("t.docx", hostile.docm()), path="/v1/create_from_template")
        )["code"]
        == "unsafe_document"
    )


def test_a_creation_that_needs_too_much_memory_is_a_clean_refusal() -> None:
    """Creation runs under the same worker limits as edits: a tiny address-space cap gives 413, never a 500."""
    with make_client(worker_memory_bytes=150 * 1024 * 1024) as tight:
        chart = {"type": "chart", "chart_type": "line", "categories": ["a"], "series": [{"name": "s", "values": [1]}]}
        r = create(tight, {"format": "pdf", "spec": {"blocks": [chart]}})
        assert r.status_code == 413 and error(r)["code"] == "resource_limit", r.text[:200]
        ok = create(
            tight, {"format": "pdf", "spec": {"blocks": [{"type": "paragraph", "text": "no chart libraries needed"}]}}
        )
        assert ok.status_code == 200  # the service keeps serving
