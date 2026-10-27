"""Slide backgrounds and free shapes in created decks (ADR 0103): the XML written, the stacking order, the
bounds check and the warnings. The layout rules of test_pptx_layout.py still hold for decks without shapes."""

from __future__ import annotations

import io
import os
from typing import Any

import pptx
import pytest
from lxml import etree
from pptx.enum.shapes import MSO_SHAPE_TYPE

from cerea_documents import create, spec
from cerea_documents.config import Limits
from cerea_documents.errors import InvalidRequest
from cerea_documents.make_pptx import contrast

from .fixtures import png
from .test_create_templates import style_template_pptx
from .test_pptx_layout import to_pdf

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
IMG = {
    "hd.png": png(160, 90, (60, 60, 60)),
    "wide.png": png(200, 100, (10, 90, 200)),
    "tall.png": png(100, 200, (200, 90, 10)),
}
CM = 360_000


def build(
    slides: list[dict[str, Any]], theme: dict[str, Any] | None = None, size: str = "16:9", template: bytes | None = None
) -> tuple[dict[str, Any], Any, bytes]:
    body: dict[str, Any] = {"slides": slides, "size": size}
    if theme:
        body["theme"] = theme
    report, data = create.create({"format": "pptx", "spec": body}, template, IMG, Limits())
    return report, pptx.Presentation(io.BytesIO(data)), data


def bg(slide: Any) -> Any:
    found = slide._element.xpath("p:cSld/p:bg")
    return found[0] if found else None


def rect(**kw: Any) -> dict[str, Any]:
    return {"kind": "rect", "x_cm": 2, "y_cm": 5, "w_cm": 6, "h_cm": 4, **kw}


# ---- backgrounds ---------------------------------------------------------------------------------


def test_a_colour_background_is_the_slides_own_solid_fill() -> None:
    _, prs, _ = build([{"layout": "blank", "background": {"color": "#1F3864"}}])
    node = bg(prs.slides[0])
    assert node is not None
    assert node.xpath("p:bgPr/a:solidFill/a:srgbClr/@val") == ["1F3864"]
    assert node.getparent().index(node) == 0  # before the shape tree, where PowerPoint expects it
    assert not list(prs.slides[0].shapes)  # a background, not a rectangle


def test_a_gradient_background_is_a_native_two_stop_fill_at_the_angle() -> None:
    _, prs, _ = build([{"layout": "blank", "background": {"gradient": ["1F3864", "2F80ED"], "angle": 45}}])
    node = bg(prs.slides[0])
    assert node.xpath("p:bgPr/a:gradFill/a:gsLst/a:gs/@pos") == ["0", "100000"]
    assert node.xpath("p:bgPr/a:gradFill/a:gsLst/a:gs/a:srgbClr/@val") == ["1F3864", "2F80ED"]
    assert node.xpath("p:bgPr/a:gradFill/a:lin/@ang") == [str(45 * 60_000)]
    _, prs, _ = build([{"layout": "blank", "background": {"gradient": ["000000", "FFFFFF"]}}])
    assert bg(prs.slides[0]).xpath("p:bgPr/a:gradFill/a:lin/@ang") == [str(90 * 60_000)]  # downwards by default


@pytest.mark.parametrize(("image", "size"), [("wide.png", (178, 100)), ("tall.png", (100, 56)), ("hd.png", (160, 90))])
def test_an_image_background_is_cropped_to_cover_not_stretched(image: str, size: tuple[int, int]) -> None:
    from PIL import Image

    _, prs, _ = build([{"layout": "blank", "background": {"image": image, "fit": "cover"}}])
    slide = prs.slides[0]
    blip = bg(slide).xpath("p:bgPr/a:blipFill")[0]
    part = slide.part.related_part(blip.xpath("a:blip/@r:embed")[0])
    assert part.content_type == "image/png"
    assert blip.xpath("a:stretch/a:fillRect")  # the cropped picture fills the slide edge to edge
    with Image.open(io.BytesIO(part.blob)) as shown:
        assert shown.width / shown.height == pytest.approx(prs.slide_width / prs.slide_height, abs=0.02)
        assert (shown.width, shown.height) == size  # the biggest slide-shaped part of the picture, centred
        assert shown.getpixel((0, 0)) == Image.open(io.BytesIO(IMG[image])).convert("RGB").getpixel((0, 0))


def test_the_cover_box_is_centred_and_has_the_slides_proportions() -> None:
    from cerea_documents.make_pptx import _cover_box

    assert _cover_box(200, 100, 16, 9) == (11, 0, 189, 100)
    assert _cover_box(100, 200, 16, 9) == (0, 72, 100, 128)
    assert _cover_box(1920, 1080, 16, 9) == (0, 0, 1920, 1080)
    for iw, ih in ((4000, 3000), (3000, 4000), (1000, 1000)):
        left, top, right, bottom = _cover_box(iw, ih, 12_192_000, 6_858_000)
        assert (right - left) / (bottom - top) == pytest.approx(16 / 9, rel=2e-3)
        assert abs(left - (iw - right)) <= 1 and abs(top - (ih - bottom)) <= 1  # centred


def test_a_picture_with_the_slides_proportions_is_embedded_as_it_came() -> None:
    _, prs, _ = build([{"layout": "blank", "background": {"image": "hd.png", "fit": "cover"}}])
    slide = prs.slides[0]
    blip = bg(slide).xpath("p:bgPr/a:blipFill")[0]
    assert slide.part.related_part(blip.xpath("a:blip/@r:embed")[0]).blob == IMG["hd.png"]


def test_a_theme_default_applies_to_slides_without_their_own_background() -> None:
    _, prs, _ = build(
        [{"title": "a"}, {"title": "b", "background": {"color": "FFFFFF"}}],
        theme={"slide_background": {"gradient": ["112233", "445566"]}},
    )
    assert bg(prs.slides[0]).xpath("p:bgPr/a:gradFill")
    assert bg(prs.slides[1]).xpath("p:bgPr/a:solidFill")  # the slide's own wins


def test_a_missing_background_image_is_a_request_error() -> None:
    with pytest.raises(InvalidRequest, match="nope.png"):
        build([{"background": {"image": "nope.png", "fit": "cover"}}])


def test_dark_text_on_a_dark_solid_background_warns_without_changing_the_text() -> None:
    report, prs, _ = build([{"title": "T", "background": {"color": "101820"}}])
    assert any("slides[0].background" in n and "contrast" in n and "text colour" in n for n in report["notes"])
    run = prs.slides[0].shapes.title.text_frame.paragraphs[0].runs[0]
    assert str(run.font.color.rgb) == "1F2933"  # the text colour is left as it was
    ok, _, _ = build([{"title": "T", "background": {"color": "FFFFFF"}}])
    assert not [n for n in ok.get("notes", []) if "contrast" in n]
    fixed, _, _ = build(
        [{"title": "T", "background": {"color": "101820"}}], theme={"colors": {"text": "FFFFFF", "heading": "FFFFFF"}}
    )
    assert not [n for n in fixed.get("notes", []) if "contrast" in n]


def test_the_contrast_ratio_is_wcag() -> None:
    assert contrast("000000", "FFFFFF") == pytest.approx(21.0)
    assert contrast("777777", "FFFFFF") == pytest.approx(4.48, abs=0.02)


# ---- shapes --------------------------------------------------------------------------------------


def test_every_shape_kind_is_written_at_its_place() -> None:
    kinds = {
        "rect": "rect",
        "rounded_rect": "roundRect",
        "ellipse": "ellipse",
        "chevron": "chevron",
        "text": "rect",
    }
    shapes = [
        {"kind": k, "x_cm": 1 + i, "y_cm": 2, "w_cm": 3, "h_cm": 2, "fill": "none" if k == "text" else "4C9F70"}
        for i, k in enumerate(kinds)
    ]
    shapes += [
        {"kind": "line", "x_cm": 1, "y_cm": 10, "w_cm": 8, "h_cm": 0, "line": {"color": "C0504D", "width_pt": 3}},
        {"kind": "arrow", "x_cm": 1, "y_cm": 12, "w_cm": 8, "h_cm": 0},
    ]
    _, prs, _ = build([{"layout": "blank", "shapes": shapes}])
    drawn = list(prs.slides[0].shapes)
    assert len(drawn) == 7
    for el, (kind, preset) in zip(drawn[:5], kinds.items(), strict=True):
        assert el._element.xpath("p:spPr/a:prstGeom/@prst") == [preset], kind
    assert [int(el.left) for el in drawn[:5]] == [(1 + i) * CM for i in range(5)]
    assert drawn[0]._element.xpath("p:spPr/a:solidFill/a:srgbClr/@val") == ["4C9F70"]
    assert drawn[4]._element.xpath("p:spPr/a:noFill")  # fill: none
    line, arrow = drawn[5], drawn[6]
    assert line.shape_type == MSO_SHAPE_TYPE.LINE and (int(line.width), int(line.height)) == (8 * CM, 0)
    assert line._element.xpath("p:spPr/a:ln/@w") == ["38100"]
    assert line._element.xpath("p:spPr/a:ln/a:solidFill/a:srgbClr/@val") == ["C0504D"]
    assert arrow._element.xpath("p:spPr/a:ln/a:tailEnd/@type") == ["triangle"]
    assert not arrow._element.xpath("p:spPr/a:ln/a:headEnd")
    assert all(not el._element.xpath("p:style") for el in drawn)  # no theme shadow or outline


def test_a_rounded_rect_takes_its_radius_in_cm() -> None:
    _, prs, _ = build([{"layout": "blank", "shapes": [{**rect(kind="rounded_rect", radius=1), "w_cm": 8, "h_cm": 4}]}])
    shape = prs.slides[0].shapes[0]
    assert shape.adjustments[0] == pytest.approx(0.25)  # 1 cm of the 4 cm short side
    _, prs, _ = build([{"layout": "blank", "shapes": [rect(kind="rounded_rect", radius=50)]}])
    assert prs.slides[0].shapes[0].adjustments[0] == 0.5  # never rounder than a pill


def test_shape_text_takes_size_alignment_and_runs() -> None:
    shapes = [
        rect(text="One\nTwo", size_pt=20, align="left", valign="top", fill="FFFFFF"),
        rect(
            x_cm=12,
            runs=[{"text": "Big ", "bold": True, "color": "C0504D"}, {"text": "small"}],
            size_pt=16,
            fill="FFFFFF",
        ),
    ]
    _, prs, _ = build([{"layout": "blank", "shapes": shapes}])
    first, second = prs.slides[0].shapes
    assert [p.text for p in first.text_frame.paragraphs] == ["One", "Two"]
    assert first.text_frame.paragraphs[0].runs[0].font.size.pt == 20
    assert first._element.xpath("p:txBody/a:bodyPr/@anchor") == ["t"]
    assert first.text_frame.paragraphs[0].alignment is not None
    runs = second.text_frame.paragraphs[0].runs
    assert runs[0].font.bold and str(runs[0].font.color.rgb) == "C0504D"
    assert second._element.xpath("p:txBody/a:bodyPr/@anchor") == ["ctr"]  # centred in a shape by default


def test_text_on_a_dark_fill_is_made_readable_unless_the_runs_say_otherwise() -> None:
    _, prs, _ = build([{"layout": "blank", "shapes": [rect(fill="101820", text="Card")]}])
    run = prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert str(run.font.color.rgb) == "FFFFFF"


def test_shapes_sit_behind_the_placeholders_in_list_order_and_front_ones_on_top() -> None:
    shapes = [
        rect(fill="111111", text="a"),
        rect(fill="222222", text="b"),
        rect(fill="333333", text="c", layer="front"),
    ]
    _, prs, _ = build(
        [
            {
                "layout": "title_content",
                "title": "Title",
                "body": [{"type": "paragraph", "text": "Body"}],
                "shapes": shapes,
            }
        ]
    )
    order = [(s.text_frame.text, s.is_placeholder) for s in prs.slides[0].shapes]
    assert order == [("a", False), ("b", False), ("Title", True), ("Body", True), ("c", False)]


def test_shapes_work_on_any_layout_and_blank_and_title_only_are_the_intended_ones() -> None:
    report, prs, _ = build(
        [
            {"layout": "blank", "shapes": [rect()]},
            {"layout": "title_only", "title": "T", "shapes": [rect(y_cm=6)]},
            {"layout": "section", "title": "S", "shapes": [rect(x_cm=25, y_cm=1, w_cm=3, h_cm=3)]},
        ]
    )
    assert [len(s.shapes) for s in prs.slides] == [1, 2, 2]
    assert not [n for n in report.get("notes", []) if "overlaps" in n]


def test_a_shape_over_a_non_empty_placeholder_warns() -> None:
    report, _, _ = build(
        [{"title": "T", "body": [{"type": "list", "items": ["one", "two"]}], "shapes": [rect(x_cm=3, y_cm=6)]}]
    )
    notes = [n for n in report["notes"] if "overlaps" in n]
    assert len(notes) == 1 and "slides[0].shapes[0]" in notes[0] and "blank or title_only" in notes[0]


def test_an_empty_placeholder_does_not_count_as_overlap() -> None:
    report, _, _ = build([{"layout": "title_content", "title": "T", "shapes": [rect(x_cm=3, y_cm=6)]}])
    assert not [n for n in report.get("notes", []) if "overlaps" in n]


# ---- bounds --------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("size", "shape", "needle"),
    [
        ("16:9", rect(x_cm=30, w_cm=6, text="a"), "33.87 x 19.05 cm (16:9)"),
        ("4:3", rect(x_cm=20, w_cm=6, text="a"), "25.40 x 19.05 cm (4:3)"),
        ("16:9", rect(y_cm=17, h_cm=3, text="a"), "19.05 cm"),
    ],
)
def test_a_shape_outside_the_slide_is_a_422_naming_it_and_the_slide_size(
    size: str, shape: dict[str, Any], needle: str
) -> None:
    with pytest.raises(InvalidRequest) as ei:
        build([{"layout": "blank"}, {"layout": "blank", "shapes": [rect(), shape]}], size=size)
    assert ei.value.status == 422
    assert "slides[1].shapes[1] (rect)" in str(ei.value) and needle in str(ei.value)


def test_a_shape_that_fits_exactly_is_accepted() -> None:
    _, prs, _ = build(
        [
            {
                "layout": "blank",
                "shapes": [{"kind": "rect", "x_cm": 0, "y_cm": 0, "w_cm": 33.867, "h_cm": 19.05, "text": "a"}],
            }
        ]
    )
    s = prs.slides[0].shapes[0]
    assert s.left + s.width <= prs.slide_width and s.top + s.height <= prs.slide_height


def test_a_shape_without_text_may_bleed_off_every_edge() -> None:
    bleeds = [
        {"kind": "ellipse", "x_cm": 28, "y_cm": -2, "w_cm": 10, "h_cm": 10},  # top right corner
        {"kind": "rect", "x_cm": -3, "y_cm": 15, "w_cm": 8, "h_cm": 8},  # bottom left corner
        {"kind": "line", "x_cm": -5, "y_cm": 4, "w_cm": 50, "h_cm": 0},  # across both side edges
        {"kind": "arrow", "x_cm": 10, "y_cm": -1, "w_cm": 0, "h_cm": 8},
    ]
    _, prs, _ = build([{"layout": "blank", "shapes": bleeds}])
    made = list(prs.slides[0].shapes)
    assert len(made) == 4
    top_right = next(m for m in made if m.name.startswith("ellipse"))
    assert top_right.top == -2 * CM and top_right.left + top_right.width > prs.slide_width


@pytest.mark.parametrize(
    "shape",
    [
        {"kind": "rect", "x_cm": 40, "y_cm": 1, "w_cm": 5, "h_cm": 5},  # entirely right of the slide
        {"kind": "rect", "x_cm": -10, "y_cm": 1, "w_cm": 5, "h_cm": 5},  # entirely left
        {"kind": "ellipse", "x_cm": 1, "y_cm": -8, "w_cm": 5, "h_cm": 5},  # entirely above
        {"kind": "rect", "x_cm": -5, "y_cm": 1, "w_cm": 5.3, "h_cm": 5},  # 0.3 cm showing
    ],
)
def test_a_bleeding_shape_must_still_show_half_a_centimetre(shape: dict[str, Any]) -> None:
    with pytest.raises(InvalidRequest) as ei:
        build([{"layout": "blank", "shapes": [rect(), shape]}])
    assert "slides[0].shapes[1]" in str(ei.value) and "lies off the slide" in str(ei.value)
    assert "33.87 x 19.05 cm (16:9)" in str(ei.value)


@pytest.mark.parametrize(
    "shape",
    [
        rect(x_cm=-1, w_cm=5, text="Title"),
        rect(y_cm=-2, w_cm=5, runs=[{"text": "x"}]),
        {"kind": "text", "x_cm": 1, "y_cm": -1, "w_cm": 5, "h_cm": 2, "text": "hi"},
        {"kind": "text", "x_cm": 30, "y_cm": 1, "w_cm": 8, "h_cm": 2, "text": "hi"},
    ],
)
def test_a_shape_with_text_never_bleeds(shape: dict[str, Any]) -> None:
    with pytest.raises(InvalidRequest) as ei:
        build([{"layout": "blank", "shapes": [shape]}])
    msg = str(ei.value)
    assert "slides[0].shapes[0]" in msg and "reaches outside the slide" in msg and "33.87 x 19.05 cm (16:9)" in msg


def test_a_position_below_minus_fifty_is_a_range_error_naming_the_range_without_the_parent_keys() -> None:
    with pytest.raises(InvalidRequest) as ei:
        build([{"layout": "blank", "shapes": [rect(y_cm=-60)]}])
    msg = str(ei.value)
    assert "y_cm: -60 is less than the minimum (-50); allowed range -50 to 100" in msg
    assert "properties" not in ei.value.details and "required" not in ei.value.details


def test_the_shape_schema_rules() -> None:
    def refused(shape: dict[str, Any], needle: str) -> None:
        with pytest.raises(InvalidRequest, match=needle):
            build([{"layout": "blank", "shapes": [shape]}])

    refused({"kind": "hexagon", "x_cm": 1, "y_cm": 1, "w_cm": 1, "h_cm": 1}, "kind")
    refused(rect(size_pt=9, text="x"), "size_pt")
    refused(rect(w_cm=0), "w_cm and h_cm above 0")
    refused({"kind": "line", "x_cm": 1, "y_cm": 1, "w_cm": 0, "h_cm": 0}, "needs w_cm or h_cm")
    refused({"kind": "arrow", "x_cm": 1, "y_cm": 1, "w_cm": 2, "h_cm": 0, "text": "x"}, "takes `line` only")
    refused(rect(kind="text", fill="FF0000"), "no fill")
    refused(rect(radius=1), "rounded_rect only")
    refused(rect(text="a", runs=[{"text": "b"}]), "one of")
    refused(rect(fill="blue"), "fill")
    with pytest.raises(InvalidRequest, match="shapes"):
        build([{"layout": "blank", "shapes": [rect()] * 31}])


def test_text_that_cannot_fit_its_shape_is_a_warning_not_an_error() -> None:
    report, _, _ = build([{"layout": "blank", "shapes": [rect(w_cm=3, h_cm=1.5, text="word " * 40, size_pt=18)]}])
    assert any("slides[0].shapes[0]" in n and "will not fit" in n for n in report["notes"])
    ok, _, _ = build([{"layout": "blank", "shapes": [rect(w_cm=8, h_cm=4, text="Short", size_pt=18)]}])
    assert not [n for n in ok.get("notes", []) if "will not fit" in n]


# ---- the deck as a whole -------------------------------------------------------------------------


def showcase() -> list[dict[str, Any]]:
    cards = [
        {
            "kind": "rounded_rect",
            "x_cm": 1.6 + i * 10.4,
            "y_cm": 5,
            "w_cm": 9.6,
            "h_cm": 9,
            "fill": fill,
            "radius": 0.8,
            "runs": [{"text": f"Step {i + 1}\n", "bold": True}, {"text": "A short sentence about the step"}],
            "size_pt": 20,
        }
        for i, fill in enumerate(("2F5597", "4C9F70", "E4893B"))
    ]
    return [
        {
            "layout": "title",
            "title": "Gradient cover",
            "subtitle": "with a native fill",
            "background": {"gradient": ["10243E", "2F5597"], "angle": 90},
            "shapes": [{"kind": "rect", "x_cm": 0, "y_cm": 17.5, "w_cm": 33.8, "h_cm": 1.5, "fill": "E4893B"}],
        },
        {"layout": "title_only", "title": "Three cards", "shapes": cards},
        {
            "layout": "blank",
            "background": {"image": "wide.png", "fit": "cover"},
            "shapes": [
                {
                    "kind": "text",
                    "x_cm": 4,
                    "y_cm": 7,
                    "w_cm": 26,
                    "h_cm": 4,
                    "text": "Section",
                    "size_pt": 44,
                    "align": "center",
                },
                {
                    "kind": "arrow",
                    "x_cm": 8,
                    "y_cm": 13,
                    "w_cm": 18,
                    "h_cm": 0,
                    "line": {"color": "FFFFFF", "width_pt": 4},
                },
            ],
        },
    ]


def test_the_showcase_deck_reopens_and_every_shape_stays_on_the_slide() -> None:
    report, prs, _data = build(showcase(), theme={"colors": {"text": "FFFFFF", "heading": "FFFFFF"}})
    assert report["verify"]["reopened_with"] == "python-pptx"
    assert len(prs.slides) == 3
    for slide in prs.slides:
        for sh in slide.shapes:
            assert sh.left >= 0 and sh.top >= 0
            assert sh.left + sh.width <= prs.slide_width and sh.top + sh.height <= prs.slide_height
    etree.fromstring(prs.slides[2].part.blob)  # the background part parses
    again = create.create({"format": "pptx", "spec": {"slides": showcase()}}, None, IMG, Limits())[1]
    assert again == build(showcase(), theme=None)[2]  # deterministic, theme left aside


def test_backgrounds_and_shapes_work_on_a_template_deck_too() -> None:
    template = style_template_pptx()
    report, prs, _ = build(
        [
            {
                "layout": "Brand Content",
                "title": "T",
                "background": {"color": "EEEEEE"},
                "shapes": [rect(x_cm=15, y_cm=12, text="a")],
            }
        ],
        theme={"template": "brand.pptx"},
        template=template,
    )
    slide = prs.slides[0]
    assert bg(slide).xpath("p:bgPr/a:solidFill/a:srgbClr/@val") == ["EEEEEE"]
    assert slide.shapes[0].is_placeholder is False  # the shape is under the template's title
    assert report["styled_from"] == "template"
    with pytest.raises(InvalidRequest, match="shapes\\[0\\]"):  # bounds use the template's own slide size
        w = prs.slide_width / CM
        build(
            [{"layout": "Brand Content", "shapes": [rect(x_cm=w - 1, text="a")]}],
            theme={"template": "brand.pptx"},
            template=template,
        )


def test_the_new_properties_reach_pptx_only() -> None:
    schemas = spec.create_schemas()
    assert {"background", "shapes"} <= set(schemas["pptx"]["slide"]["properties"])
    assert "slide_background" in str(schemas["pptx"]["document"]["properties"]["theme"])
    for fmt in ("docx", "pdf", "xlsx"):
        assert "slide_background" not in str(schemas[fmt]) and "shapes" not in str(schemas[fmt]["document"])
    assert "slide_background" not in str(spec.THEME)


def test_the_outline_lists_the_shapes_of_a_created_deck() -> None:
    report, _, _ = build([{"layout": "blank", "shapes": [rect(), rect(kind="ellipse", x_cm=12)]}])
    assert report["outline"]["slides"][0]["shapes"]


@pytest.mark.skipif(
    not os.environ.get("GOTENBERG_URL"), reason="needs a Gotenberg (LibreOffice) service; CI provides one"
)
def test_a_deck_with_backgrounds_and_shapes_converts_to_a_pdf_page_per_slide() -> None:
    import pypdf

    slides = [*showcase(), {"title": "Plain", "body": [{"type": "paragraph", "text": "after the designed ones"}]}]
    _, data = create.create({"format": "pptx", "spec": {"slides": slides}}, None, IMG, Limits())
    assert len(pypdf.PdfReader(io.BytesIO(to_pdf(data))).pages) == len(slides)
