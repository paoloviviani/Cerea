"""Create a .pptx from a spec with python-pptx.

The built-in set is PowerPoint's default master and layouts with geometry rebuilt for the slide size
(16:9 or 4:3) and restyled through the master (fonts, colours, title style). With a template the uploaded file's master and layouts are used as
they are; `layout` then names one of its layouts. Native charts are python-pptx charts."""

from __future__ import annotations

import io
import math
import re
from typing import Any

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Emu, Pt

from . import charts
from .determinism import FIXED, fix_zip
from .errors import InvalidRequest
from .images_ref import ImageSet
from .theme import DEFAULT_COLORS, Theme

EMU_CM = 360_000
MARGIN_CM = 1.6
TITLE_PT = 36  # title style on the master; long titles step down through TITLE_STEPS
TITLE_STEPS = (36, 32, 28, 24)
BODY_LEVEL_PT = (24, 22, 20, 18, 18)  # the master's body style
BODY_STEPS = (24, 22, 20, 18, 16, 14)  # sizes a body may take; below the last it is too dense for a slide
SMALL_PT = 16  # below this a note suggests splitting the slide
TABLE_STEPS = (18, 16, 14, 12, 11, 10)
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
BUILTIN = {"title": 0, "title_content": 1, "section": 2, "two_content": 3, "title_only": 5, "blank": 6}
TABLE_STYLE_IDS = {
    "banded": "{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}",
    "grid": "{5940675A-B579-460E-94D1-54222C63F5DA}",
    "minimal": "{2D5ABB26-0587-4C30-8999-92F81FD0307C}",
    "plain": "{2D5ABB26-0587-4C30-8999-92F81FD0307C}",
}
TEXT_KINDS = ("paragraph", "list", "quote", "code")
Box = tuple[int, int, int, int]
_PPTX_MAIN = b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"


def _as_plain_pptx(data: bytes) -> bytes:
    """python-pptx opens a plain .pptx only; present a template (.potx) or show (.ppsx) to it as one."""
    import zipfile

    with zipfile.ZipFile(io.BytesIO(data)) as z:
        ct = z.read("[Content_Types].xml")
    fixed = re.sub(
        rb"application/vnd\.openxmlformats-officedocument\.presentationml\.(?:template|slideshow)\.main\+xml",
        _PPTX_MAIN,
        ct,
    )
    if fixed == ct:
        return data
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as zi, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for i in zi.infolist():
            zo.writestr(i.filename, fixed if i.filename == "[Content_Types].xml" else zi.read(i.filename))
    return out.getvalue()


SHAPE_TYPES = {
    "rect": MSO_SHAPE.RECTANGLE,
    "rounded_rect": MSO_SHAPE.ROUNDED_RECTANGLE,
    "ellipse": MSO_SHAPE.OVAL,
    "chevron": MSO_SHAPE.CHEVRON,
}
SHAPE_TEXT_PT = 18
MIN_CONTRAST = 4.5
EDGE_CM = 0.01  # tolerance when a shape touches the slide edge
OVERLAP_CM = 0.2  # a shape that covers less than this of a placeholder (either way) is not an overlap


def _luminance(h: str) -> float:
    def lin(v: int) -> float:
        c = v / 255
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4

    return 0.2126 * lin(int(h[0:2], 16)) + 0.7152 * lin(int(h[2:4], 16)) + 0.0722 * lin(int(h[4:6], 16))


def contrast(a: str, b: str) -> float:
    """WCAG contrast ratio of two hex colours."""
    hi, lo = sorted((_luminance(a), _luminance(b)), reverse=True)
    return (hi + 0.05) / (lo + 0.05)


def _rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.lstrip("#").upper())


class PptxBuilder:
    def __init__(self, spec: dict[str, Any], images: ImageSet, theme: Theme, template: bytes | None) -> None:
        self.spec, self.images, self.t = spec, images, theme
        self.notes: list[str] = []
        self._warned_contrast: set[str] = set()
        self.charts: list[dict[str, Any]] = []
        self.has_template = template is not None
        if template:
            self.prs = Presentation(io.BytesIO(_as_plain_pptx(template)))
            lst = self.prs.slides._sldIdLst
            for sld in list(lst):  # the template's own slides are not content
                self.prs.part.drop_rel(sld.rId)
                lst.remove(sld)
        else:
            self.prs = Presentation()
            self._lay_out_base(12_192_000 if spec.get("size", "16:9") == "16:9" else 9_144_000)
            self._restyle()
        self.W, self.H = int(self.prs.slide_width), int(self.prs.slide_height)

    # ---- the built-in set ---------------------------------------------------------------------

    def _lay_out_base(self, width: int) -> None:
        """Give the default deck's master and every layout geometry for the chosen slide size.

        python-pptx's default deck is 4:3 and most of its layout placeholders have no geometry of their own
        (they inherit the master's). Rescaling them in place wrote the inherited numbers back twice scaled and
        with a zero height, so the geometry is set here from the slide size instead: one margin, a title band,
        a content area, a footer row, and the same figures in every layout."""
        self.prs.slide_width = Emu(width)
        m = int(EMU_CM * MARGIN_CM)
        gap = int(EMU_CM)
        inner = width - 2 * m
        title: Box = (m, int(EMU_CM * 0.9), inner, int(EMU_CM * 3.0))
        body: Box = (m, int(EMU_CM * 4.4), inner, int(EMU_CM * 12.2))
        col = (inner - gap) // 2
        foot_y, foot_h = int(EMU_CM * 17.4), int(EMU_CM * 1.0)
        foot_w = int(EMU_CM * 7.5)
        footers: dict[PP_PLACEHOLDER, Box] = {
            PP_PLACEHOLDER.DATE: (m, foot_y, foot_w, foot_h),
            PP_PLACEHOLDER.FOOTER: ((width - int(EMU_CM * 11)) // 2, foot_y, int(EMU_CM * 11), foot_h),
            PP_PLACEHOLDER.SLIDE_NUMBER: (width - m - foot_w, foot_y, foot_w, foot_h),
        }
        wide = int(EMU_CM * 3)
        centred_w = width - 2 * wide
        cap_w = int(EMU_CM * 10.5)
        pic_w = min(int(EMU_CM * 20), inner)
        pic_x = (width - pic_w) // 2
        by_layout: dict[str, list[Box | None]] = {  # None: inherit from the master
            "Title Slide": [
                (wide, int(EMU_CM * 5.2), centred_w, int(EMU_CM * 4.4)),
                (wide + int(EMU_CM * 2), int(EMU_CM * 10.2), centred_w - int(EMU_CM * 4), int(EMU_CM * 3.4)),
            ],
            "Title and Content": [None, None],
            "Section Header": [
                (wide, int(EMU_CM * 5.6), centred_w, int(EMU_CM * 4.4)),
                (wide, int(EMU_CM * 10.2), centred_w, int(EMU_CM * 3.0)),
            ],
            "Two Content": [None, (m, body[1], col, body[3]), (m + col + gap, body[1], col, body[3])],
            "Comparison": [
                None,
                (m, body[1], col, int(EMU_CM * 1.8)),
                (m, body[1] + int(EMU_CM * 1.9), col, body[3] - int(EMU_CM * 1.9)),
                (m + col + gap, body[1], col, int(EMU_CM * 1.8)),
                (m + col + gap, body[1] + int(EMU_CM * 1.9), col, body[3] - int(EMU_CM * 1.9)),
            ],
            "Title Only": [None],
            "Blank": [],
            "Content with Caption": [
                (m, int(EMU_CM * 1.6), cap_w, int(EMU_CM * 3.4)),
                (m + cap_w + gap, int(EMU_CM * 1.6), inner - cap_w - gap, int(EMU_CM * 14.8)),
                (m, int(EMU_CM * 5.2), cap_w, int(EMU_CM * 11.2)),
            ],
            "Picture with Caption": [
                (pic_x, int(EMU_CM * 12.6), pic_w, int(EMU_CM * 1.7)),
                (pic_x, int(EMU_CM * 1.6), pic_w, int(EMU_CM * 10.8)),
                (pic_x, int(EMU_CM * 14.4), pic_w, int(EMU_CM * 2.2)),
            ],
            "Title and Vertical Text": [None, None],
            "Vertical Title and Text": [
                (width - m - int(EMU_CM * 7.5), int(EMU_CM * 1.6), int(EMU_CM * 7.5), int(EMU_CM * 15.0)),
                (m, int(EMU_CM * 1.6), inner - int(EMU_CM * 8.5), int(EMU_CM * 15.0)),
            ],
        }
        master: Any = self.prs.slide_master
        for ph in master.placeholders:
            kind = ph.placeholder_format.type
            box = footers.get(kind) or (title if kind == PP_PLACEHOLDER.TITLE else body)
            self._set_box(ph, box)
        layouts: list[Any] = list(self.prs.slide_layouts)
        for lay in layouts:
            boxes = list(by_layout.get(lay.name, []))
            for ph in lay.placeholders:
                kind = ph.placeholder_format.type
                if kind in footers:
                    self._set_box(ph, None)  # the master's footer row
                elif boxes:
                    self._set_box(ph, boxes.pop(0))
                else:
                    self._set_box(ph, None)

    @staticmethod
    def _set_box(ph: Any, box: Box | None) -> None:
        sp_pr = ph._element.spPr
        for xfrm in sp_pr.findall(f"{{{A}}}xfrm"):
            sp_pr.remove(xfrm)
        if box is not None:
            ph.left, ph.top, ph.width, ph.height = (Emu(v) for v in box)

    def _restyle(self) -> None:
        t = self.t
        master = self.prs.slide_master
        theme = master.part.part_related_by(RT.THEME)
        root = etree.fromstring(theme.blob)
        ns = {"a": A}
        for tag, font in (("majorFont", t.heading_font), ("minorFont", t.body_font)):
            for el in root.xpath(f"//a:fontScheme/a:{tag}/a:latin", namespaces=ns):
                el.set("typeface", font)
        for el in root.xpath("//a:clrScheme/a:accent1/a:srgbClr", namespaces=ns):
            el.set("val", t.colors["accent"])
        for i, colour in enumerate(t.palette[1:6], start=2):
            for el in root.xpath(f"//a:clrScheme/a:accent{i}/a:srgbClr", namespaces=ns):
                el.set("val", colour)
        theme._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        styles = master._element.find(f"{{{P}}}txStyles")
        if styles is None:
            return
        title = styles.find(f"{{{P}}}titleStyle/{{{A}}}lvl1pPr")
        if title is not None:
            title.set("algn", "l")
            rpr = title.find(f"{{{A}}}defRPr")
            if rpr is not None:
                rpr.set("sz", str(TITLE_PT * 100))
                rpr.set("b", "1")
                for fill in rpr.findall(f"{{{A}}}solidFill"):
                    rpr.remove(fill)
                fill = etree.SubElement(rpr, f"{{{A}}}solidFill")
                etree.SubElement(fill, f"{{{A}}}srgbClr").set("val", t.colors["heading"])
                rpr.insert(0, fill)
        for lvl, size in zip(range(1, 6), BODY_LEVEL_PT, strict=True):
            el = styles.find(f"{{{P}}}bodyStyle/{{{A}}}lvl{lvl}pPr/{{{A}}}defRPr")
            if el is not None:
                el.set("sz", str(int(size * 100)))
        self.prs.core_properties  # noqa: B018 - touches the part so it exists before properties are set

    # ---- layouts ------------------------------------------------------------------------------

    def layouts(self) -> dict[str, Any]:
        return {lay.name.lower(): lay for lay in self.prs.slide_layouts}

    def pick(self, name: str, where: str) -> Any:
        if not self.has_template:
            if name in BUILTIN:
                return self.prs.slide_layouts[BUILTIN[name]]
            raise InvalidRequest(
                f"{where}: unknown layout {name!r}; available: {sorted(BUILTIN)}. "
                "Named layouts need `theme.template` (an uploaded presentation).",
                reason="schema",
                details={"available": sorted(BUILTIN)},
            )
        named = self.layouts()
        if name.lower() in named:
            return named[name.lower()]
        if name in BUILTIN:
            found = self._by_shape(name)
            if found is not None:
                return found
        raise InvalidRequest(
            f"{where}: the template has no layout {name!r}; its layouts: {[lay.name for lay in self.prs.slide_layouts]}",
            reason="schema",
            details={"available": [lay.name for lay in self.prs.slide_layouts]},
        )

    def _by_shape(self, kind: str) -> Any | None:
        def types(lay: Any) -> list[Any]:
            return [p.placeholder_format.type for p in lay.placeholders]

        body = (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT)
        for lay in self.prs.slide_layouts:
            ts = types(lay)
            has_title = PP_PLACEHOLDER.TITLE in ts or PP_PLACEHOLDER.CENTER_TITLE in ts
            n_body = sum(x in body for x in ts)
            if kind == "title" and PP_PLACEHOLDER.CENTER_TITLE in ts:
                return lay
            if kind == "section" and has_title and "section" in lay.name.lower():
                return lay
            if kind == "title_content" and has_title and n_body == 1 and PP_PLACEHOLDER.CENTER_TITLE not in ts:
                return lay
            if kind == "two_content" and has_title and n_body >= 2:
                return lay
            if kind == "title_only" and has_title and n_body == 0 and PP_PLACEHOLDER.SUBTITLE not in ts:
                return lay
            if kind == "blank" and not ts:
                return lay
        return None

    # ---- build --------------------------------------------------------------------------------

    def build(self) -> tuple[bytes, dict[str, Any]]:
        self._check_shape_bounds()
        for i, s in enumerate(self.spec["slides"]):
            self.slide(i, s)
        cp = self.prs.core_properties
        cp.author = cp.last_modified_by = "Cerea"
        cp.title = str(self.spec.get("title") or "")
        cp.created = cp.modified = FIXED
        cp.revision = 1
        if self.spec.get("language"):
            cp.language = self.spec["language"]
        out = io.BytesIO()
        self.prs.save(out)
        return fix_zip(out.getvalue()), {"charts": self.charts, "notes": self.notes, "slides": len(self.spec["slides"])}

    def slide(self, i: int, s: dict[str, Any]) -> None:
        where = f"slides[{i}]"
        visuals_only = bool(s.get("images") or s.get("charts"))
        default = (
            "title" if (i == 0 or s.get("subtitle")) and not s.get("body") and not visuals_only
            else "two_content" if s.get("body_right")
            else "title_content" if s.get("body") or visuals_only
            else "title_only"
        )  # fmt: skip
        layout = self.pick(s.get("layout") or default, where)
        slide = self.prs.slides.add_slide(layout)
        self._background(slide, s.get("background"), self.t.slide_background, where)
        mine = self._shapes(slide, s.get("shapes") or [], "back", where)
        phs = list(slide.placeholders)
        title_ph = next(
            (p for p in phs if p.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)), None
        )
        sub_ph = next((p for p in phs if p.placeholder_format.type == PP_PLACEHOLDER.SUBTITLE), None)
        bodies = [p for p in phs if p.placeholder_format.type in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT)]
        pictures = [p for p in phs if p.placeholder_format.type == PP_PLACEHOLDER.PICTURE]
        centered = (s.get("layout") or default) in ("title", "section") and not self.has_template
        if s.get("title"):
            if title_ph is None:
                self.notes.append(f"{where}: layout {layout.name!r} has no title placeholder; the title was not shown")
            else:
                self._text(
                    title_ph,
                    [{"text": s["title"]}],
                    0,
                    size=self._title_size(s["title"], title_ph, 44 if centered else TITLE_PT),
                    bold=True if centered else None,
                    align=PP_ALIGN.CENTER if centered else None,
                    heading=True,
                )
        if s.get("subtitle"):
            target = sub_ph or (bodies.pop(0) if bodies and layout.name.lower().startswith("section") else None)
            if target is None:
                self.notes.append(
                    f"{where}: layout {layout.name!r} has no subtitle placeholder; the subtitle was not shown"
                )
            else:
                self._text(
                    target,
                    [{"text": s["subtitle"]}],
                    0,
                    size=22 if not self.has_template else None,
                    align=PP_ALIGN.CENTER if centered else None,
                )
        left, right = [*bodies, None, None][:2]
        area = self._area(slide, title_ph, left, right if s.get("body_right") is None else None)
        left_area, right_area = area, area
        if s.get("body_right"):
            if left is not None and right is not None and left.width and right.width:
                left_area, right_area = self._box(left), self._box(right)  # the layout's own columns
            else:
                left_area, right_area = self._half(area, 0), self._half(area, 1)
        self.fill(slide, s.get("body") or [], left, left_area, s, where, extra=True)
        if s.get("body_right"):
            self.fill(slide, s["body_right"], right, right_area, s, where + ".body_right", extra=False)
        for ph in pictures:
            self._remove(ph)
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]
        for ph in list(slide.placeholders):  # nothing may be left as "Click to add text"
            if ph.has_text_frame and not ph.text_frame.text.strip() and ph.element.getparent() is not None:
                self._remove(ph)
        content = [sh for sh in slide.shapes if sh.shape_id not in {m.shape_id for m in mine}]
        mine += self._shapes(slide, s.get("shapes") or [], "front", where)
        self._warn_overlaps(s.get("shapes") or [], mine, content, where)

    # ---- backgrounds and free shapes -----------------------------------------------------------

    def _slide_cm(self) -> tuple[float, float]:
        return self.W / EMU_CM, self.H / EMU_CM

    def _check_shape_bounds(self) -> None:
        """Every shape lies on the slide (all of them, before anything is built): a 422 names the shape and the size."""
        sw, sh = self._slide_cm()
        for i, s in enumerate(self.spec["slides"]):
            for j, shape in enumerate(s.get("shapes") or []):
                x, y, w, h = (float(shape[k]) for k in ("x_cm", "y_cm", "w_cm", "h_cm"))
                if x + w > sw + EDGE_CM or y + h > sh + EDGE_CM:
                    size = (
                        "16:9" if abs(sw / sh - 16 / 9) < 0.01 else "4:3" if abs(sw / sh - 4 / 3) < 0.01 else "custom"
                    )
                    over = [
                        f"{ax} {p:g} + {dim} {d:g} = {p + d:.2f} cm, past the slide's {lim:.2f} cm"
                        for ax, dim, p, d, lim in (("x", "w", x, w, sw), ("y", "h", y, h, sh))
                        if p + d > lim + EDGE_CM
                    ]
                    raise InvalidRequest(
                        f"slides[{i}].shapes[{j}] ({shape['kind']}) reaches outside the slide: "
                        + " and ".join(over)
                        + f", but the slide is {sw:.2f} x {sh:.2f} cm ({size}). Move it or make it smaller.",
                        reason="schema",
                        details={"slide_cm": [round(sw, 2), round(sh, 2)], "slide": i, "shape": j},
                    )

    def _background(self, slide: Any, own: dict[str, Any] | None, default: dict[str, Any] | None, where: str) -> None:
        bg = own or default
        if not bg:
            return
        label = f"{where}.background" if own else "theme.slide_background"
        colors: list[str]
        if "color" in bg:
            colors = [bg["color"]]
            fill = f'<a:solidFill xmlns:a="{A}"><a:srgbClr val="{_hex(bg["color"])}"/></a:solidFill>'
        elif "gradient" in bg:
            colors = list(bg["gradient"])
            stops = "".join(
                f'<a:gs pos="{pos}"><a:srgbClr val="{_hex(c)}"/></a:gs>'
                for pos, c in zip((0, 100_000), colors, strict=True)
            )
            ang = int(float(bg.get("angle", 90)) * 60_000) % 21_600_000
            fill = f'<a:gradFill xmlns:a="{A}" rotWithShape="1"><a:gsLst>{stops}</a:gsLst><a:lin ang="{ang}" scaled="0"/></a:gradFill>'
        else:
            colors = []
            img = self.images.get(bg["image"])
            _part, rid = slide.part.get_or_add_image_part(io.BytesIO(_cover(img.data, img.info.ext, self.W, self.H)))
            fill = (
                f'<a:blipFill xmlns:a="{A}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
                f'dpi="0" rotWithShape="1"><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></a:blipFill>'
            )
        c_sld = slide._element.find(f"{{{P}}}cSld")
        for old in c_sld.findall(f"{{{P}}}bg"):
            c_sld.remove(old)
        node = etree.fromstring(f'<p:bg xmlns:p="{P}"><p:bgPr>{fill}<a:effectLst xmlns:a="{A}"/></p:bgPr></p:bg>')
        c_sld.insert(0, node)
        self._contrast_note(colors, label)

    def _contrast_note(self, colors: list[str], label: str) -> None:
        """Text colour is never changed on its own: say when the text colours cannot be read on a solid background."""
        text = {k: self.t.colors[k] for k in ("text", "heading") if self.t.uses(f"color_{k}")}
        for c in colors:
            for name, fg in text.items():
                ratio = contrast(_hex(fg), _hex(c))
                key = f"{label}|{c}|{name}"
                if ratio < MIN_CONTRAST and key not in self._warned_contrast:
                    self._warned_contrast.add(key)
                    self.notes.append(
                        f"{label}: the {name} colour {_hex(fg)} has a contrast of {ratio:.1f}:1 with {_hex(c)} "
                        f"(under {MIN_CONTRAST}:1); set theme.colors.{name} to a lighter or darker colour, "
                        "or colour the text runs"
                    )

    def _shapes(self, slide: Any, shapes: list[dict[str, Any]], layer: str, where: str) -> list[Any]:
        """Draw the shapes of one layer in list order: back ones go under the layout's placeholders, front ones on top."""
        tree = slide.shapes._spTree
        made: list[Any] = []
        first = 2  # after nvGrpSpPr and grpSpPr
        for j, spec in enumerate(shapes):
            if (spec.get("layer") or "back") != layer:
                continue
            shape = self._shape(slide, spec, f"{where}.shapes[{j}]")
            if layer == "back":
                tree.remove(shape._element)
                tree.insert(first, shape._element)
                first += 1
            made.append(shape)
        return made

    def _shape(self, slide: Any, spec: dict[str, Any], where: str) -> Any:
        kind = spec["kind"]
        x, y, w, h = (round(float(spec[k]) * EMU_CM) for k in ("x_cm", "y_cm", "w_cm", "h_cm"))
        x, y = min(x, self.W), min(y, self.H)
        w, h = min(w, self.W - x), min(h, self.H - y)  # rounding must not push the edge past the slide
        style = spec.get("line") or {}
        if kind in ("line", "arrow"):
            shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x, y, x + w, y + h)
            ln = shape.line
            ln.color.rgb = _rgb(style.get("color") or self.t.colors["text"])
            ln.width = Pt(style.get("width_pt", 2))
            if kind == "arrow":
                etree.SubElement(ln._get_or_add_ln(), f"{{{A}}}tailEnd").set("type", "triangle")
            self._drop_style(shape)
            return shape
        if kind == "text":
            shape = slide.shapes.add_textbox(x, y, w, h)
            fill: str | None = None
        else:
            shape = slide.shapes.add_shape(SHAPE_TYPES[kind], x, y, w, h)
            given = spec.get("fill")
            fill = None if given == "none" else _hex(given or self.t.colors["accent"])
            if fill:
                shape.fill.solid()
                shape.fill.fore_color.rgb = _rgb(fill)
            else:
                shape.fill.background()
            if kind == "rounded_rect":
                radius = float(spec.get("radius", 0.5)) * EMU_CM
                shape.adjustments[0] = min(0.5, radius / max(1, min(w, h)))
        if style:
            shape.line.color.rgb = _rgb(style.get("color") or self.t.colors["text"])
            shape.line.width = Pt(style.get("width_pt", 1))
        else:
            shape.line.fill.background()
        self._drop_style(shape)
        shape.name = f"{kind} {where.rsplit('[', 1)[1].rstrip(']')}"
        if "text" in spec or "runs" in spec:
            self._shape_text(shape, spec, fill, w, h, where)
        return shape

    @staticmethod
    def _drop_style(shape: Any) -> None:
        """The theme's own shape style adds shadows and its line; every property is set explicitly instead."""
        for st in shape._element.findall(f"{{{P}}}style"):
            shape._element.remove(st)

    def _shape_text(self, shape: Any, spec: dict[str, Any], fill: str | None, w: int, h: int, where: str) -> None:
        size = float(spec.get("size_pt", SHAPE_TEXT_PT))
        runs = spec.get("runs") or [{"text": spec.get("text", "")}]
        paras: list[list[dict[str, Any]]] = [[]]
        for r in runs:  # a line break in the text starts a new paragraph
            for n, part in enumerate(str(r["text"]).split("\n")):
                if n:
                    paras.append([])
                if part:
                    paras[-1].append({**r, "text": part})
        tf = shape.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[
            spec.get("valign") or ("middle" if spec["kind"] != "text" else "top")
        ]
        align = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[
            spec.get("align") or ("center" if spec["kind"] != "text" else "left")
        ]
        default_color = self._text_on(fill)
        for n, runs_in in enumerate(paras):
            p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
            p.alignment = align
            for r in runs_in:
                self._run(p, r, size)
                if not r.get("color") and default_color:
                    p.runs[-1].font.color.rgb = _rgb(default_color)
        est = self._text_height([(r, 0, "paragraph") for r in paras], w / 12700 + 22, math.ceil(size))
        if est > h / 12700:
            self.notes.append(
                f"{where}: the text will not fit in a {w / EMU_CM:.1f} x {h / EMU_CM:.1f} cm shape at {size:g} pt; "
                "make the shape larger, shorten the text or use a smaller size_pt (10 at least)"
            )

    def _text_on(self, fill: str | None) -> str | None:
        """Default text colour inside a filled shape: the theme's text colour, or white when that is hard to read on the fill."""
        if not fill:
            return None
        text = self.t.colors["text"] if self.t.uses("color_text") else DEFAULT_COLORS["text"]
        return (
            text
            if contrast(_hex(text), fill) >= MIN_CONTRAST
            else max(("FFFFFF", "000000"), key=lambda c: contrast(c, fill))
        )

    def _warn_overlaps(self, specs: list[dict[str, Any]], mine: list[Any], content: list[Any], where: str) -> None:
        """A shape over the title or body (or a table, picture, chart) hides or collides with it."""
        for shape in mine:
            x, y, w, h = box_of(shape)
            for other in content:
                if None in (other.left, other.top, other.width, other.height):
                    continue
                ox, oy, ow, oh = box_of(other)
                if (
                    min(x + w, ox + ow) - max(x, ox) > OVERLAP_CM * EMU_CM
                    and min(y + h, oy + oh) - max(y, oy) > OVERLAP_CM * EMU_CM
                ):
                    what = f"the {'placeholder' if other.is_placeholder else 'content'} {other.name!r}"
                    self.notes.append(
                        f"{where}.shapes[{shape.name.rsplit(' ', 1)[1]}] overlaps {what}; "
                        "use the blank or title_only layout for slides built from shapes, or move the shape"
                    )
                    break

    @staticmethod
    def _remove(shape: Any) -> None:
        shape.element.getparent().remove(shape.element)

    @staticmethod
    def _box(ph: Any) -> Box:
        return int(ph.left), int(ph.top), int(ph.width), int(ph.height)

    def _title_size(self, text: str, title_ph: Any, start: int) -> int | None:
        """A size for a title that would not fit its placeholder in two lines at the default; None keeps the style's."""
        if self.has_template:
            return None
        if title_ph is None or not title_ph.width or not title_ph.height:
            return start if start != TITLE_PT else None
        width_pt, height_pt = title_ph.width / 12700, title_ph.height / 12700
        for size in (s for s in TITLE_STEPS if s <= start):
            lines = math.ceil(len(text) * size * 0.55 / max(width_pt, 1))
            if lines * size * 1.2 <= height_pt:
                return size if size != TITLE_PT else None
        return TITLE_STEPS[-1]

    def _area(self, slide: Any, title_ph: Any, body_ph: Any, second_ph: Any = None) -> Box:
        """The content area: the body placeholder's box (joined with the second column's when that is unused)."""
        if body_ph is not None and body_ph.width and body_ph.height:
            x, y, w, h = self._box(body_ph)
            if second_ph is not None and second_ph.width and second_ph.height and second_ph.left > x:
                w = int(second_ph.left + second_ph.width) - x
                h = max(h, int(second_ph.top + second_ph.height) - y)
            return x, y, w, h
        margin = int(EMU_CM * MARGIN_CM)
        title_bottom = (
            int(title_ph.top + title_ph.height)
            if title_ph is not None and title_ph.top is not None
            else int(EMU_CM * 3)
        )
        top = title_bottom + int(EMU_CM * 0.4)
        return margin, top, self.W - 2 * margin, max(self.H - top - int(EMU_CM * 2.4), int(EMU_CM * 4))

    @staticmethod
    def _half(area: Box, side: int) -> Box:
        x, y, w, h = area
        gap = int(EMU_CM)
        return (x + side * (w + gap) // 2, y, (w - gap) // 2, h)

    # ---- content ------------------------------------------------------------------------------

    def fill(
        self, slide: Any, blocks: list[dict[str, Any]], ph: Any, area: Box, s: dict[str, Any], where: str, extra: bool
    ) -> None:
        text = [b for b in blocks if b["type"] in TEXT_KINDS]
        visuals: list[dict[str, Any]] = [b for b in blocks if b["type"] not in TEXT_KINDS]
        if extra:  # explicit images and charts of the slide
            visuals += [{"type": "image", "_slide_image": im} for im in s.get("images") or []]
            visuals += [{"type": "chart", **c} for c in s.get("charts") or []]
        x, y, w, h = area
        text_area, visual_area = area, area
        if text and visuals:
            gap = int(EMU_CM)
            tw = int((w - gap) * 0.42)
            text_area, visual_area = (x, y, tw, h), (x + tw + gap, y, w - tw - gap, h)
        if text:
            target = ph
            if target is None:
                target = slide.shapes.add_textbox(*text_area)
            elif text_area != self._box(target):  # only when the text shares the area with a table, picture or chart
                target.left, target.top, target.width, target.height = text_area
            self._body(target, text, text_area, where)
        elif ph is not None:
            self._remove(ph)
        self._visuals(slide, visuals, visual_area, where)

    def _body(self, shape: Any, blocks: list[dict[str, Any]], area: Box, where: str) -> None:
        paras: list[tuple[list[dict[str, Any]], int, str]] = []  # runs, level, kind
        for b in blocks:
            if b["type"] == "list":
                self._flatten(b["items"], 0, bool(b.get("ordered")), paras)
            elif b["type"] == "code":
                paras.extend(([{"text": line, "code": True}], 0, "code") for line in b["text"].split("\n"))
            else:
                paras.append((b.get("runs") or [{"text": b.get("text", "")}], 0, b["type"]))
        size, fits = self._fit_size(paras, area)
        if not fits:
            self.notes.append(
                f"{where}: the text is too long for one slide even at {size} pt; split it over two slides or shorten it"
            )
        elif size < SMALL_PT:
            self.notes.append(f"{where}: the text was set at {size} pt to fit; consider splitting it over two slides")
        tf = shape.text_frame
        tf.word_wrap = True
        body_pr = tf._txBody.find(f"{{{A}}}bodyPr")
        for old in [e for e in body_pr if e.tag in (f"{{{A}}}noAutofit", f"{{{A}}}spAutoFit", f"{{{A}}}normAutofit")]:
            body_pr.remove(old)
        etree.SubElement(body_pr, f"{{{A}}}normAutofit")  # PowerPoint shrinks further if the text is edited
        for n, (runs, level, kind) in enumerate(paras):
            p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
            if kind in ("paragraph", "quote", "code"):
                self._nobullet(p)
            p.level = min(level, 4)
            for r in runs:
                self._run(p, r, size - 2 * min(level, 2), italic=kind == "quote", mono=kind == "code")
            p.space_after = Pt(max(4, size // 3))

    def _flatten(self, items: list[Any], level: int, ordered: bool, out: list[Any]) -> None:
        for n, it in enumerate(items, 1):
            block = {"text": it} if isinstance(it, str) else it
            runs = list(block.get("runs") or [{"text": block.get("text", "")}])
            if ordered:
                runs = [{"text": f"{n}. "}, *runs]
            out.append((runs, level, "list"))
            if not isinstance(it, str) and it.get("children"):
                self._flatten(it["children"], level + 1, ordered, out)

    def _nobullet(self, p: Any) -> None:
        ppr = p._p.get_or_add_pPr()
        ppr.set("marL", "0")
        ppr.set("indent", "0")
        if ppr.find(f"{{{A}}}buNone") is None:
            etree.SubElement(ppr, f"{{{A}}}buNone")

    def _fit_size(self, paras: list[Any], area: Box) -> tuple[int, bool]:
        """The largest body size (24 down to 14 pt) at which the text is estimated to fit; False if none does."""
        for size in BODY_STEPS:
            if self._text_height(paras, area[2] / 12700, size) <= area[3] / 12700 * 0.92:
                return size, True
        return BODY_STEPS[-1], False

    @staticmethod
    def _text_height(paras: list[Any], width_pt: float, size: int) -> float:
        total = 0.0
        for runs, lvl, _k in paras:
            sz = size - 2 * min(lvl, 2)
            indent = 22 + 18 * lvl  # bullet hanging indent and nesting, in pt
            per_line = max(8, int((width_pt - indent - 14) / (sz * 0.44)))
            lines = max(1, math.ceil(sum(len(r["text"]) for r in runs) / per_line))
            total += lines * sz * 1.2 + max(4, size // 3)
        return total + 14  # the frame's top and bottom insets

    def _run(self, p: Any, r: dict[str, Any], size: float | None, italic: bool = False, mono: bool = False) -> None:
        run = p.add_run()
        run.text = r["text"]
        f = run.font
        if size:
            f.size = Pt(size)
        if r.get("bold"):
            f.bold = True
        if r.get("italic") or italic:
            f.italic = True
        if r.get("underline"):
            f.underline = True
        if r.get("code") or mono:
            f.name = "Consolas"
        elif self.t.uses("body_font") and self.has_template:
            f.name = self.t.body_font
        if r.get("color"):
            f.color.rgb = _rgb(r["color"])
        elif self.t.uses("color_text") and not (r.get("link")):
            f.color.rgb = _rgb(self.t.colors["text"])
        if r.get("link"):
            run.hyperlink.address = r["link"]

    def _text(
        self,
        shape: Any,
        runs: list[dict[str, Any]],
        level: int,
        size: float | None = None,
        bold: bool | None = None,
        align: Any = None,
        heading: bool = False,
    ) -> None:
        tf = shape.text_frame
        p = tf.paragraphs[0]
        if align is not None:
            p.alignment = align
        for r in runs:
            self._run(p, r, size)
            if bold is not None:
                p.runs[-1].font.bold = bold
        if heading:
            for run in p.runs:
                if self.has_template and self.t.uses("heading_font"):
                    run.font.name = self.t.heading_font
                if self.t.uses("color_heading"):
                    run.font.color.rgb = _rgb(self.t.colors["heading"])

    # visuals: tables, images, charts laid out in a grid inside the area
    def _visuals(self, slide: Any, visuals: list[dict[str, Any]], area: Box, where: str) -> None:
        if not visuals:
            return
        placed = [v for v in visuals if self._has_box(v)]
        flow = [v for v in visuals if not self._has_box(v)]
        for v in placed:
            self._place(slide, v, self._clamp(self._box_of(v), where), where)
        if not flow:
            return
        x, y, w, h = area
        cols = 1 if len(flow) == 1 else 2
        rows = math.ceil(len(flow) / cols)
        gap = int(EMU_CM * 0.5)
        cw, ch = (w - gap * (cols - 1)) // cols, (h - gap * (rows - 1)) // rows
        for n, v in enumerate(flow):
            r, c = divmod(n, cols)
            self._place(slide, v, (x + c * (cw + gap), y + r * (ch + gap), cw, ch), where)

    def _clamp(self, box: Box, where: str) -> Box:
        """An explicit box is kept inside the slide (and said so when it had to move or shrink)."""
        x, y, w, h = box
        nw, nh = min(w, self.W), min(h, self.H)
        nx = x if x < 0 else min(max(x, 0), self.W - (nw or self.W // 10))
        ny = y if y < 0 else min(max(y, 0), self.H - (nh or self.H // 10))
        if (nx, ny, nw, nh) != box:
            self.notes.append(f"{where}: an explicit box reached past the slide edge and was fitted inside it")
        return nx, ny, nw, nh

    @staticmethod
    def _has_box(v: dict[str, Any]) -> bool:
        src = v.get("_slide_image") or v
        return not src.get("fit") and any(k in src for k in ("x_cm", "y_cm", "width_cm", "height_cm"))

    def _box_of(self, v: dict[str, Any]) -> Box:
        """An explicit box in EMU; a missing position is -1 (centred for pictures, the content corner otherwise)."""
        src = v.get("_slide_image") or v
        x = int(src["x_cm"] * EMU_CM) if "x_cm" in src else -1
        y = int(src["y_cm"] * EMU_CM) if "y_cm" in src else -1
        w = int(src["width_cm"] * EMU_CM) if "width_cm" in src else 0
        h = int(src["height_cm"] * EMU_CM) if "height_cm" in src else 0
        return x, y, w, h

    def _place(self, slide: Any, v: dict[str, Any], box: Box, where: str) -> None:
        kind = v["type"]
        if kind == "table":
            self._table(slide, v, box)
        elif kind == "chart":
            self._chart(slide, v, box)
        else:
            self._image(slide, v.get("_slide_image") or v, box)

    def _image(self, slide: Any, im: dict[str, Any], box: Box) -> None:
        img = self.images.get(im["image"])
        x, y, w, h = box
        aspect = img.info.width / img.info.height
        if not w and not h:
            w = int(self.W * 0.5)
        if not w:
            w = int(h * aspect)
        if not h:
            h = int(w / aspect)
        w, h = min(w, self.W - max(x, 0)), min(h, self.H - max(y, 0))  # never past the slide edge
        if w / h > aspect:  # fit inside the box without distortion, centred in it
            w2, h2 = int(h * aspect), h
        else:
            w2, h2 = w, int(w / aspect)
        px = (self.W - w2) // 2 if x < 0 else x + (w - w2) // 2
        py = (self.H - h2) // 2 if y < 0 else y + (h - h2) // 2
        pic = slide.shapes.add_picture(io.BytesIO(img.data), px, py, w2, h2)
        pic._element.nvPicPr.cNvPr.set("descr", im.get("alt") or im["image"])

    def _chart(self, slide: Any, c: dict[str, Any], box: Box) -> None:
        x, y, w, h = box
        x, y = (int(EMU_CM * MARGIN_CM) if x < 0 else x), (int(EMU_CM * 4.4) if y < 0 else y)
        w = w or int(self.W * 0.5)
        h = h or min(int(w * 9 / 16), self.H - y - int(EMU_CM * 1.5))
        if c.get("render", "native") == "native":
            try:
                charts.add_pptx_chart(slide.shapes, c, self.t, x, y, w, h)
                self.charts.append(charts.describe(c, "native"))
                return
            except Exception as exc:
                note = f"native chart could not be written ({type(exc).__name__}); placed as an image instead"
                self.notes.append(note)
                self.charts.append(charts.describe(c, "image", note))
        else:
            self.charts.append(charts.describe(c, "image"))
        png = charts.render_png(c, self.t, w / EMU_CM, h / EMU_CM)
        pic = slide.shapes.add_picture(io.BytesIO(png), x, y, w, h)
        pic._element.nvPicPr.cNvPr.set("descr", c.get("title") or f"{c['chart_type']} chart")

    @staticmethod
    def _table_fit(data: list[Any], cols: int, col_w: list[int], height: int) -> tuple[int, list[int], bool]:
        """The largest of 18..10 pt at which the rows (with wrapping) are estimated to fit the height."""
        size, heights = TABLE_STEPS[-1], list[int]()
        for size in TABLE_STEPS:
            heights = []
            for row in data:
                lines = 1
                for ci in range(cols):
                    raw = row[ci] if ci < len(row) else ""
                    text = str(raw.get("text", "") if isinstance(raw, dict) else "" if raw is None else raw)
                    per_line = max(4, int((col_w[ci] / 12700 - 14) / (size * 0.52)))
                    lines = max(lines, math.ceil(len(text) / per_line) if text else 1)
                heights.append(int(Pt(lines * size * 1.2 + 8)))
            if sum(heights) <= height:
                return size, heights, True
        return size, heights, False

    def _table(self, slide: Any, b: dict[str, Any], box: Box) -> None:
        t = self.t
        style = b.get("style") or t.table_style
        header, rows = b.get("header"), b["rows"]
        cols = max([len(header or []), *(len(r) for r in rows)])
        data = ([header] if header else []) + rows
        x, y, w, h = box
        x, y = (int(EMU_CM * MARGIN_CM) if x < 0 else x), (int(EMU_CM * 4.4) if y < 0 else y)
        w = w or self.W - 2 * x
        h = h or max(self.H - y - int(EMU_CM * 2.4), int(EMU_CM * 4))
        widths = b.get("widths") or [1] * cols
        col_w = [int(w * wd / sum(widths)) for wd in widths]
        size, heights, fits = self._table_fit(data, cols, col_w, h)
        if not fits:
            self.notes.append(
                f"slide table with {len(data)} rows is taller than its area even at {size} pt; split it over two slides"
            )
        frame = slide.shapes.add_table(len(data), cols, x, y, w, min(h, sum(heights)) if fits else sum(heights))
        table = frame.table
        sid = frame._element.find(f".//{{{A}}}tableStyleId")
        if sid is not None:
            sid.text = TABLE_STYLE_IDS[style]
        table.first_row = bool(header)
        for i, cw in enumerate(col_w):
            table.columns[i].width = cw
        aligns = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
        for ri, row in enumerate(data):
            is_head = bool(header) and ri == 0
            table.rows[ri].height = heights[ri]
            for ci in range(cols):
                raw = row[ci] if ci < len(row) else ""
                spec = raw if isinstance(raw, dict) else {"text": raw}
                v = spec["text"]
                cell = table.cell(ri, ci)
                cell.text = "" if v is None else ("true" if v is True else "false" if v is False else str(v))
                p = cell.text_frame.paragraphs[0]
                p.alignment = aligns[spec.get("align") or (b.get("align") or ["left"] * cols)[ci]]
                f = p.runs[0].font if p.runs else None
                if f is None:
                    continue
                f.size = Pt(size)
                f.bold = bool(spec.get("bold")) or is_head
                fill = spec.get("fill")
                colour = spec.get("color")
                if is_head and style in ("banded", "grid"):
                    fill, colour = fill or t.colors["table_header_fill"], colour or t.colors["table_header_text"]
                elif style == "banded" and not fill and (ri - (1 if header else 0)) % 2 == 1:
                    fill = t.colors["table_band"]
                elif style == "banded" and not fill:
                    fill = "FFFFFF"
                if fill:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = _rgb(fill)
                if colour:
                    f.color.rgb = _rgb(colour)
                elif not (is_head and style in ("banded", "grid")):
                    f.color.rgb = _rgb(t.colors["text"])


def _hex(value: str) -> str:
    return value.lstrip("#").upper()


def box_of(shape: Any) -> Box:
    return int(shape.left), int(shape.top), int(shape.width), int(shape.height)


def _cover_box(iw: int, ih: int, sw: int, sh: int) -> tuple[int, int, int, int]:
    """The centred (left, top, right, bottom) pixel box of a picture that has the slide's proportions."""
    if iw * sh > ih * sw:  # wider than the slide: the sides go
        width = max(1, round(ih * sw / sh))
        left = (iw - width) // 2
        return left, 0, left + width, ih
    height = max(1, round(iw * sh / sw))
    top = (ih - height) // 2
    return 0, top, iw, top + height


def _cover(data: bytes, ext: str, sw: int, sh: int) -> bytes:
    """The picture cropped to the slide's proportions, so that it fills the slide without being stretched.

    The crop is made in the pixels because LibreOffice ignores the crop rectangle of a slide background and
    would stretch the whole picture; a cropped file looks the same in every application."""
    from PIL import Image

    with Image.open(io.BytesIO(data)) as im:
        box = _cover_box(im.width, im.height, sw, sh)
        if box == (0, 0, im.width, im.height):
            return data
        out = io.BytesIO()
        cropped = im.crop(box)
        if ext == "jpeg":
            cropped.convert("RGB").save(out, "JPEG", quality=95, icc_profile=im.info.get("icc_profile"))
        else:  # PNG, and GIF (a still frame)
            cropped.save(out, "PNG")
        return out.getvalue()


def build_pptx(
    spec: dict[str, Any], images: ImageSet, theme: Theme, template: bytes | None
) -> tuple[bytes, dict[str, Any]]:
    return PptxBuilder(spec, images, theme, template).build()
