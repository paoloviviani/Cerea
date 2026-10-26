"""Check an edit: the result must be a sound package that differs from the source only where the
engine says it edited, and reopen in an independent library."""

from __future__ import annotations

import io
import re
import zipfile
import zlib
from types import ModuleType
from typing import Any

from .config import Limits
from .errors import InvalidDocument, PatchError
from .pkg import Pkg, owner
from .xmlsafe import parse

_XML = (".xml", ".rels")


_WML_MAIN = "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"


def _as_plain_docx(data: bytes) -> bytes:
    """python-docx only opens plain .docx; present a macro-enabled or template file to it as one
    (the check copy is discarded; the delivered bytes are never changed)."""
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        ct = z.read("[Content_Types].xml")
    fixed = re.sub(
        rb'(ContentType=")application/vnd\.(?:ms-word\.document\.macroEnabled\.main\+xml|openxmlformats-officedocument\.wordprocessingml\.template\.main\+xml|ms-word\.template\.macroEnabledTemplate\.main\+xml)"',
        rb"\1" + _WML_MAIN.encode() + b'"',
        ct,
    )
    if fixed == ct:
        return data
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as zi, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for i in zi.infolist():
            zo.writestr(i.filename, fixed if i.filename == "[Content_Types].xml" else zi.read(i.filename))
    return out.getvalue()


def _reopen(fmt: str, data: bytes) -> str:
    try:
        if fmt == "docx":
            import docx

            docx.Document(io.BytesIO(_as_plain_docx(data)))
            return "python-docx"
        if fmt == "pptx":
            import pptx

            pptx.Presentation(io.BytesIO(data))
            return "python-pptx"
        import openpyxl

        openpyxl.load_workbook(io.BytesIO(data), read_only=True).close()
        return "openpyxl"
    except Exception as exc:
        raise PatchError(
            f"the result does not reopen cleanly ({type(exc).__name__}: {exc}); it was not delivered"
        ) from exc


def verify(
    src: Pkg, out_bytes: bytes, limits: Limits, engine: ModuleType, expected: set[str], fmt: str
) -> tuple[dict[str, Any], Pkg]:
    """Raise PatchError unless the output is a sound edit of `src`. Returns (report, output package)."""
    try:
        out = Pkg(out_bytes, limits)
    except InvalidDocument as exc:
        raise PatchError(f"the edited package is not valid: {exc.message}") from exc
    # every part is read back: CRC and size are checked by the zip layer
    surviving = [n for n in src.names if out.has(n)]
    added = [n for n in out.names if not src.has(n)]
    removed = [n for n in src.names if not out.has(n)]
    diff: list[str] = []
    try:
        for n in surviving:
            if src.read(n) != out.read(n):
                diff.append(n)
        for n in added:
            out.read(n)
    except (zipfile.BadZipFile, zlib.error) as exc:
        raise PatchError(f"the edited package is damaged: {exc}") from exc
    if surviving != [n for n in out.names if src.has(n)]:
        raise PatchError("the order of the surviving parts changed; the result was not delivered")
    editable: re.Pattern[str] = engine.EDITABLE
    changed = set(diff) | set(added) | set(removed)
    bad = sorted(n for n in changed if not editable.match(n))
    if bad:
        raise PatchError(f"parts changed that should not have: {bad}. The result was not delivered.")
    stray = sorted(changed - expected)
    if stray:
        raise PatchError(f"parts changed that the edit did not account for: {stray}. The result was not delivered.")
    for n in [*diff, *added]:
        if n.endswith(_XML):
            parse(out.read(n), n)
    for n in out.names:
        if n.endswith(".rels"):
            for _i, _t, tgt, mode in out.rels(owner(n)):
                if mode != "External" and not out.has(tgt):
                    raise PatchError(f"{n} points at a missing part {tgt}")
    ct = out.tree("[Content_Types].xml")
    defaults = {(d.get("Extension") or "").lower() for d in ct if isinstance(d.tag, str) and d.tag.endswith("}Default")}
    overrides = {o.get("PartName") for o in ct if isinstance(o.tag, str) and o.tag.endswith("}Override")}
    for n in out.names:
        if (
            n != "[Content_Types].xml"
            and not n.endswith("/")
            and "/" + n not in overrides
            and n.rsplit(".", 1)[-1].lower() not in defaults
        ):
            raise PatchError(f"{n} has no content type; the file would need repair in Office")
    for o in overrides:
        if o and o.lstrip("/") not in out.names:
            raise PatchError(f"[Content_Types].xml lists {o} but the part is gone")
    reopened = _reopen(fmt, out_bytes)
    report: dict[str, Any] = {
        "ok": True,
        "parts": len(src.names),
        "byte_identical": len(surviving) - len(diff),
        "changed": sorted(diff),
        "added": sorted(added),
        "removed": sorted(removed),
        "reopened_with": reopened,
        **engine.summary(out),
    }
    return report, out
