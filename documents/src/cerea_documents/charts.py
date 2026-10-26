"""Charts from a spec (ADR 0103 section 2.1): a PNG from matplotlib (Agg, a fixed clean style), or a
native chart. python-pptx writes native charts for PowerPoint and, through the same chart part, for
Word; openpyxl writes them for Excel (see make_xlsx). Everything here is deterministic."""

from __future__ import annotations

import io
from typing import Any

from .theme import Theme

DEFAULT_CM = (16.0, 9.0)
Chart = dict[str, Any]


def categories(c: Chart) -> list[str]:
    return [str(x) for x in c.get("categories") or []]


def wants_legend(c: Chart) -> bool:
    return bool(c.get("legend", len(c["series"]) > 1 or c["chart_type"] in ("pie", "doughnut")))


def describe(c: Chart, render: str, note: str | None = None) -> dict[str, Any]:
    """The chart's line in the result: what was asked and what was written."""
    out: dict[str, Any] = {
        "type": c["chart_type"],
        "title": c.get("title"),
        "series": len(c["series"]),
        "requested": c.get("render", "native"),
        "render": render,
    }
    if note:
        out["note"] = note
    return out


def render_png(c: Chart, theme: Theme, width_cm: float | None = None, height_cm: float | None = None) -> bytes:
    """A deterministic PNG of the chart: the same spec gives the same bytes."""
    import matplotlib

    matplotlib.use("Agg", force=True)
    from matplotlib import rc_context
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib.figure import Figure

    w, h = width_cm or DEFAULT_CM[0], height_cm or DEFAULT_CM[1]
    colors = ["#" + x for x in theme.palette]
    text = "#" + theme.colors["text"]
    muted = "#" + theme.colors["muted"]
    style = {
        "font.family": "DejaVu Sans",
        "font.size": 9,
        "text.color": text,
        "axes.labelcolor": text,
        "xtick.color": muted,
        "ytick.color": muted,
        "axes.edgecolor": "#" + theme.colors["border"],
        "axes.spines.top": False,
        "axes.spines.right": False,
        "axes.axisbelow": True,
        "savefig.facecolor": "white",
    }
    with rc_context(style):
        fig = Figure(figsize=(w / 2.54, h / 2.54), dpi=200, facecolor="white", layout="constrained")
        FigureCanvasAgg(fig)
        ax = fig.add_subplot(111)
        _draw(ax, c, colors, theme)
        if c.get("title"):
            ax.set_title(str(c["title"]), loc="left", fontweight="bold", fontsize=11, color=text)
        buf = io.BytesIO()
        fig.savefig(buf, format="png", metadata={"Software": None})
    return buf.getvalue()


def _draw(ax: Any, c: Chart, colors: list[str], theme: Theme) -> None:
    kind, series, cats = c["chart_type"], c["series"], categories(c)
    n = len(series)
    grid = "#" + theme.colors["border"]
    if kind in ("pie", "doughnut"):
        s = series[0]
        ax.pie(
            s["values"],
            labels=cats,
            colors=[colors[i % len(colors)] for i in range(len(s["values"]))],
            autopct="%1.0f%%",
            startangle=90,
            counterclock=False,
            wedgeprops={"width": 0.45, "edgecolor": "white"} if kind == "doughnut" else {"edgecolor": "white"},
            textprops={"fontsize": 8},
        )
        ax.set_aspect("equal")
        for side in ax.spines.values():
            side.set_visible(False)
        return
    if kind == "scatter":
        for i, s in enumerate(series):
            ax.scatter(
                [p[0] for p in s["points"]],
                [p[1] for p in s["points"]],
                s=14,
                color=colors[i % len(colors)],
                label=s["name"],
            )
        ax.grid(color=grid, linewidth=0.5)
    elif kind in ("column", "stacked", "bar"):
        horizontal = kind == "bar"
        idx = list(range(len(cats)))
        width = 0.8 / (1 if kind == "stacked" else n)
        base = [0.0] * len(cats)
        for i, s in enumerate(series):
            if kind == "stacked":
                pos, bottom = [float(x) for x in idx], base[:]
            else:
                pos, bottom = [x - 0.4 + width * (i + 0.5) for x in idx], [0.0] * len(cats)
            draw = ax.barh if horizontal else ax.bar
            kwargs = {"left": bottom} if horizontal else {"bottom": bottom}
            draw(pos, s["values"], width, color=colors[i % len(colors)], label=s["name"], **kwargs)
            if kind == "stacked":
                base = [b + v for b, v in zip(base, s["values"], strict=True)]
        (ax.set_yticks if horizontal else ax.set_xticks)(idx)
        (ax.set_yticklabels if horizontal else ax.set_xticklabels)(cats)
        if horizontal:
            ax.invert_yaxis()
        ax.grid(axis="x" if horizontal else "y", color=grid, linewidth=0.5)
    else:  # line, area
        idx = list(range(len(cats)))
        for i, s in enumerate(series):
            col = colors[i % len(colors)]
            if kind == "area":
                ax.fill_between(idx, s["values"], color=col, alpha=0.55, label=s["name"], linewidth=0)
                ax.plot(idx, s["values"], color=col, linewidth=1.4)
            else:
                ax.plot(idx, s["values"], color=col, linewidth=1.8, marker="o", markersize=3.5, label=s["name"])
        ax.set_xticks(idx)
        ax.set_xticklabels(cats)
        ax.grid(axis="y", color=grid, linewidth=0.5)
    if len(cats) > 8 and kind != "scatter":
        for lab in ax.get_xticklabels():
            lab.set_rotation(30)
            lab.set_ha("right")
    if c.get("x_title"):
        ax.set_xlabel(str(c["x_title"]))
    if c.get("y_title"):
        ax.set_ylabel(str(c["y_title"]))
    if wants_legend(c):
        ax.legend(frameon=False, fontsize=8)


# ---- native charts via python-pptx -----------------------------------------------------------------


def add_pptx_chart(shapes: Any, c: Chart, theme: Theme, x: int, y: int, cx: int, cy: int) -> Any:
    """Add a native chart to a python-pptx shape collection; returns the graphic frame."""
    from pptx.chart.data import CategoryChartData, XyChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Pt

    kind = c["chart_type"]
    types = {
        "bar": XL_CHART_TYPE.BAR_CLUSTERED,
        "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "stacked": XL_CHART_TYPE.COLUMN_STACKED,
        "line": XL_CHART_TYPE.LINE_MARKERS,
        "area": XL_CHART_TYPE.AREA,
        "pie": XL_CHART_TYPE.PIE,
        "doughnut": XL_CHART_TYPE.DOUGHNUT,
        "scatter": XL_CHART_TYPE.XY_SCATTER,
    }
    data: Any
    if kind == "scatter":
        data = XyChartData()
        for s in c["series"]:
            ser = data.add_series(s["name"])
            for px, py in s["points"]:
                ser.add_data_point(px, py)
    else:
        data = CategoryChartData()
        data.categories = categories(c)
        for s in c["series"]:
            data.add_series(s["name"], s["values"])
    frame = shapes.add_chart(types[kind], x, y, cx, cy, data)
    chart = frame.chart
    chart.font.size = Pt(10)
    chart.font.name = theme.body_font
    chart.font.color.rgb = RGBColor.from_string(theme.colors["text"])
    if c.get("title"):
        chart.has_title = True
        chart.chart_title.text_frame.text = str(c["title"])
        run = chart.chart_title.text_frame.paragraphs[0].runs[0]
        run.font.size, run.font.bold = Pt(12), True
    else:
        chart.has_title = False
    chart.has_legend = wants_legend(c)
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    if kind not in ("pie", "doughnut"):
        for axis, key in ((chart.category_axis, "x_title"), (chart.value_axis, "y_title")):
            if c.get(key):
                axis.has_title = True
                axis.axis_title.text_frame.text = str(c[key])
        chart.value_axis.major_gridlines.format.line.color.rgb = RGBColor.from_string(theme.colors["border"])
    colors = [RGBColor.from_string(x_) for x_ in theme.palette]
    plot = chart.plots[0]
    if kind in ("pie", "doughnut"):
        for i, point in enumerate(plot.series[0].points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = colors[i % len(colors)]
    else:
        for i, ser in enumerate(plot.series):
            col = colors[i % len(colors)]
            if kind in ("line", "scatter"):
                ser.format.line.color.rgb = col
                if kind == "line":
                    ser.smooth = False
                ser.marker.format.fill.solid()
                ser.marker.format.fill.fore_color.rgb = col
                if kind == "scatter":
                    ser.format.line.fill.background()  # points only
            else:
                ser.format.fill.solid()
                ser.format.fill.fore_color.rgb = col
    book = chart.part.chart_workbook  # XlsxWriter stamps 'now' into the embedded workbook: fix it
    book.update_from_xlsx_blob(normalise_xlsx(bytes(book.xlsx_part.blob)))
    return frame


def native_part_for_word(c: Chart, theme: Theme) -> tuple[bytes, bytes]:
    """The chart XML and its embedded workbook, as python-pptx writes them (Word reads the same chart
    part). Built on a scratch presentation that is thrown away."""
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    frame = add_pptx_chart(slide.shapes, c, theme, 0, 0, 5_760_000, 3_240_000)
    part = frame.chart.part
    return bytes(part.blob), bytes(part.chart_workbook.xlsx_part.blob)


def normalise_xlsx(blob: bytes) -> bytes:
    """Fixed timestamps in an embedded workbook (XlsxWriter stamps 'now')."""
    from .determinism import fix_zip

    return fix_zip(blob, core_part="docProps/core.xml")
