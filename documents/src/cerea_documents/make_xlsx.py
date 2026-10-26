"""Create an .xlsx from a spec with openpyxl: cells (values, formulas, formats), column widths, freeze panes,
autofilter, tables, native charts and images. Formulas are written, never evaluated; the workbook asks Excel
to recalculate when it opens."""

from __future__ import annotations

import datetime as dt
import io
import re
from typing import Any

import openpyxl
from openpyxl.chart import AreaChart, BarChart, DoughnutChart, LineChart, PieChart, Reference, ScatterChart, Series
from openpyxl.chart.marker import DataPoint
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.table import Table, TableStyleInfo

from . import charts
from .determinism import FIXED, fix_zip
from .errors import InvalidRequest
from .images_ref import ImageSet
from .spec import split_ref
from .theme import Theme

_BAD_FORMULA = re.compile(r"(?i)\b(cmd|powershell|mshta|msexcel)\s*\||\\\\|file:|\[[^\]]*\.(xls|xlsx|xlsm|xlsb)\]")
_TABLE_STYLES = {
    "banded": "TableStyleMedium2",
    "grid": "TableStyleMedium15",
    "minimal": "TableStyleLight1",
    "plain": "TableStyleLight1",
}
_DATE_FMT = "yyyy-mm-dd"


def _argb(h: str) -> str:
    return "FF" + h.lstrip("#").upper()


def _formula(ref: str, f: str) -> str:
    f = f if f.startswith("=") else "=" + f
    if _BAD_FORMULA.search(f):
        raise InvalidRequest(
            f"cell {ref}: this formula refers to an external file or a command, which is not accepted; "
            "use references inside the workbook",
            reason="schema",
        )
    return f


class XlsxBuilder:
    def __init__(self, spec: dict[str, Any], images: ImageSet, theme: Theme, template: bytes | None) -> None:
        self.spec, self.images, self.t = spec, images, theme
        self.notes: list[str] = []
        self.charts: list[dict[str, Any]] = []
        self.has_template = template is not None
        if template:
            self.wb = openpyxl.load_workbook(io.BytesIO(template))
            self.wb.defined_names.clear()
            if hasattr(self.wb, "template"):
                self.wb.template = False
            self.old = list(self.wb.worksheets)
        else:
            self.wb = openpyxl.Workbook()
            self.old = list(self.wb.worksheets)

    def font(self, **kw: Any) -> Font | None:
        if self.has_template and not (kw or self.t.explicit & {"body_font", "body_size"}):
            return None
        base: dict[str, Any] = {}
        if self.t.uses("body_font"):
            base["name"] = self.t.body_font
        if self.t.uses("body_size"):
            base["size"] = self.t.body_size
        return Font(**{**base, **kw})

    def build(self) -> tuple[bytes, dict[str, Any]]:
        sheets = self.spec["sheets"]
        used: list[Any] = []
        for sh in sheets:
            ws = self.wb.create_sheet(sh["name"])
            used.append(ws)
            self.sheet(ws, sh)
        for ws in self.old:
            self.wb.remove(ws)
        self.wb.active = 0
        p = self.wb.properties
        p.creator = p.lastModifiedBy = "Cerea"
        p.title = str(self.spec.get("title") or "")
        p.created = p.modified = FIXED
        self.wb.calculation = CalcProperties(fullCalcOnLoad=True)
        out = io.BytesIO()
        self.wb.save(out)
        return fix_zip(out.getvalue()), {
            "charts": self.charts,
            "notes": self.notes,
            "sheets": [{"name": s["name"], "used_range": ws.dimensions} for s, ws in zip(sheets, used, strict=True)],
        }

    # ---- one sheet ------------------------------------------------------------------------------

    def put(self, ws: Any, ref: str, raw: Any, **fmt: Any) -> None:
        cell = ws[ref]
        value = raw
        if isinstance(raw, dict) and raw.get("text") is not None:
            value = raw["text"]
        if "formula" in fmt:
            value = _formula(ref, fmt.pop("formula"))
        elif "date" in fmt:
            d = fmt.pop("date")
            try:
                value = dt.datetime.fromisoformat(d) if "T" in d or " " in d else dt.date.fromisoformat(d)
            except ValueError as exc:
                raise InvalidRequest(
                    f"cell {ref}: date must be ISO 8601 such as 2026-01-31 ({exc})", reason="schema"
                ) from exc
            fmt.setdefault("number_format", _DATE_FMT)
        elif isinstance(value, str) and value.startswith("="):
            value = _formula(ref, value)
        cell.value = value
        style = fmt.pop("style", None)
        if style:
            if style not in self.wb.named_styles:
                raise InvalidRequest(
                    f"cell {ref}: the named style {style!r} is not in the workbook; available: {sorted(self.wb.named_styles)}",
                    reason="schema",
                )
            cell.style = style
        kw = {k: fmt[k] for k in ("bold", "italic") if fmt.get(k)}
        if fmt.get("color"):
            kw["color"] = _argb(fmt["color"])
        f = self.font(**kw) if not (style and not kw) else None
        if f is not None:
            cell.font = f
        if fmt.get("fill"):
            cell.fill = PatternFill("solid", start_color=_argb(fmt["fill"]), end_color=_argb(fmt["fill"]))
        if fmt.get("align") or fmt.get("wrap"):
            cell.alignment = Alignment(
                horizontal=fmt.get("align"),
                wrap_text=bool(fmt.get("wrap")),
                vertical="top" if fmt.get("wrap") else None,
            )
        if fmt.get("number_format"):
            cell.number_format = fmt["number_format"]
        elif isinstance(value, dt.date | dt.datetime):
            cell.number_format = _DATE_FMT

    def sheet(self, ws: Any, sh: dict[str, Any]) -> None:
        t = self.t
        c0, r0 = split_ref(sh.get("start", "A1"))
        for ri, row in enumerate(sh.get("rows") or []):
            for ci, raw in enumerate(row):
                if raw is None or raw == "":
                    continue
                ref = f"{get_column_letter(c0 + ci)}{r0 + ri}"
                head = bool(sh.get("header")) and ri == 0
                fmt: dict[str, Any] = {}
                if head:
                    fmt = {"bold": True, "fill": t.colors["table_header_fill"], "color": t.colors["table_header_text"]}
                if isinstance(raw, dict):
                    fmt |= {k: v for k, v in raw.items() if k in ("bold", "align", "fill", "color")}
                self.put(ws, ref, raw, **fmt)
        for c in sh.get("cells") or []:
            fmt = {k: v for k, v in c.items() if k not in ("ref", "value")}
            self.put(ws, c["ref"], c.get("value"), **fmt)
        self.widths(ws, sh)
        if sh.get("freeze"):
            ws.freeze_panes = sh["freeze"]
        for rng in sh.get("merge") or []:
            ws.merge_cells(rng)
        self.tables(ws, sh)
        if sh.get("autofilter"):
            ws.auto_filter.ref = sh["autofilter"]
        for ch in sh.get("charts") or []:
            self.chart(ws, ch)
        for im in sh.get("images") or []:
            img = self.images.get(im["image"])
            xl = XlImage(io.BytesIO(img.data))
            self.size_image(xl, img.info.width, img.info.height, im.get("width_px"), im.get("height_px"))
            ws.add_image(xl, im["at"])

    @staticmethod
    def size_image(xl: Any, w0: int, h0: int, w: int | None, h: int | None) -> None:
        if w is None and h is None:
            w = min(w0, 640)
        if w is None:
            w = round((h or 1) * w0 / h0)
        if h is None:
            h = round(w * h0 / w0)
        xl.width, xl.height = w, h

    def widths(self, ws: Any, sh: dict[str, Any]) -> None:
        wanted = sh.get("column_widths") or {}
        longest: dict[int, int] = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    text = str(cell.value) if not str(cell.value).startswith("=") else "x" * 10
                    longest[cell.column] = max(
                        longest.get(cell.column, 0), max((len(x) for x in text.split("\n")), default=0)
                    )
        for col, n in longest.items():
            letter = get_column_letter(col)
            ws.column_dimensions[letter].width = wanted.get(letter, min(max(n * 1.15 + 2, 8), 60))
        for letter, w in wanted.items():
            ws.column_dimensions[letter].width = w

    def tables(self, ws: Any, sh: dict[str, Any]) -> None:
        for tb in sh.get("tables") or []:
            cells = ws[tb["range"]]
            header = (
                [c.value for c in (cells[0] if isinstance(cells[0], tuple) else cells)]
                if isinstance(cells, tuple)
                else [cells.value]
            )
            if any(not isinstance(h, str) or not h.strip() or h.startswith("=") for h in header) or len(
                set(header)
            ) != len(header):
                raise InvalidRequest(
                    f"table {tb['name']!r}: the first row of {tb['range']} must be unique, non-empty header text "
                    f"(found {header}); write the headers first",
                    reason="schema",
                )
            if sh.get("autofilter") == tb["range"]:
                raise InvalidRequest(
                    f"table {tb['name']!r} already has filter buttons; drop `autofilter` for the same range",
                    reason="schema",
                )
            table = Table(displayName=tb["name"], ref=tb["range"])
            style = self.t.table_style
            table.tableStyleInfo = TableStyleInfo(name=_TABLE_STYLES[style], showRowStripes=style == "banded")
            ws.add_table(table)

    # ---- charts ---------------------------------------------------------------------------------

    def chart(self, ws: Any, ch: dict[str, Any]) -> None:
        c1, r1, c2, r2 = _box(ch["data_range"])
        if c2 <= c1 or r2 <= r1:
            raise InvalidRequest(
                f"chart {ch.get('title') or ch['chart_type']!r}: data_range {ch['data_range']} needs a category column, "
                "at least one series column and a header row",
                reason="schema",
            )
        w, h = ch.get("width_cm", 16.0), ch.get("height_cm", 9.0)
        if ch.get("render", "native") == "image":
            data = self._cells_as_chart(ws, ch, c1, r1, c2, r2)
            png = charts.render_png(data, self.t, w, h)
            xl = XlImage(io.BytesIO(png))
            xl.width, xl.height = round(w / 2.54 * 96), round(h / 2.54 * 96)
            ws.add_image(xl, ch["anchor"])
            self.charts.append(charts.describe(data, "image"))
            return
        kind = ch["chart_type"]
        chart: Any = {
            "bar": BarChart, "column": BarChart, "stacked": BarChart, "line": LineChart, "area": AreaChart,
            "pie": PieChart, "doughnut": DoughnutChart, "scatter": ScatterChart,
        }[kind]()  # fmt: skip
        if isinstance(chart, BarChart):
            chart.type = "bar" if kind == "bar" else "col"
            if kind == "stacked":
                chart.grouping, chart.overlap = "stacked", 100
        colors = self.t.palette
        if kind == "scatter":
            xs = Reference(ws, min_col=c1, min_row=r1 + 1, max_row=r2)
            for i, col in enumerate(range(c1 + 1, c2 + 1)):
                s = Series(
                    Reference(ws, min_col=col, min_row=r1 + 1, max_row=r2),
                    xs,
                    title_from_data=False,
                    title=str(ws.cell(r1, col).value),
                )
                s.marker.symbol, s.marker.size = "circle", 6
                s.marker.graphicalProperties.solidFill = colors[i % len(colors)]
                s.graphicalProperties.line.noFill = True
                chart.series.append(s)
        else:
            last = c1 + 1 if kind in ("pie", "doughnut") else c2
            chart.add_data(Reference(ws, min_col=c1 + 1, max_col=last, min_row=r1, max_row=r2), titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=c1, min_row=r1 + 1, max_row=r2))
            for k, s in enumerate(chart.series):
                colour = colors[k % len(colors)]
                if kind in ("pie", "doughnut"):
                    for j in range(r2 - r1):
                        pt = DataPoint(idx=j)
                        pt.graphicalProperties.solidFill = colors[j % len(colors)]
                        s.dPt.append(pt)
                elif kind == "line":
                    s.graphicalProperties.line.solidFill = colour
                    s.smooth = False
                else:
                    s.graphicalProperties.solidFill = colour
        if ch.get("title"):
            chart.title = ch["title"]
        if not isinstance(chart, PieChart | DoughnutChart):
            chart.x_axis.delete = chart.y_axis.delete = False  # openpyxl 3.1 hides axes unless told otherwise
            if ch.get("x_title"):
                chart.x_axis.title = ch["x_title"]
            if ch.get("y_title"):
                chart.y_axis.title = ch["y_title"]
        n_series = len(chart.series)
        if not ch.get("legend", n_series > 1 or kind in ("pie", "doughnut")):
            chart.legend = None
        else:
            chart.legend.position = "b"
        chart.width, chart.height = w, h
        ws.add_chart(chart, ch["anchor"])
        self.charts.append(charts.describe({**ch, "series": [0] * n_series}, "native"))

    def _cells_as_chart(self, ws: Any, ch: dict[str, Any], c1: int, r1: int, c2: int, r2: int) -> dict[str, Any]:
        """The plotted cells as a chart spec, for an image render. Needs literal values: formulas are not evaluated."""

        def num(r: int, c: int) -> float:
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.startswith("="):
                raise InvalidRequest(
                    f"chart {ch.get('title') or ch['chart_type']!r}: {ws.cell(r, c).coordinate} holds a formula, and an image "
                    'chart is drawn from values; use render "native" (Excel computes it) or write the values',
                    reason="schema",
                )
            if isinstance(v, bool) or not isinstance(v, int | float):
                raise InvalidRequest(
                    f"chart {ch.get('title') or ch['chart_type']!r}: {ws.cell(r, c).coordinate} is {v!r}, not a number",
                    reason="schema",
                )
            return float(v)

        rows = range(r1 + 1, r2 + 1)
        cols = range(c1 + 1, (c1 + 1 if ch["chart_type"] in ("pie", "doughnut") else c2) + 1)
        base = {k: v for k, v in ch.items() if k not in ("data_range", "anchor")}
        if ch["chart_type"] == "scatter":
            series: list[dict[str, Any]] = [
                {"name": str(ws.cell(r1, c).value), "points": [[num(r, c1), num(r, c)] for r in rows]} for c in cols
            ]
            return {**base, "series": series}
        cats = [ws.cell(r, c1).value for r in rows]
        series = [{"name": str(ws.cell(r1, c).value), "values": [num(r, c) for r in rows]} for c in cols]
        return {**base, "categories": cats, "series": series}


def _box(rng: str) -> tuple[int, int, int, int]:
    a, _, b = rng.partition(":")
    ca, ra = split_ref(a)
    cb, rb = split_ref(b or a)
    return ca, ra, cb, rb


def build_xlsx(
    spec: dict[str, Any], images: ImageSet, theme: Theme, template: bytes | None
) -> tuple[bytes, dict[str, Any]]:
    return XlsxBuilder(spec, images, theme, template).build()
