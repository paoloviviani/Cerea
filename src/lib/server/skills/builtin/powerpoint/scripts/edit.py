"""edit.py: inventory and text replacement for an existing deck, in Cerea's sandbox.

Distilled for Cerea from pptx-from-layouts (MIT, Tristan McInnis). Small text
fixes only — reorder slides, retype paragraphs. Layout changes, added or
removed slides, or churn across more than ~30% of slides belong to
`generate.py` regeneration, not editing. No argv: import this file and call
the functions with paths.
"""

from __future__ import annotations

from pptx import Presentation


def inventory(path: str) -> dict:
	"""Every text-bearing shape on every slide, addressed as slide-N / shape-M.

	Paragraph entries carry an `id` (`p-<index>`) and their current text;
	`replace` consumes changes keyed exactly like this output.
	"""
	prs = Presentation(path)
	result: dict = {}
	for s, slide in enumerate(prs.slides):
		shapes: dict = {}
		for m, shape in enumerate(slide.shapes):
			if not shape.has_text_frame:
				continue
			paragraphs = []
			for index, paragraph in enumerate(shape.text_frame.paragraphs):
				text = "".join(run.text for run in paragraph.runs) or paragraph.text
				paragraphs.append({"id": f"p-{index}", "text": text})
			shapes[f"shape-{m}"] = {
				"name": shape.name,
				"is_placeholder": shape.is_placeholder,
				"placeholder_idx": shape.placeholder_format.idx if shape.is_placeholder else None,
				"paragraphs": paragraphs,
			}
		if shapes:
			result[f"slide-{s + 1}"] = shapes
	return result


def replace(path: str, changes: dict, out_path: str) -> dict:
	"""Apply text edits: {\"slide-2\": {\"shape-3\": {\"paragraphs\": [{\"id\": \"p-0\", \"text\": \"Q2 2026\"}]}}}.

	Only the paragraphs listed change, and only their text: the first run
	keeps its formatting, later runs in the paragraph are cleared.
	"""
	prs = Presentation(path)
	applied = 0
	missing = []
	for s, slide in enumerate(prs.slides):
		slide_changes = changes.get(f"slide-{s + 1}")
		if not slide_changes:
			continue
		shapes = list(slide.shapes)
		for shape_key, shape_changes in slide_changes.items():
			if not str(shape_key).startswith("shape-"):
				missing.append(f"slide-{s + 1}/{shape_key}")
				continue
			index = int(str(shape_key).split("-", 1)[1])
			if index >= len(shapes):
				missing.append(f"slide-{s + 1}/{shape_key}")
				continue
			shape = shapes[index]
			if not shape.has_text_frame:
				missing.append(f"slide-{s + 1}/{shape_key}")
				continue
			paragraphs = shape.text_frame.paragraphs
			for change in shape_changes.get("paragraphs", []):
				p_index = int(str(change.get("id", "p-0")).split("-", 1)[1])
				if p_index >= len(paragraphs):
					missing.append(f"slide-{s + 1}/{shape_key}/{change.get('id')}")
					continue
				paragraph = paragraphs[p_index]
				text = str(change.get("text", ""))
				if paragraph.runs:
					paragraph.runs[0].text = text
					for run in paragraph.runs[1:]:
						run.text = ""
				else:
					paragraph.text = text
				applied += 1
	prs.save(out_path)
	return {"out_path": out_path, "applied": applied, "missing": missing}


def reorder(path: str, order: list[int], out_path: str) -> dict:
	"""Write the slides in a new order; `order` is 1-based."""
	prs = Presentation(path)
	xml = prs.slides._sldIdLst
	id_elements = list(xml)
	count = len(id_elements)
	if sorted(order) != list(range(1, count + 1)):
		raise ValueError(f"order must be a permutation of 1..{count}, got {order}")
	for element in id_elements:
		xml.remove(element)
	for index in order:
		xml.append(id_elements[index - 1])
	prs.save(out_path)
	return {"out_path": out_path, "slides": count}
