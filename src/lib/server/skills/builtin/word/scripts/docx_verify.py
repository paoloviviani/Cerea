"""docx_verify - check an edit made with docx_edit / docx_tables, and show what it did.

    import docx_verify as dv
    rep = dv.verify("original.docx", "edited.docx")   # raises PatchError unless sound
    print("\n".join(rep["text_diff"]))

Licence: Apache-2.0, as the rest of this repository. Original work.
"""

import difflib

from lxml import etree

from docx_edit import (P, PatchError, _content_parts, _main, _Pkg, _ptext, _rels, _resolve, w)


def outline_text(o, width=70):
    """The outline as compact text, for printing."""
    cut = lambda s: s if len(s) <= width else s[: width - 3] + "..."
    lines = []
    for b in o["blocks"]:
        if b["type"] == "table":
            lines.append("%s table %dx%d [%s]" % (b["id"], b["rows"], b["cols"], b["style"]))
            lines += ["    r%d: %s" % (i, cut(" | ".join(r).replace("\n", " / "))) for i, r in enumerate(b["cells"])]
        else:
            lines.append("%s [%s] %s" % (b["id"], b["style"], cut(b["text"]) or "(empty)"))
    for kind in ("headers", "footers"):
        for part, ps in o[kind].items():
            lines.append("%s: %s" % (part, cut(" / ".join(x for x in ps if x))))
    lines.append("placeholders: %s; comments: %d; tracked: %s" % (o["placeholders"], o["comments"], o["tracked_changes"]))
    return "\n".join(lines)


def _texts(pkg):
    return {n: [_ptext(p) for p in pkg.tree(n).iter(P)] for n in _content_parts(pkg)}


def diff_text(src, dst):
    """Lines describing what changed in the text (deleted-by-track-changes text
    is hidden, tracked insertions are shown), grouped by part: "- old" / "+ new"."""
    a, b = _texts(_Pkg(src)), _texts(_Pkg(dst))
    out = []
    for part in sorted(set(a) | set(b)):
        d = [x for x in difflib.unified_diff(a.get(part, []), b.get(part, []), lineterm="", n=0)
             if not x.startswith(("---", "+++", "@@"))]
        if d:
            out += [part + ":"] + [x[0] + " " + x[1:] for x in d]
    return out


def verify(src, dst, changed_parts=None):
    """Raise PatchError unless `dst` is a sound edit of `src`: valid zip (CRC),
    same parts in the same order, only body/header/footer parts (or the
    `changed_parts` you pass) differ, every other part byte-identical, changed
    XML well-formed, relationships intact, and python-docx reopens it. Returns
    {"parts", "byte_identical", "changed", "text_diff", "tracked"}."""
    a, b = _Pkg(src), _Pkg(dst)
    if a.names != b.names:
        raise PatchError("the part list or order differs from the source: lost %s, added %s" % (
            sorted(set(a.names) - set(b.names)), sorted(set(b.names) - set(a.names))))
    diff = [n for n in a.names if a.data[n] != b.data[n]]
    allowed = set(changed_parts) if changed_parts is not None else set(_content_parts(a))
    if set(diff) - allowed:
        raise PatchError("parts changed that should not have: %s. Do not deliver this file; redo the edit with docx_edit "
                         "instead of python-docx save()." % sorted(set(diff) - allowed))
    for n in diff:
        if n.endswith((".xml", ".rels")):
            try:
                etree.fromstring(b.data[n])
            except etree.XMLSyntaxError as e:
                raise PatchError("%s is not well-formed XML: %s" % (n, e))
    for n in b.names:
        if n.endswith(".rels"):
            owner = n.rpartition("_rels/")[0] + n.rpartition("_rels/")[2][:-5]
            for typ, tgt, mode in _rels(b, owner):
                if mode != "External" and _resolve(owner, tgt) not in b.data:
                    raise PatchError("%s points at a missing part %s" % (n, tgt))
    try:
        import docx
    except ImportError:
        docx = None
    if docx is not None:
        try:
            docx.Document(dst)
        except Exception as e:
            raise PatchError("python-docx cannot reopen the result: %s" % e)
    root = b.tree(_main(b))
    return {"parts": len(a.names), "byte_identical": len(a.names) - len(diff), "changed": sorted(diff),
            "text_diff": diff_text(src, dst),
            "tracked": {"ins": len(root.findall(".//" + w("ins"))), "del": len(root.findall(".//" + w("del")))}}
